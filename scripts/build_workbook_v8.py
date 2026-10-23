#!/usr/bin/env python3
"""Rebuild the reviewable HARBORLIGHT8 workbook from unchanged version 7.

Requires openpyxl. By default LibreOffice recalculates every formula into cached
values in a temporary profile; --no-recalc is useful only for editing drafts.
The original workbook is never written. Run from any directory.
"""
from __future__ import annotations
import argparse
from copy import copy
from datetime import datetime
import hashlib
import json
import re
from pathlib import Path
import shutil
import subprocess
import tempfile
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/arkusze/HARBORLIGHT7-zasady.xlsx'
OUTPUT = ROOT / 'docs/arkusze/HARBORLIGHT8-zasady.xlsx'
SOURCE_SHA = 'bf49eef3b43ba81e0492c6eb187940c99177bbb8ae925ceca2b09ab0a2b9e945'
NAVY, TEAL, GOLD = '123140', '146D75', 'F0CE80'
LIGHT, BLUE, MUTED = 'F0F6F8', '135CAA', 'E6EDF0'
SOURCE_NOTE = 'Rozmowa po Excelu 7; zaakceptowany pakiet frontend v11 oraz późniejszy pakiet zatwierdzony „go”. Nowsza decyzja zastępuje poprzednią wersję tego samego widoku.'

# These are the accepted outcomes, rather than a chronology of abandoned UI proposals.
RULES = [
('Wersja', 'Wszystkie', 'V8 opiera się na niezmienionym pliku V7 i późniejszych uzgodnieniach. 17 decyzji D01–D17 nadal obowiązuje poza jawnymi nowszymi korektami.', 'Decyzje v7 są zachowane; źródła historyczne nie są nowymi wymaganiami.'),
('Start', 'Nowy gracz', 'Ratusz 1, Dom 1, 12 mieszkańców w domu, 0 pracowników, 0% szczęścia; każdy z 6 zasobów po 100. Ścieżki: wszystkie 100 odcinków na 0.', 'Normalny start nie przyznaje testowego magazynu ani 10 000 monet.'),
('Start', 'Podgląd testowy', 'W dotychczasowym podglądzie jednorazowo uzupełniamy centralę do limitu na rodzaj i łącznie oraz do minimum 10 000 monet; większe saldo i postęp zachowane.', 'Znacznik zapisu zapobiega ponownemu grantowi po odświeżeniu. To migracja piaskownicy, nie zasada startu.'),
('Start', 'Oddzielny nowy gracz', 'Link profile=new-player używa oddzielnego zapisu i normalnych wartości startowych. Nie kasuje bieżącej wyspy; późniejsze otwarcie kontynuuje zapis tego profilu.', 'Link i stan wdrożenia: docs/IMPLEMENTATION-LOG.md.'),
('Produkcja', 'Pięciu producentów', 'Lokalny magazyn = 200 × WŁASNY poziom producenta. Następny poziom ma 200 × (L+1). Wczytanie nie usuwa starszej nadwyżki.', 'Zastępuje zależność lokalnego bufora od ratusza.'),
('Ludzie', 'Osiem miejsc pracy', 'Limit pozostaje 12 × własny poziom zakładu; przydział jest ograniczony wspólną pulą wolnych osób w domu.', 'Wyświetlanie trzech liczb nie zmienia poziomu wyznaczającego limit.'),
('Frontend', 'Wszystkie budynki', 'U góry panelu status odpowiedni do funkcji. Producent: Producing / Waiting for workers / Storage full / Waiting for resources.', 'Status odświeża się według rzeczywistego stanu.'),
('Frontend', 'Produkcja', 'Producing wskazuje produkt, liczbę najbliższych sztuk i czas ich złożenia w lokalnym magazynie.', 'Uwzględnia rzeczywiste fazy pracy, dojście, głód, surowiec wejściowy i wolne miejsce.'),
('Frontend', 'Storage', 'Trzy wartości: In stock / Current cap. / Next level cap. Krótkie podpisy znajdują się pod każdą liczbą.', 'Np. 50 / 200 / 400; ostatnia liczba oznacza najbliższe ulepszenie, nie końcowy limit gry.'),
('Frontend', 'Workers', 'Trzy wartości: Assigned / Free at home / Max workers; małe +/− korzystają z rzeczywistej dostępności mieszkańców.', 'Np. 5 / 0 / 12. Osoba wracająca lub idąca nie jest wolna w domu.'),
('Frontend', 'Podatki', 'Usuń wszystkie informacje o podatkach z interfejsu: nagłówka, ratusza, kart budynków, zarządzania ludźmi i objaśnień. Sam rachunek podatków pozostaje bez zmiany.', 'Silnik nadal: 1 moneta/os./h po dotarciu, także przy przestoju; kaplica czasowo ×2. Excel dokumentuje tę regułę.'),
('Frontend', 'Kompaktowy układ', 'Level jest mniejszym polem po lewej; Storage po prawej, pod nim Workers. Po usunięciu podatków nie zostaje pole Taxes.', 'Najnowsze usunięcie podatków zastępuje wcześniejszy pomysł Taxes pod Level.'),
('Frontend', 'House', 'Housing capacity pokazuje populację, limit ratusza, miejsca do spania i przyrost/spadek. Opisy Housing capacity oraz House happiness są pod i.', 'Pula, terminy i przyciski domu pozostają widoczne.'),
('Frontend', 'Budowa i ludzie', 'Construction and queue oraz licznik występują tylko w Town Hall; Manage all workers tylko z House. Lokalne +/− zostają w miejscach pracy.', 'Residents pozostaje informacyjnym przeglądem z przejściem do domu.'),
('Frontend', 'Building upgrade', 'Pięciu producentów ma Upgrade, małe osobne czasy transportu i budowy oraz sąsiednią ramkę Premium z dwoma dotychczasowymi przyciskami.', 'Buy building package / Deliver local stock now. Opisy Coins/Instant pod i; bez osobnej dolnej sekcji Premium u producentów.'),
('Frontend', 'ETA', 'Etykieta budowy na wyspie ma dwa wiersze: Transport · ETA oraz Construction · ETA. Czasy uwzględniają aktualną obsadę i rozkład dostaw.', 'Bez zbierania materiałów i oczekiwania w kolejce; brak wykonalnej estymaty ma jawną przyczynę.'),
('Transport', 'Warehouse', 'Każdy magazynier ma widoczne przypisanie, towar, liczbę sztuk, źródło → cel i pozostały czas; osobno dojście i oczekiwanie.', 'Odwołanie nie przypisuje starego ładunku nowej osobie. Rozpoczęty ładunek dochodzi dokładnie raz.'),
('Priorytety', 'Towary i Construction', 'Można zaznaczyć maksymalnie 3 priorytety spośród pięciu produktów oraz Construction. Aktywne są podświetlone i numerowane 1/2/3.', 'Construction uczestniczy w tej samej puli i podziale kursów, nie jest dodatkowym czwartym priorytetem.'),
('Priorytety', 'Podział kursów', '1 priorytet: 8 kursów +2 losowe; 2: 5 pierwszego +3 drugiego +2 losowe; 3: 5/3/2. Bloki 10 rzeczywistych kursów powtarzają się bezterminowo.', 'Losowe wybierają dostępne pozostałe źródła; brak wybranego towaru nie zatrzymuje dostępnej pracy. Ratunkowe Food zachowuje ochronę.'),
('Priorytety', 'Zapis', 'Wybór, kolejność, pozycja cyklu i stan losowania są zapisywane. Zmiana priorytetów nie przerywa załadowanego kursu.', 'Nie ma wyłączenia po 10 kursach.'),
('Zadania', 'Jedno aktywne', 'Zawsze 3 oferty, ale wykonywać można 1 zadanie. Aktywna karta podświetlona w obu widokach, pozostałe wyszarzone bez startu ani wpłaty.', 'Pozostałe wymagania i nagrody można obejrzeć.'),
('Zadania', 'Postęp i migracja', 'Każdy wymagany produkt pokazuje dostarczone / wymagane na liście i po prawej. Zachowujemy stare wpłaty, termin oraz zapisany mnożnik.', 'Ukończenie lub wygaśnięcie zwalnia aktywne miejsce.'),
('Zadania', 'Norma', 'Wymagania pozostają 12 × własny poziom producenta × wydajność/os./h × wylosowane godziny; bez chwilowych bonusów.', 'Nie zmieniamy istniejących wymagań z powodu nowych podpisów czy stanu obsady.'),
('Animacje', 'Świat', 'Woda, dym, wiatrak, mewy, mieszkańcy, zwierzęta, chmury i budowa poruszają się przy włączonej opcji gry.', 'Świadome włączenie animacji ma pierwszeństwo przed systemowym ograniczeniem ruchu i jest zapisywane; stare pause/speed nie zamrażają obrazu.'),
('Frontend', 'Nagłówek', 'Marka, Level i lewy górny początek nawigacji mają oddzielne obszary; stary obraz i złote obramowanie nie przebijają.', 'Układ na komputerze i telefonie.'),
('Resources', 'Dolny pasek', 'Przełącza cały rząd Buildings na dokładnie 5 produktów: Wood, Planks, Stone, Grain, Food. Monety nie są szóstym produktem.', 'Buildings przywraca rząd budynków.'),
('Resources', 'Górny licznik', 'Mały numerek nad produktem to liczba faktycznie pracujących w jego zakładzie, po dojściu i bez blokady produkcji.', 'Głodny pracownik produkujący z połową wydajności nadal jest aktywny.'),
('Resources', 'Dolne liczby', 'Pod produktem: zapas u producenta / zapas tego samego produktu w magazynie głównym.', 'Zastępuje wolnych w domu / limit pracy tylko na kartach Resources; Workers zachowuje swoje trzy wartości.'),
('Resources', 'Kliknięcie', 'Wood→Lumber Camp; Planks→Sawmill; Stone→Quarry; Grain→Farm; Food→Windmill.', 'Ten sam panel i stan co kliknięcie budynku na wyspie.'),
('Zadania', 'Ukończenie', 'Ukończenie zadania otwiera popup z fajerwerkami; telefon otrzymuje krótką wibrację, jeśli urządzenie ją obsługuje.', 'Nagroda zadania jest przyznana przed opcjonalnym ekranem społecznościowym.'),
('Zadania', 'Share on Facebook', 'Po ukończeniu ekran Share your achievement: Share on Facebook lub Skip. Kliknięcie Share przyznaje 500 monet, jeden raz na ukończone zadanie.', 'NOWA odpowiedź użytkownika: 500 monet. Na razie bez integracji Facebook i bez deklarowania rzeczywistej publikacji.'),
('Zadania', 'Skip i powtórzenia', 'Skip zachowuje całą normalną nagrodę. Odświeżenie ani ponowne kliknięcie Share nie przyznaje ponownie premii za to samo zadanie.', 'Bonus za zadanie jest odrębny od starszych premii za grę/poziom/posty i polecenia.'),
('Atrakcje', 'Katalog i pasek', 'Decorations przełącza dolny pasek na 5 atrakcji z katalogu 20 dekoracji i wydarzeń; nie otwiera dotychczasowego okna listy.', 'Cyrk i teatr są wydarzeniami, nie budynkami produkcyjnymi.'),
('Atrakcje', 'Rotacja', 'Zestaw 5 ofert zmienia się co 24 h czasu zegarowego i zachowuje termin po odświeżeniu oraz offline. Pięć nowych rodzajów różni się od poprzedniej piątki.', 'Nie losujemy nowych cen ani okna czasu przy każdym kliknięciu karty.'),
('Atrakcje', 'Y i cena', 'Dla oferty Y jest całkowitą liczbą 5…100. Cena = Y × aktualny poziom ratusza L monet.', 'Y oferty jest zapisane do jej wymiany; koszt reaguje na osiągnięty poziom ratusza.'),
('Atrakcje', 'Szczęście', 'Jedno użycie daje Y × współczynnik przyzwyczajenia pp dla L mieszkańców przez 6 h. Szczęście nadal indywidualne, ograniczone do ±100, a populacja jest średnią.', 'Wybór implementacyjny: najmniej szczęśliwi niezamrożeni mieszkańcy; remisy według ID. Gdy populacja mniejsza od L, najwyżej tyle osób, ile istnieje.'),
('Atrakcje', 'Przyzwyczajanie', 'Każdy rodzaj ma własny licznik: 100%, 75%, 50%, 25%, 0% początkowej premii. Przy 0 zakup zablokowany bez pobrania monet.', 'Koszt nie maleje. Użycia innych atrakcji nie zwiększają tego licznika.'),
('Atrakcje', 'Przerwa i kumulacja', 'Po 24 h bez zakupu TEJ atrakcji wraca pełna skuteczność. Każde skuteczne użycie ma osobny zegar 6 h; nakładające się premie sumują się przed ograniczeniem H.', 'Dokładny składnik Y×współczynnik może być ułamkowy; prezentacja szczęścia jest zaokrąglona w górę do pełnego procentu.'),
('Ścieżki', 'Odcinki i poziomy', 'Jest 100 odcinków; start poziom 0, maksimum 5. Paths odsłania dolny poziomy pasek z poziomami i postępem, np. 37/100.', 'Nie tworzymy 100 osobnych budynków gospodarczych.'),
('Ścieżki', 'Brama poziomu', 'Aby ulepszać do poziomu N, wszystkie 100 odcinków musi mieć co najmniej N−1. Częściowe ulepszenie nie uruchamia premii N.', 'Ukończony poziom = minimum poziomów wszystkich odcinków.'),
('Ścieżki', 'Koszt i czas', 'Jedno ulepszenie jednego odcinka o poziom kosztuje 100 kamienia z centrali i działa po kliknięciu. Cały poziom: 10 000; 0→5: 50 000 kamienia.', 'Nie ma wymagania jednorazowego zgromadzenia 10 000 — gracz kupuje pojedyncze odcinki.'),
('Ścieżki', 'Szybkość', 'Ukończone poziomy 1/2/3/4/5: szybkość kursów +10/+20/+30/+40/+50%. Czas kursu = bazowy czas / (1+0,1×poziom).', 'Poziom 1: 300/1,1 s; nie odejmujemy po prostu 10% czasu.'),
('Ścieżki', 'Ładowność', 'Od ukończonego poziomu 3 dodatkowa pojemność: +10%, poziom 4 +20%, poziom 5 +30%; liczba sztuk całkowita w dół.', 'Ładunek własnego poziomu miejsca pracy oraz premia dróg to odrębne mnożniki.'),
('Ścieżki', 'Zakres', 'Premie działają na kursy tragarzy magazynu, targu i portu, także dostawy materiałów budowy.', 'Nie skracają rejsu statku, dojścia do pracy/powrotu, pobytu w latarni ani zegarów konsumpcji. Nie dają dodatkowego pasywnego H.'),
('Efekty', 'Town Hall', 'Na górze karty ratusza lista wszystkich aktywnych błogosławieństw, premii i kar: źródło, wartość, odbiorcy i czas albo warunek zakończenia.', 'Premie dróg są stałe, czasowe atrakcje mają własne terminy. Informacje o podatkach pozostają ukryte w UI.'),
('Efekty', 'Powiadomienia', 'Każde uruchomienie, zmiana i zakończenie efektu ma wpis w powiadomieniach; bez wielokrotnego wpisu co klatkę/odświeżenie.', 'Dotyczy wszystkich źródeł: domu, jedzenia/pracy, głodu/bezrobocia, przestoju, przeludnienia, kaplicy, latarni, atrakcji i dróg.'),
('Migracja', 'Ścieżki wcześniejszego zapisu', 'Nowy gracz ma wszystkie odcinki na 0. Wcześniej kupione odcinki ścieżek zachowują wydany kamień jako odpowiadające im odcinki poziomu 1.', 'Migracja zachowuje płatny postęp starej wyspy; nie przyznaje całego poziomu, jeśli nie ukończono wszystkich 100 odcinków.'),
('Stan', 'Język i zapis', 'Aktywny interfejs po angielsku. Zachowujemy istniejące zasoby, mieszkańców, budowy, ładunki i zadania; osobne profile nie nadpisują sobie zapisów.', 'Oryginalny Excel 7 pozostaje niezmieniony. Prawdziwe konta, OAuth, płatności i reklamy nadal wymagają osobnych integracji.'),
]

ATTRACTIONS = [
('flower-beds','Flower beds','Rabaty kwiatowe','Dekoracja'),('village-benches','Village benches','Ławki we wsi','Dekoracja'),
('lantern-strings','Lantern strings','Girlandy lampionów','Dekoracja'),('harbor-flags','Harbor flags','Flagi przystani','Dekoracja'),
('village-fountain','Village fountain','Fontanna wiejska','Dekoracja'),('garden-pond','Garden pond','Oczko ogrodowe','Dekoracja'),
('wooden-sculptures','Wooden sculptures','Drewniane rzeźby','Dekoracja'),('seaside-mural','Seaside mural','Mural nad morzem','Dekoracja'),
('lighthouse-garden','Lighthouse garden','Ogród przy latarni','Dekoracja'),('ornamental-sails','Ornamental sails','Ozdobne żagle','Dekoracja'),
('travelling-circus','Travelling circus','Wędrowny cyrk','Wydarzenie'),('theatre-troupe','Theatre troupe','Grupa teatralna','Wydarzenie'),
('folk-concert','Folk concert','Koncert ludowy','Wydarzenie'),('puppet-show','Puppet show','Teatr lalek','Wydarzenie'),
('harvest-festival','Harvest festival','Dożynki','Wydarzenie'),('fishing-contest','Fishing contest','Zawody wędkarskie','Wydarzenie'),
('storytelling-evening','Storytelling evening','Wieczór opowieści','Wydarzenie'),('village-dance','Village dance','Wiejska potańcówka','Wydarzenie'),
('fireworks-night','Fireworks night','Noc fajerwerków','Wydarzenie'),('harbor-fair','Harbor fair','Jarmark portowy','Wydarzenie')]


def sheet(wb, name, title, description, headers, rows, widths=None):
    s=wb.create_sheet(name)
    n=len(headers)
    s.merge_cells(start_row=1,start_column=1,end_row=1,end_column=n)
    s.merge_cells(start_row=2,start_column=1,end_row=2,end_column=n)
    s.cell(1,1,title);s.cell(2,1,description)
    s.row_dimensions[1].height=32;s.row_dimensions[2].height=64
    for r in (1,2):
        c=s.cell(r,1);c.font=Font(name='Calibri',size=18 if r==1 else 11,bold=r==1,color=GOLD if r==1 else NAVY)
        c.fill=PatternFill('solid',fgColor=NAVY if r==1 else LIGHT);c.alignment=Alignment(vertical='center',wrap_text=True)
    for col,h in enumerate(headers,1):
        c=s.cell(5,col,h);c.font=Font(name='Calibri',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=TEAL);c.alignment=Alignment(wrap_text=True,vertical='center')
    s.row_dimensions[5].height=34
    for r, values in enumerate(rows,6):
        for col,value in enumerate(values,1):
            c=s.cell(r,col,value);c.font=Font(name='Calibri',size=10,color=NAVY);c.fill=PatternFill('solid',fgColor=LIGHT if r%2==0 else 'FFFFFF');c.alignment=Alignment(vertical='top',wrap_text=True)
            c.border=Border(bottom=Side(style='hair',color='D4E3E7'))
        longest=max((len(str(v)) for v in values if v is not None),default=0)
        s.row_dimensions[r].height=min(150,max(32,17*(longest//65+1)))
    for c in range(1,n+1):s.column_dimensions[get_column_letter(c)].width=(widths or {}).get(c,20)
    s.freeze_panes='C6';s.auto_filter.ref=f'A5:{get_column_letter(n)}{5+len(rows)}'
    s.sheet_view.showGridLines=False;s.sheet_properties.pageSetUpPr.fitToPage=True
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0
    s.print_title_rows='1:5';s.print_options.horizontalCentered=True
    s.sheet_properties.tabColor=TEAL
    return s


def append_row(s, values, height=64):
    r=s.max_row+1
    for col,value in enumerate(values,1):
        target=s.cell(r,col,value);model=s.cell(r-1,col)
        target._style=copy(model._style);target.alignment=Alignment(vertical='top',wrap_text=True)
    s.row_dimensions[r].height=height
    if s.auto_filter.ref:s.auto_filter.ref=f'A5:{get_column_letter(max(s.max_column,len(values)))}{r}'
    return r


def build():
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_SHA,'Source version 7 changed; inspect before rebuilding.'
    w=load_workbook(SOURCE)
    history=w.copy_worksheet(w['Ustalenia - uwagi']);history.title='Rejestr historyczny v7';history.sheet_properties.tabColor='94A3A8'
    # Preserve source/comment cells and original history verbatim; only current rules change.
    current=[s for s in w if not s.title.startswith(('Twoje','Zrodlo','Rejestr')) and s.title!='Decyzje v7']
    for s in current:
        for row in s:
            for c in row:
                if isinstance(c.value,str) and c.data_type!='f':
                    c.value=c.value.replace('200 × Ratusz','200 × własny poziom producenta').replace('200×Ratusz','200×własny poziom producenta').replace('200×ratusz','200×własny poziom producenta').replace('200×poziom Ratusza','200×własny poziom producenta')
    start=w['Start']
    start['A1']='HARBORLIGHT V8 • pełne zasady i zmiany interfejsu'
    start['A2']='Źródło: niezmieniony HARBORLIGHT7-zasady.xlsx, 17 decyzji D01–D17 oraz WSZYSTKIE późniejsze zaakceptowane zmiany frontendu i gospodarki. Nowe arkusze: Zmiany v8, Frontend v8, Atrakcje, Atrakcje - kalkulator, Sciezki, Sciezki - transport, Efekty i notyfikacje, Sprawdzenia v8. Archiwalne źródła i 48 komentarzy użytkownika zachowano. Nowsze decyzje w Zmiany v8 mają pierwszeństwo.'
    start.row_dimensions[2].height=90
    start['B6']='12 budynków + asset Człowiek; własne bufory producentów, dotychczasowe limity pracy i najnowsze widoki bez podatków.'
    start['B7']='Pełny bazowy kurs 5 min; ścieżki przyspieszają kursy i od poziomu 3 zwiększają ładunek. Budynki 2T; latarnia 5T; kaplica 4T/7T/10T.'
    start['B9']='Limit 12 × własny poziom, produkcja 10/12/13/14…, 5 min do pracy. Podatki naliczane bez zmian, ale informacja o nich usunięta z UI.'
    start['B12']='Do trzech trwałych numerowanych priorytetów, w tym Construction. Podział kursów 8/2 losowe, 5/3/2 losowe lub 5/3/2. Ratunkowe Food zachowuje ochronę.'
    start['B21']='Zawsze 3 oferty, tylko 1 aktywne zadanie. Podświetlony postęp każdego produktu w obu widokach; pozostałe wyszarzone. Po sukcesie fajerwerki, wibracja telefonu i opcjonalny Share on Facebook za 500 monet.'
    start['B26']='Angielskie etykiety liczb, statusy, ETA transportu i budowy; bez podatków; Resources = zapas lokalny / centralny; zarządzanie budową tylko ratusz, pracownikami tylko dom.'
    start['C26']='Frontend v8 / Zmiany v8; stare notatki w kolumnach E/F są źródłowe.'
    start['B29']='Reguły w V8 są zatwierdzone. Prototyp frontendowy ma lokalny zapis; konta, OAuth, płatności i reklamy nie są integrowane. Końcowy commit, publikację i wyniki testów gry zapisujemy w docs/IMPLEMENTATION-LOG.md.'
    for title,desc,where in [
        ('Atrakcje','20 dekoracji/wydarzeń, 5 ofert co 24 h; Y×L monet; premia dla L osób na 6 h; przyzwyczajanie osobno na rodzaj.','Atrakcje / Atrakcje - kalkulator'),
        ('Ścieżki','100 odcinków 0–5, każdy krok 100 kamienia; pełny poziom daje +10% szybkości, od 3 dodatkową ładowność.','Sciezki / Sciezki - transport'),
        ('Efekty','Błogosławieństwa, premie i kary w powiadomieniach oraz u góry karty Town Hall; zmiany bez spamu.','Efekty i notyfikacje'),
        ('Nowy gracz / testy','Osobny link nowego gracza z normalnym startem; dotychczasowy podgląd ma jednorazowy pełny magazyn i minimum 10 000 monet.','Frontend v8 / Zmiany v8'),
        ('Archiwum','36 zakładek źródła zachowanych; dawne rejestry i wpisy źródłowe są historią. Aktualny rejestr poprawiono i rozszerzono.','Rejestr historyczny v7 / Rejestr historyczny v6'),
    ]:append_row(start,[title,desc,where],64)
    b=w['Budynki'];b['A1']='BUDYNKI V8 • aktualne reguły'
    b['A2']='12 budynków + asset Człowiek. Aktualne zasady: V7 + Zmiany v8. Kolumna K i komentarze są zachowaną historią użytkownika. Bufory producentów: 200 × własny poziom. Rachunek podatków zachowany, ich informacje usunięte z interfejsu. Czasy porównawcze zależą od parametrów: transport + budowa, bez zbierania i czekania w kolejce.'
    b['B6']='Wyznacza poziom gry i limity rozwoju; skarbiec monet. U góry karty pokazuje aktywne błogosławieństwa, premie i kary.'
    b['D6']='Maksymalna populacja 24×L, wyższy limit ulepszeń, L budów + L miejsc kolejki. Lokalny magazyn producenta zależy od jego własnego poziomu, nie od ratusza.'
    for r in range(6,19):
        b.cell(r,7,'Uzgodnione w V8; prototyp frontendowy. Szczegóły bieżącego wdrożenia i testów: docs/IMPLEMENTATION-LOG.md. Backend i integracje: Na pozniej.')
        v=b.cell(r,8).value or ''
        b.cell(r,8,v+' Premie/kary: powiadomienie przy aktywacji, zmianie i końcu oraz aktywna lista u góry Town Hall.')
    b['D7'].value+=' Manage all workers dostępne wyłącznie z domu. Cztery wskaźniki: populacja, limit ratusza, łóżka, przyrost/spadek; opisy pod i.'
    b['I13'].value+=' Podgląd każdej osoby: ładunek, źródło→cel i ETA. Trzy trwałe priorytety spośród produktów i Construction, podział kursów 8/2, 5/3/2 lub 5/3/2.'
    w['Czlowiek']['A2'].value+=' V8 dodaje czasowe atrakcje dla wybranych osób; podatki nadal naliczane, ich informacji nie pokazujemy w UI. Workers ma podpisane Assigned / Free at home / Max workers.'
    w['Czlowiek']['D14'].value+=' Ścieżki po ukończeniu całego poziomu zwiększają szybkość; od poziomu 3 również ładunek.'
    w['Czlowiek']['D15'].value+=' V8: wartość i opis podatku ukryte we wszystkich widokach gry; to zmiana prezentacji, nie rachunku.'
    append_row(w['Czlowiek'],['Atrakcje','Każdy zakup dodaje czasową premię wybranym mieszkańcom, nie całej średniej automatycznie.','Y×współczynnik pp dla L osób przez 6 h; L=aktualny poziom ratusza.','Przyzwyczajanie osobno; limit osobistego H nadal100; zamrożenie nowej osoby zachowane.','POTWIERDZONE V8'])
    append_row(w['Czlowiek'],['Informacja o efektach','Wszystkie aktywne premie i kary mają źródło, wartość, odbiorców i czas/warunek zakończenia.','Powiadomienia + góra Town Hall.','Wpis na zmianę stanu, nie co klatkę.','POTWIERDZONE V8'])
    t=w['Transport i priorytety'];t['A2']='V8: fizyczne kursy, drogi 0–5, trzy trwałe priorytety towarów/Construction i podgląd rzeczywistych tragarzy. Całość działa zegarowo także offline; rejs statku i dojście do pracy to osobne zegary.'
    t['B6']='Bazowo pełny kurs z powrotem 5 min. Ścieżki: czas = 5/(1+0,1×ukończony poziom). Jeden rodzaj na kurs. Ładunek z własnego poziomu magazynu/targu/portu, dodatkowo premia ścieżek od poziomu3.'
    t['C11']='Maksymalnie 3 numerowane priorytety, w tym Construction. Bez priorytetów najpierw centrala, potem budowy; ratunkowy Food ma ochronę.'
    t['B12']='ETA osobno dla transportu i budowy przy obecnej obsadzie, ładunkach i priorytetach. Etykieta świata: Transport · ETA / Construction · ETA.'
    for vals in [
        ['Wybrane priorytety','5 produktów i Construction w jednej puli; najwyżej3 aktywne, podświetlone i numerowane.','Construction może mieć numer1,2 lub3.','POTWIERDZONE V8'],
        ['Podział kursów','1:8+2 losowe; 2:5+3+2 losowe; 3:5+3+2.','Bloki rzeczywistych kursów powtarzają się bezterminowo. Niepełny kurs liczy się jako jeden.','POTWIERDZONE V8'],
        ['Pusty priorytet','Brak towaru nie zatrzymuje tragarza, jeśli jest inna wykonalna dostawa.','Zapis zachowuje wybór, licznik cyklu i stan losowania.','POTWIERDZONE V8'],
        ['Podgląd osób','Przydział, rodzaj i ilość ładunku, skąd→dokąd, czas końca; osobno droga/bezczynność.','Odwołany ładunek kończy kurs niezależnie; bez kopiowania na nową osobę.','POTWIERDZONE V8'],
        ['Ścieżki','100 odcinków×5 poziomów; korzyść dopiero po ukończeniu całej setki poziomu.','Szybkość +10%/poziom; ładunek od3: +10/+20/+30%. Zobacz Sciezki.','POTWIERDZONE V8'],
    ]:append_row(t,vals)
    # Append parameters without moving any v7 cell referenced by existing formulas.
    p=w['Parametry'];p['A1']='PARAMETRY V8 • zatwierdzone liczby i scenariusze'
    params=[
      ('local_capacity','Lokalna pojemność / własny poziom',200,'szt./L','Pięciu producentów; nie ratusz.'),
      ('road_sections','Liczba odcinków ścieżek',100,'odcinków','Wszystkie zaczynają na0.'),
      ('road_max','Maksymalny poziom ścieżek',5,'poziom','Kolejny poziom po ukończeniu wszystkich100 poprzedniego.'),
      ('road_cost','Koszt jednego odcinka o jeden poziom',100,'kamienia','Natychmiast po kliknięciu; koszt z centrali.'),
      ('road_speed','Bonus szybkości / ukończony poziom',0.1,'udział','Czas dzielony przez mnożnik szybkości.'),
      ('road_load','Bonus ładunku / poziom od3',0.1,'udział','Ukończone3/4/5: +10/+20/+30%.'),
      ('road_calculation','Ukończony poziom dróg w kalkulatorach',0,'0…5','Parametr scenariusza, NIE zmiana startu gry.'),
      ('attraction_min','Atrakcja: minimum Y',5,'monet/pp','KosztY×L; dokładna premiaY×przyzwyczajenie.'),
      ('attraction_max','Atrakcja: maksimum Y',100,'monet/pp','Y całkowite, zapisane w ofercie.'),
      ('attraction_hours','Czas premii atrakcji',6,'h','Osobny zegar każdego użycia.'),
      ('attraction_reset','Przerwa resetująca daną atrakcję',24,'h','Od jej ostatniego zakupu, oddzielnie na rodzaj.'),
      ('attraction_rotation','Rotacja pięciu ofert',24,'h','Działa zegarowo i nie resetuje się przy odświeżeniu.'),
      ('attraction_catalog','Katalog atrakcji',20,'rodzajów','10 dekoracji i10 wydarzeń.'),
      ('attraction_visible','Widoczne atrakcje',5,'kart','W dolnym pasku, zamiast starego okna.'),
      ('task_share','Share ukończenia zadania',500,'monet','Raz na ukończone zadanie; Skip nie zmienia normalnej nagrody. Bez integracjiFB.'),
      ('preview_coins','Minimum monet w testowym podglądzie',10000,'monet','Jednorazowa migracja; oddzielny nowy gracz bez grantu.'),
      ('priorities_max','Maksymalnie aktywnych priorytetów',3,'pozycje','Towary i Construction we wspólnej puli.'),
    ]
    refs={}
    for ident,label,value,unit,note in params:
        row=append_row(p,[ident,label,value,unit,'POTWIERDZONE V8',note],48);refs[ident]=f'Parametry!$C${row}'
        p.cell(row,3).font=Font(name='Calibri',color=BLUE,size=11)
        p.cell(row,3).number_format='0.0%' if ident in ('road_speed','road_load') else '0'
    roadref=refs['road_calculation'];loadref=refs['road_load'];speedref=refs['road_speed']
    # Adapt the CURRENT calculators to roads, default0 preserving all v7 examples.
    for s in (w['Produkcja'],w['Czas budowy'],w['Premium - budynek']):
        for row in s:
            for c in row:
                if c.data_type=='f' and c.value.startswith('=ROUNDDOWN(Parametry!$C$9*Parametry!$C$10^('):
                    inner=c.value[1:]
                    c.value=f'=ROUNDDOWN({inner}*(1+{loadref}*MAX(0,{roadref}-2)),0)'
    for row in range(6,w['Czas budowy'].max_row+1):
        w['Czas budowy'].cell(row,22,f'=U{row}*Parametry!$C$11/(1+{speedref}*{roadref})')
        w['Czas budowy'].cell(row,27,'Tylko ta budowa; 6 dostępnych nakarmionych tragarzy; poziom magazynu według F. Poziom ścieżek z Parametry; domyślnie0. Rywalizujące kursy i priorytety w grze mogą wydłużyć ETA.')
    w['Czas budowy']['A2']='74 warianty transportu i budowy. Młyn/drwale/tartak/kamieniołom koszt×2; farma bez zmian; kolejka+10%. Uwzględniono edytowalny ukończony poziom dróg (Parametry; domyślnie0), jego szybkość i ładunek. 6 dostępnych nakarmionych tragarzy. Bez zbierania i czekania; rzeczywisty scheduler może mieć dodatkowe konkurujące kursy.'
    w['Produkcja']['A2'].value+=' V8: lokalna pojemność =200×własny poziom producenta; kolumnaE uwzględnia poziom ścieżek ze scenariusza Parametry.'
    w['Premium - budynek']['A2'].value+=' Ładunek wyceny uwzględnia aktualny parametr ścieżek, ale budynek nadal jest gotowy po10s.'
    for r in range(6,w['Premium - transport'].max_row+1):
        base=w['Premium - transport'].cell(r,4).value
        if isinstance(base,(int,float)):w['Premium - transport'].cell(r,4,f'=ROUNDDOWN({base}*(1+{loadref}*MAX(0,{roadref}-2)),0)')
    w['Premium - transport']['A2'].value+=' V8: ładunek przykładów uwzględnia premię ścieżek z Parametry (domyślnie0); liczba kursów i cena przeliczają się razem.'
    q=w['Zadania latarni']
    q['A1']='ZADANIA LATARNI V8 • jedno aktywne, postęp i celebracja'
    q['C6']=(q['C6'].value or '')+' Jednocześnie wykonywane tylko1; pozostałe oferty wyszarzone.'
    q['C24']='Powiadomienia o aktywacji i mnożniku, wpłacie, ukończeniu i wygaśnięciu. Ukończenie: popup z fajerwerkami, krótka wibracja telefonu i opcjonalny ekran Share your achievement.'
    append_row(q,['Aktywność','POTWIERDZONE V8','Tylko jedno aktywne zadanie. Active i postęp każdego produktu w obu widokach; pozostałe oferty szare, bez startu/wpłaty.','Wymagania i nagrody można nadal obejrzeć.','Stare wpłaty, terminy i utrwalone mnożniki zachowane.'])
    append_row(q,['Bonus udostępnienia','POTWIERDZONE V8','Share on Facebook przyznaje500monet raz na ukończone zadanie; Skip zachowuje zwykłą nagrodę.','Bez integracjiFacebook; kliknięcie prototypowego przycisku, nie potwierdzenie publikacji.','Najnowsza odpowiedź użytkownika:500monet.'])
    social=w['Social i polecenia'];append_row(social,['Ukończenie zadania: Share on Facebook',500,'Kliknięcie opcjonalnego Share po rzeczywistym ukończeniu zadania; jeden raz na ID ukończonego zadania.','Bez integracjiFB. Skip nie odbiera nagrody. To osobna akcja od pozostałych kwot w tej tabeli.',f'={refs["task_share"]}/Parametry!$C$41'])
    w['Na pozniej']['B8'].value+=' Nowy wyjątekV8: lokalny przycisk Share ukończonego zadania, bonus500 raz/zadanie, bez integracjiFB.'
    w['Na pozniej']['D8']='Starsze integracje i limity nadal odłożone; bonus500 za kliknięcie zadania jest jawną nową regułą prototypu.'
    w['Blogoslawienstwa']['A2'].value+=' V8: aktywne efekty widoczne także na górze Town Hall; rozpoczęcie, zmiana i koniec w powiadomieniach. Samych informacji podatkowych nie pokazujemy w UI.'
    for r in range(6,w['Blogoslawienstwa'].max_row+1):
        w['Blogoslawienstwa'].cell(r,2,(w['Blogoslawienstwa'].cell(r,2).value or '')+' Powiadomienia + aktywna lista ratusza.')
    registry=w['Ustalenia - uwagi'];registry['A1']='USTALENIA V8 • bieżący rejestr reguł'
    registry['A2']='Zachowane217wpisówV7, poprawione według nowych zasad, oraz pełny pakietV8. Rejestr historyczny v7 zachowuje poprzednie wartości. Źródła historyczne i komentarze nie stanowią nowych decyzji; najnowszy wpis ma pierwszeństwo.'
    registry['E106']='Najwyżej3 trwałe priorytety spośród produktów i Construction, numery1/2/3. Domyślnie centrala przed budowami; ratunkoweFood chronione. Bloki10kursów:8/2losowe,5/3/2losowe,5/3/2.'
    registry['E114']='Panel i etykieta budowy pokazują dostarczone/wymagane oraz osobne Transport · ETA / Construction · ETA. Aktualna obsada, ładunki, ścieżki i priorytety wpływają na wycenę czasu.'
    for i,(area,assets,rule,note) in enumerate(RULES,1):
        append_row(registry,[f'V8-{i:03}',area,assets,'POTWIERDZONE V8',rule,note+' Źródło: '+SOURCE_NOTE],80)
    changes=sheet(w,'Zmiany v8','ZMIANY V8 • komplet późniejszych uzgodnień',SOURCE_NOTE+' Nowsze postanowienie o usunięciu podatków i stanach magazynów na Resources zastępuje wcześniejszy układ.',
                  ['ID','Obszar','Asset / widok','Ostateczna zasada','Uwagi / źródło'],
                  [(f'V8-{i:03}',a,b,c,d) for i,(a,b,c,d) in enumerate(RULES,1)],{1:14,2:21,3:27,4:85,5:80})
    ui=sheet(w,'Frontend v8','FRONTEND V8 • finalne zachowanie interfejsu','Angielskie krótkie etykiety; polski opis zasad. Wcześniejsze pola Taxes i dolne free/max na Resources są zastąpione późniejszą instrukcją. Rachunek podatków zachowany wyłącznie w silniku.',
      ['Widok','Element / podpis','Wartości i zachowanie','Reguła / wyjątek'],[
      ('Panel producenta','Level','Mniejsze pole po lewej.','Bez pola Taxes pod spodem.'),
      ('Panel producenta','Storage','In stock / Current cap. / Next level cap.','Stan lokalny /200×L /200×(L+1).'),
      ('Miejsce pracy','Workers','Assigned / Free at home / Max workers','Trzy podpisy pod liczbami, np.5/0/12.'),
      ('Miejsce pracy','+/−','Zmiana przydziału z faktycznie dostępnych osób.','Droga5min; brak wolnych blokujeplus.'),
      ('Każdy budynek','Status','Produkcja/blokada albo status odpowiedni do roli.','Pod status produkcji produkt+ilość+ETA lokalnego złożenia.'),
      ('Wszystkie widoki','Podatki','Brak wyświetlania podatków i ich opisu.','Nie zmienia naliczania ani premii kaplicy.'),
      ('House','Housing capacity','Populacja / limit ratusza / łóżka / przyrost-spadek.','Opisy pod i.'),
      ('House','House happiness','Pula i zegar widoczne; opis pod i.','Manage all workers otwierane tylko tutaj.'),
      ('Town Hall','Construction and queue','Jedyny panel z licznikiem i przyciskiem zarządzania kolejką.','Inne budynki zachowują własny Upgrade.'),
      ('Town Hall','Active effects','Aktywne błogosławieństwa, premie i kary na górze karty.','Źródło,wartość,odbiorcy,czas/warunek; bez podatków.'),
      ('Producent','Building upgrade','Upgrade + dwa małe czasy + sąsiednia ramkaPremium z2przyciskami.','InformacjeCoins/Instant pod i; brak dolnej osobnej sekcji.'),
      ('Etykieta budynku','Transport · ETA / Construction · ETA','Dwa osobne wiersze obok budowanego obiektu.','Rzeczywiste ustawienia dostaw; brak estymaty wyjaśniony.'),
      ('Warehouse','Warehouse workers','Każda osoba,ładunek,ilość,źródło→cel,pozostałyczas.','Dochodzący i bezczynny odróżnieni.'),
      ('Warehouse','Transport priority','5produktów orazConstruction, najwyżej3 wybrane, numery1/2/3.','8/2,5/3/2,5/3/2 powtarzane bezterminowo.'),
      ('Resources','Pięć kart produktów','Wood / Planks / Stone / Grain / Food.','Klik otwiera właściwego producenta; Buildings wraca.'),
      ('Resources','Górny numerek','Faktycznie pracujący producenta.','Po dojściu,bezblokady; głodnyprodukujący liczy się.'),
      ('Resources','Dolna para','Stan lokalny / stan centrali tego samego produktu.','Np.50/800; już NIE free/max.'),
      ('Tasks','Active / postęp','Jedno aktywne,zielono-złote; postęp każdego produktu w obu widokach.','Pozostałe2szare,bezstartu/wpłaty.'),
      ('Task complete','Share your achievement','Fajerwerki; mobilekrótkawibracja. ShareonFacebook / Skip.','500monetraz/zadanie; zwykłanagrodanienaruszona.'),
      ('Decorations','Pięć kart atrakcji','Dolny poziomy pasek,5z20,rotacja24h.','Cena,premia,osoby,czas,przyzwyczajenie.'),
      ('Paths','Pięć poziomów','Dolny poziomy pasek poziomów1…5 i postęp0…100.','Blokadanastępnego poziomu dopełnych100poprzedniego.'),
      ('Nagłówek','Marka / Level / nawigacja','Rozłączne obszary bez przebijania starychobrazów.','Desktop/mobile.'),
      ('Settings','Animations and highlights','Jawnewłączenie nadpisuje systemowereduced-motion,zapiswyboru.','Efektywody,dymu,postaci,wiatraka,mew,zwierząt,chmur,budowy.'),
      ('New player','profile=new-player','Oddzielny zapis,normalnystart100zasobów/12osób.','Nieusuwa istniejącejtestowejwyspy; bezgrantu10k.'),
    ],{1:24,2:43,3:85,4:75})
    a=sheet(w,'Atrakcje','ATRAKCJE • 20 dekoracji i wydarzeń','5ofert co24h. Y5…100 zapisane w ofercie; cenaY×bieżący poziom ratusza. Każde skuteczne użycie: Y×przyzwyczajenie pp dla L osób na6h. Licznik osobny na rodzaj, reset po24h bez jego zakupu; przy0zakupzablokowany.',
       ['ID','Nazwa w grze','Znaczenie','Typ','Cena','Efekt','Czas','Przyzwyczajanie'],
       [(i,en,pl,t,'Y × L monet','Y × współczynnik pp dla L osób',6,'100% →75% →50% →25% →0%; osobny reset24h') for i,en,pl,t in ATTRACTIONS],{1:25,2:28,3:28,4:17,5:22,6:40,7:14,8:55})
    for row in range(6,26):a.cell(row,7).number_format='0" h"'
    ac_rows=[]
    for level in (1,2,3,5,10):
      for y in (5,25,50,100):
       for use in (1,2,3,4,5):
        r=6+len(ac_rows)
        ac_rows.append([level,y,use,f'=MAX(0,1-(C{r}-1)*0.25)',f'=A{r}*B{r}',f'=B{r}*D{r}',f'=ROUNDUP(F{r},0)',f'=A{r}',f'={refs["attraction_hours"]}',f'=IF(D{r}>0,"MOŻNA KUPIĆ","ZABLOKOWANE — bez opłaty")','Osobno na typ; każdy zakup osobne6h. Gdy0,cena wE jest nominalna,niepobierana.'])
    ac=sheet(w,'Atrakcje - kalkulator','ATRAKCJE • cena i przyzwyczajanie','Edytowalne A/B/C: ratuszL,ofertoweY i numerużycia. F zawiera dokładny składnikH; G jego prezentację w pełnychpp. CałeH mieszkańca nadalograniczone±100. Wiersze są niezależnymi przykładami, nie obietnicą cen jednej oferty.',
         ['Ratusz L','Y','Użycie typu','Współczynnik','Cena nominalna','Dokładna premia pp','Prezentacja pp','Maks.osób','Czas h','Dostępność','Uwagi'],ac_rows,{1:12,2:12,3:14,4:18,5:19,6:22,7:20,8:14,9:12,10:37,11:65})
    for row in range(6,ac.max_row+1):
      ac.cell(row,4).number_format='0%';ac.cell(row,6).number_format='0.00';ac.row_dimensions[row].height=38
      for col in (1,2,3):ac.cell(row,col).font=Font(name='Calibri',size=10,color=BLUE)
    roadrows=[]
    for tier in range(6):
      r=6+tier;roadrows.append([tier,100 if tier else 0,f'={refs["road_cost"]}',f'=IF(A{r}=0,0,{refs["road_sections"]}*C{r})',f'=A{r}*{refs["road_sections"]}*C{r}',f'=A{r}*{speedref}',f'=1+F{r}',f'=MAX(0,A{r}-2)*{loadref}',f'=1+H{r}',f'=300/G{r}','Premia dopiero po100/100 tego poziomu; prawo do zakupu wymaga100poprzedniego.'])
    road=sheet(w,'Sciezki','ŚCIEŻKI • 100 odcinków, poziomy0–5','Na starcie100odcinkówL0. Jedenklik=+1poziomjednegoodcinka za100kamienia. Żadenodcineknie może wejśćnaN,zanimkażdymaN−1. Ukończony poziom=minimum odcinków. Bonusytransportu nie są premiąszczęścia.',
       ['Ukończony poziom','Odcinki ukończone','Cena kroku','Koszt całego poziomu','Koszt od0','Bonus szybkości','Mnożnik szybkości','Bonus ładunku','Mnożnik ładunku','Kurs bazowy s','Warunek'],roadrows,{1:18,2:20,3:16,4:25,5:21,6:20,7:22,8:19,9:21,10:23,11:80})
    for r in range(6,12):
      for c in (6,8):road.cell(r,c).number_format='0%'
      road.cell(r,10).number_format='0.00" s"'
    rr=[]
    for level in (1,2,3,4):
      for tier in range(6):
        r=6+len(rr);rr.append([level,tier,f'=ROUNDDOWN(Parametry!$C$9*Parametry!$C$10^(A{r}-1),0)',f'=1+{loadref}*MAX(0,B{r}-2)',f'=ROUNDDOWN(C{r}*D{r},0)',f'=1+{speedref}*B{r}',f'=300/F{r}',f'=G{r}/0.5',100,f'=ROUNDUP(I{r}/E{r},0)',f'=J{r}*G{r}',f'=J{r}*Parametry!$C$43','1osoba,jedenrodzaj,gotowy100szt.,bez konkurującychkursów. Rejsstatku/droga do pracybezskrócenia.'])
    tr=sheet(w,'Sciezki - transport','ŚCIEŻKI • działanie na czas i ładunek','Ładunek: FLOOR(FLOOR(60×1,1^(własnyL−1))×mnożnikdróg). Czas:300/(1+0,1×poziomdróg); głód osobno×0,5szybkości. ParametrdrógB jest niezależny od poziomumagazynuA. Premiumcena kursu nadal10monet.',
       ['Własny L','Drogi0…5','Bazowy ładunek','Mnożnik dróg','Ładunek wynik','Mnożnik szybkości','Kurs nakarmiony s','Kurs głodny s','Towar do przewozu','Potrzebne kursy','Transport s','Premium monety','Założenia'],rr,{1:14,2:14,3:18,4:18,5:18,6:22,7:23,8:23,9:20,10:18,11:20,12:20,13:78})
    for r in range(6,tr.max_row+1):
      for c in (7,8,11):tr.cell(r,c).number_format='0.00'
    effects=[
      ('Jedzenie','+1 pp/os./h','Najedzeni mieszkańcy','Dopóki mają posiłek','Zmiana stanu karmienia','Tak,bez informacji podatkowych'),
      ('Praca','+1 pp/os./h','Osoby po dotarciu,faktyczniepracujące','Dopóki pracują bez blokady','Start/koniec i zmiana liczby osób','Tak'),
      ('Głód','Po12h −5pp/os./h; wydajność×0,5odrazu','Głodni mieszkańcy; zamrożoneH chronione','Do posiłku; licznikgłodu resetowany','Rozpoczęcie,prógkary,koniec','Tak'),
      ('Bezrobocie','Po12h −5pp/os./h','Niepracujący; osobny licznik','Do podjęciapracy','Rozpoczęcie,prógkary,koniec','Tak'),
      ('Przestój','−10pp odrazu; kolejne−10co24h,minimum−50składnika','Przypisani do zablokowanegoproducenta','Do odblokowania lubzwolnienia','Zmiana blokady lubstopniakary','Tak'),
      ('Przeludnienie','−1pp/os./h zaosobębezłóżka','Całapopulacja','Do zapewnienia miejsc; aktualna liczba braków','Początek,zmianaliczbybraków,koniec','Tak'),
      ('Pula domu','Ręczne dodatniepp; automatreszty23h58','Wybrani/najmniejszczęśliwi','Punkty trwałe; pulaodnawia24h','Przydział/automat/odnowienie','Tak,zdarzeniaw historii'),
      ('Kaplica:rozbudowa','+10×nowypoziompp','Populacja','10h','Początek,odświeżenie,koniec','Tak; ukrytymechaniczny×2podatków nadalobowiązuje'),
      ('Kaplica:błogosławieństwo','+50każdegozasobu,+20pp','Populacja','1h; kolejnepłatne4h odkońca','Początek,koniec','Tak'),
      ('Latarnia:rozbudowa','100%H; produkcjamateriałów×2','Populacja; noweosobymają0%przez24h','24h×osiągniętypoziom','Początek,odświeżenie,koniec','Tak'),
      ('Latarnia:uczestnikwizyty','100%H','Uczestnicy,popowrocie','24h odpowrotu; odwiedziny4h','Powrót,koniec','Tak'),
      ('Latarnia:premiagrupy','+10pp dla populacji','Całapopulacja','6h odpowrotukażdejgrupy','Każdy powrót,koniecgrupy','Tak; grupy niezależnesumująsię'),
      ('Atrakcja','Y×przyzwyczajeniepp','Lwybranychmieszkańców','6h osobnozakup','Początek,koniec','Tak'),
      ('Przyzwyczajenie','100/75/50/25/0%skuteczności','Konkretnyrodzajatrakcji','Reset po24hbezjegozakupu','Zmianastopnia,reset','Tak; toparametroferty,nieosobnakaraH'),
      ('Ścieżki:szybkość','+10%za ukończonypoziom','Tragarzemagazynu,targu,portu','Bezterminowo','Ukończenie100/100poziomu','Tak'),
      ('Ścieżki:ładunek','Od3:+10%,4:+20%,5:+30%','Tragarzemagazynu,targu,portu','Bezterminowo','Ukończeniecałegopoziomu','Tak'),
      ('Premiareakcjizadania','×3<15s,×2<5min,inaczej×1','Nagrodaprzyjętegozadania','Do sukcesu/wygaśnięcia','Utrwalenie,zakończenie','Tak'),
      ('Sharezadania','+500monetraz','Skarbiecgracza','Jednorazowezdarzenie','PoudanymkliknięciuShare','Powiadomienie; nieudajemypasywnegoefektu'),
      ('Wygaśnięciezadania','Utratanieodwracalnychwpłat','Zadanie','Wchwiliterminu','RaznawygasłeID','Powiadomienie'),
    ]
    sheet(w,'Efekty i notyfikacje','EFEKTY • powiadomienia i góra Town Hall','Każdy efekt maźródło,wartość,odbiorców i czas/warunek końca. Powiadomienie przy aktywacji,zmianie i zakończeniu, bez powtórzeńcoodświeżenie. Ratuszpokazuje aktywneefekty; jednorazowezdarzenia pozostają whistorii. SamychpodatkówniepokazujemywUI.',
       ['Źródło','Wartość / skutek','Kogo dotyczy','Czas / koniec','Powiadomienie','Widoczność'],effects,{1:30,2:57,3:46,4:50,5:50,6:65})
    # Build a transparent formula audit sheet, validated again from cached results.
    checks=[]
    def check(name,formula,expected,note=''):
      r=6+len(checks);checks.append([name,formula,expected,f'=B{r}=C{r}',note])
    check('ŹródłoV7 liczbaosób', '=Parametry!C6',12)
    check('Własnylimitpracy', '=Parametry!C8',12)
    check('Podateksilnikniezmieniony','=Parametry!C27',1,'UIbezpodatków,rachuneknadal1moneta.')
    check('PremiaShare500',f'={refs["task_share"]}',500)
    check('Katalog20', '=COUNTA(Atrakcje!A6:A25)',20)
    check('Widoczneoferty5',f'={refs["attraction_visible"]}',5)
    check('Rotacja24h',f'={refs["attraction_rotation"]}',24)
    check('Przerwa24h',f'={refs["attraction_reset"]}',24)
    check('Premia6h',f'={refs["attraction_hours"]}',6)
    for tier in range(6):
      rr0=6+tier
      check(f'Drogi{tier}:szybkość',f'=Sciezki!G{rr0}',1+.1*tier)
      check(f'Drogi{tier}:ładunek',f'=Sciezki!I{rr0}',1+.1*max(0,tier-2))
      check(f'Drogi{tier}:kosztod0',f'=Sciezki!E{rr0}',10000*tier)
    check('Wszystkie5poziomówkamień','=Sciezki!E11',50000)
    check('Pierwszepoziom+10%nie−10%czasu','=ROUND(Sciezki!J7,6)',round(300/1.1,6))
    for index,(level,y,use) in enumerate(( (l,y,u) for l in (1,2,3,5,10) for y in (5,25,50,100) for u in (1,2,3,4,5) )):
      r=6+index
      check(f'AtrakcjaL{level}Y{y}u{use}:cena',f"='Atrakcje - kalkulator'!E{r}",level*y)
      check(f'AtrakcjaL{level}Y{y}u{use}:premia',f"='Atrakcje - kalkulator'!F{r}",y*max(0,1-.25*(use-1)))
    check('Budowaratusza1→2bezdróg','=\'Czas budowy\'!Y6',80,'5rodzajów×ceil(500/60)=45kursów;6osób8fal×5min×2=80min.')
    check('Latarnianie2T','=\'Czas budowy\'!W72',5,'Dokładnywierszzostaniezweryfikowanyprzeznazwę/etapwskrypcie.')
    # Identify Lighthouse row rather than relying on an old hardcoded row number.
    lighthouse=next(r for r in range(6,w['Czas budowy'].max_row+1) if w['Czas budowy'].cell(r,1).value=='Latarnia')
    checks[-1][1]=f"='Czas budowy'!W{lighthouse}"
    cs=sheet(w,'Sprawdzenia v8','SPRAWDZENIA V8 • przeliczone przykłady','KolumnaD=TRUEoznacza zgodnośćwynikuformuły z niezależniepoliczoną wartościąoczekiwaną. Wszystkie formuły są przeliczane LibreOfficeprzed publikacją; walidator sprawdza brak błędów/cachedNone, komentarze,źródła ihashV7.',
         ['Sprawdzenie','Wynik formuły','Oczekiwane','Zgodne','Uwagi'],checks,{1:54,2:25,3:25,4:17,5:85})
    for r in range(6,cs.max_row+1):cs.row_dimensions[r].height=28
    # Readability pass applies only to newly authored sheets, never source cells.
    vocabulary = {
      'poziomów1':'poziomów 1', 'poziomu3':'poziomu 3', 'Lwybranychmieszkańców':'L wybranych mieszkańców',
      '100%skuteczności':'100% skuteczności', 'Przyzwyczajanie':'Przyzwyczajanie',
      'własnyL':'własny L', 'poziomdróg':'poziom dróg', 'mnożnikdróg':'mnożnik dróg',
      'Wszystkieformuły':'Wszystkie formuły', 'Wszystkie formuły':'Wszystkie formuły',
      'dopełnych':'do pełnych', 'podatkównadal':'podatków nadal', 'nieodwracalnychwpłat':'nieodwracalnych wpłat',
      'wybranychmieszkańców':'wybranych mieszkańców', 'niezamrożonychmieszkańców':'niezamrożonych mieszkańców',
      'Najedzeni mieszkańcy':'Najedzeni mieszkańcy', 'faktyczniepracujące':'faktycznie pracujące',
      'po dotarciu,faktycznie':'po dotarciu, faktycznie', 'wydajność×':'wydajność ×',
      'odrazu':'od razu', 'zamrożoneH':'zamrożone H', 'Do posiłku; licznikgłodu resetowany':'Do posiłku; licznik głodu resetowany',
      'Do podjęciapracy':'Do podjęcia pracy', 'Przypisani do zablokowanegoproducenta':'Przypisani do zablokowanego producenta',
      'Do odblokowania lubzwolnienia':'Do odblokowania lub zwolnienia', 'stopniakary':'stopnia kary',
      'zaosobębezłóżka':'za osobę bez łóżka', 'Całapopulacja':'Cała populacja',
      'Początek,zmianaliczbybraków,koniec':'Początek, zmiana liczby braków, koniec',
      'Pula domu':'Pula domu', 'dodatniepp':'dodatnie pp', 'automatreszty':'automat reszty',
      'Wybrani/najmniejszczęśliwi':'Wybrani / najmniej szczęśliwi', 'pulaodnawia':'pula odnawia',
      'zdarzeniaw historii':'zdarzenia w historii', 'nowypoziompp':'nowy poziom pp',
      'ukrytymechaniczny':'ukryty mechaniczny ', 'nadalobowiązuje':'nadal obowiązuje',
      'każdegozasobu':'każdego zasobu', 'popowrocie':'po powrocie',
      'noweosobymają':'nowe osoby mają ', 'przez24':'przez 24', 'osiągniętypoziom':'osiągnięty poziom',
      'końcagrupy':'końca grupy', 'koniecgrupy':'koniec grupy', 'Każdy powrót,koniec':'Każdy powrót, koniec',
      'grupy niezależnesumująsię':'grupy niezależne sumują się', 'Konkretnyrodzajatrakcji':'Konkretny rodzaj atrakcji',
      'Zmianastopnia':'Zmiana stopnia', 'toparametroferty,nieosobnakaraH':'to parametr oferty, nie osobna kara H',
      'za ukończonypoziom':'za ukończony poziom', 'Tragarzemagazynu,targu,portu':'Tragarze magazynu, targu, portu',
      'Ukończeniecałegopoziomu':'Ukończenie całego poziomu', 'Premiareakcjizadania':'Premia reakcji zadania',
      'Nagrodaprzyjętegozadania':'Nagroda przyjętego zadania', 'Utrwalenie,zakończenie':'Utrwalenie, zakończenie',
      'Sharezadania':'Share zadania', 'Skarbiecgracza':'Skarbiec gracza', 'Jednorazowezdarzenie':'Jednorazowe zdarzenie',
      'PoudanymkliknięciuShare':'Po udanym kliknięciu Share', 'nieudajemypasywnegoefektu':'nie jest pasywnym efektem',
      'Wygaśnięciezadania':'Wygaśnięcie zadania', 'Utratanieodwracalnych':'Utrata nieodwracalnych',
      'Wchwiliterminu':'W chwili terminu', 'RaznawygasłeID':'Raz na wygasłe ID',
      'bezpodatków':'bez podatków', 'rachuneknadal':'rachunek nadal',
      'Katalog20':'Katalog 20', 'Widoczneoferty5':'Widoczne oferty 5',
      'Wszystkie5poziomówkamień':'Wszystkie 5 poziomów — kamień',
      'Osobno na typ; każdy zakup osobne6h.':'Osobno na typ; każdy zakup ma osobne 6 h.',
      'cena wE jest nominalna,niepobierana':'cena w E jest nominalna, niepobierana',
      'Premia dopiero po100/100 tego poziomu; prawo do zakupu wymaga100poprzedniego.':'Premia dopiero po 100/100 tego poziomu; zakup wymaga 100 odcinków poprzedniego poziomu.',
      'nakarmiony':'nakarmiony', 'jednorazowezdarzenia':'jednorazowe zdarzenia',
      'Powiadomienie przy aktywacji,zmianie i zakończeniu':'Powiadomienie przy aktywacji, zmianie i zakończeniu',
      'bez powtórzeńcoodświeżenie':'bez powtórzeń co odświeżenie',
      'Ratuszpokazuje aktywneefekty':'Ratusz pokazuje aktywne efekty',
      'pozostają whistorii':'pozostają w historii', 'SamychpodatkówniepokazujemywUI':'Samych podatków nie pokazujemy w UI',
      'źródło,wartość,odbiorców':'źródło, wartość, odbiorców', 'czas/warunek końca':'czas / warunek końca',
      '1osoba,jedenrodzaj,gotowy100szt.,bez konkurującychkursów. Rejsstatku/droga do pracybezskrócenia.':'1 osoba, jeden rodzaj, gotowe 100 szt., bez konkurujących kursów. Rejs statku i droga do pracy bez skrócenia.',
      '1 osoba,jedenrodzaj,gotowy100szt.,bez konkurującychkursów. Rejsstatku/droga do pracybezskrócenia.':'1 osoba, jeden rodzaj, gotowe 100 szt., bez konkurujących kursów. Rejs statku i droga do pracy bez skrócenia.',
      'Do sukcesu/wygaśnięcia':'Do sukcesu / wygaśnięcia', 'czasowedodatki':'czasowe dodatki',
      '5produktów':'5 produktów', 'Construction, najwyżej3 wybrane':'Construction, najwyżej 3 wybrane',
      'numery1/2/3':'numery 1/2/3', 'powtarzane bezterminowo':'powtarzane bezterminowo',
      'Świadomewłączenie':'Świadome włączenie', 'systemowereduced-motion':'systemowe reduced-motion',
      'efektywody':'efekty wody', 'bezblokady':'bez blokady', 'poziomuratusza':'poziomu ratusza',
      'startu/wpłaty':'startu / wpłaty', 'Możnaobejrzeć':'Można obejrzeć',
      'Bez integracjiFacebook':'Bez integracji Facebook', 'Bez integracjiFB':'Bez integracji FB',
      'zwykłanagrodanienaruszona':'zwykła nagroda zachowana', 'bezstartu/wpłaty':'bez startu / wpłaty',
      'Stare wpłaty,terminy':'Stare wpłaty, terminy', 'mobilekrótkawibracja':'mobile: krótka wibracja',
      'Fajerwerki; mobilekrótkawibracja':'Fajerwerki; mobile: krótka wibracja',
      'ShareonFacebook':'Share on Facebook', 'jedno aktywne,zielono-złote':'jedno aktywne, zielono-złote',
      'Pozostałe2szare':'Pozostałe 2 szare', 'Źródło,wartość,odbiorcy,czas/warunek':'Źródło, wartość, odbiorcy, czas / warunek',
      'Dwa osobne wiersze obok budowanego obiektu.':'Dwa osobne wiersze obok budowanego obiektu.',
      'Rzeczywiste ustawienia dostaw; brak estymaty wyjaśniony.':'Rzeczywiste ustawienia dostaw; brak estymaty wyjaśniony.',
      'ładunek,ilość,źródło→cel,pozostałyczas':'ładunek, ilość, źródło → cel, pozostały czas',
      'Dokładnywierszzostaniezweryfikowanyprzeznazwę/etapwskrypcie.':'Wiersz latarni wskazywany automatycznie według nazwy.',
      'stanu ratusza':'stanu ratusza', 'ŹródłoV7 liczbaosób':'Źródło V7 — liczba osób',
      'Własnylimitpracy':'Własny limit pracy', 'Podateksilnikniezmieniony':'Podatek — silnik bez zmiany',
      'PremiaShare500':'Premia Share 500', 'bez jego zakupu':'bez jego zakupu',
      'przeliczane LibreOfficeprzed publikacją':'przeliczane LibreOffice przed publikacją',
      'KolumnaD=TRUEoznacza zgodnośćwynikuformuły':'Kolumna D = TRUE oznacza zgodność wyniku formuły',
      'z niezależniepoliczoną wartościąoczekiwaną':'z niezależnie policzoną wartością oczekiwaną',
      'brak błędów/cachedNone, komentarze,źródła ihashV7':'brak błędów i pustych wyników, komentarze, źródła i hash V7',
      'orazConstruction':'oraz Construction', 'postęp0…100':'postęp 0…100',
      'Blokadanastępnego poziomu':'Blokada następnego poziomu', 'dopełnych100poprzedniego':'do pełnych 100 poprzedniego',
      'Oddzielny zapis,normalnystart':'Oddzielny zapis, normalny start', 'Nieusuwa istniejącejtestowejwyspy':'Nie usuwa istniejącej testowej wyspy',
      'bezgrantu':'bez grantu ', 'niedodajemy':'nie dodajemy', 'stanlokalny':'stan lokalny',
      'OpisyCoins/Instant':'Opisy Coins / Instant', 'ramkaPremium':'ramka Premium', 'przyciskami':'przyciskami',
      'liczbąosób':'liczbą osób', 'Stanlokalny':'Stan lokalny',
    }
    new_names = ['Zmiany v8','Frontend v8','Atrakcje','Atrakcje - kalkulator','Sciezki','Sciezki - transport','Efekty i notyfikacje','Sprawdzenia v8']
    for name in new_names:
      for row in w[name]:
        for cell in row:
          if isinstance(cell.value,str) and cell.data_type != 'f':
            text = cell.value
            for old, new in vocabulary.items():
              text = text.replace(old,new)
            text = re.sub(r'(?<=\d)(?=[a-ząćęłńóśźż])',' ',text)
            text = re.sub(r'(?<=[a-ząćęłńóśźż])(?=\d)',' ',text)
            text = re.sub(r',(?=[^\s\d])',', ',text)
            cell.value = text

    # Explicit readable copy for review; identifiers/formulas/history stay unchanged.
    descriptions = {
      'Zmiany v8': SOURCE_NOTE + ' Nowsze usunięcie podatków oraz zapasy lokalne/centralne na kartach Resources zastępują wcześniejsze wersje tych samych elementów.',
      'Frontend v8': 'Krótkie angielskie podpisy w grze, polski opis zasad. Późniejsza instrukcja usuwa Taxes ze wszystkich widoków i zastępuje dolne free/max na Resources zapasami lokalny/centralny. Rachunek podatków pozostaje bez zmiany.',
      'Atrakcje': 'Pięć ofert zmienia się co 24 h. Y jest całkowite od 5 do 100 i zapisane w ofercie. Cena: Y × bieżący poziom ratusza L. Premia: Y × współczynnik przyzwyczajenia pp dla L osób przez 6 h. Każdy rodzaj ma osobny licznik; reset po 24 h bez jego zakupu. Przy zerowej premii zakup zablokowany.',
      'Atrakcje - kalkulator': 'Niebieskie pola A/B/C: poziom ratusza L, ofertowe Y i numer użycia. F zawiera dokładny składnik szczęścia; G prezentację w pełnych pp, zaokrągloną w górę. Osobiste H nadal ograniczamy do ±100%. Wiersze są niezależnymi przykładami, a nie cenami jednej oferty.',
      'Sciezki': 'Na starcie jest 100 odcinków na poziomie 0. Kliknięcie ulepsza jeden odcinek o jeden poziom za 100 kamienia. Poziom N jest dostępny dopiero, gdy wszystkie odcinki mają N−1. Ukończony poziom to minimum poziomów wszystkich odcinków; dopiero wtedy działa premia. Drogi nie dodają pasywnego szczęścia.',
      'Sciezki - transport': 'Ładunek = FLOOR(FLOOR(60 × 1,1^(własny L−1)) × mnożnik dróg). Czas kursu = 300 / (1 + 0,1 × poziom dróg). Głód osobno zmniejsza szybkość do połowy. B to poziom dróg, A to własny poziom miejsca pracy. Cena Premium nadal wynosi 10 monet za potrzebny kurs.',
      'Efekty i notyfikacje': 'Każdy efekt ma źródło, wartość, odbiorców oraz czas albo warunek zakończenia. Powiadomienia pojawiają się przy aktywacji, zmianie i końcu, bez powtórzeń przy odświeżaniu. Ratusz pokazuje aktywne efekty; jednorazowe zdarzenia pozostają w historii. Informacji o podatkach nie pokazujemy w UI.',
      'Sprawdzenia v8': 'TRUE w kolumnie D oznacza zgodność wyniku formuły z niezależnie obliczoną wartością oczekiwaną. Przed publikacją wszystkie formuły przeliczono w LibreOffice. Walidator sprawdza brak błędów i pustych wyników, zachowanie komentarzy i źródeł oraz niezmienioną sumę SHA-256 pliku V7.',
    }
    for name,text in descriptions.items():
      w[name]['A2']=text
      w[name].row_dimensions[2].height=76
    ui_copy = [
      ['Panel producenta','Level','Mniejsze pole po lewej.','Pole Taxes zostało usunięte.'],
      ['Panel producenta','Storage','In stock / Current cap. / Next level cap.','Zapas lokalny / 200 × L / 200 × (L+1).'],
      ['Miejsce pracy','Workers','Assigned / Free at home / Max workers','Trzy podpisy pod liczbami, np. 5 / 0 / 12.'],
      ['Miejsce pracy','+/−','Przydział z faktycznie dostępnych osób.','Droga trwa 5 min; brak wolnych osób blokuje plus.'],
      ['Każdy budynek','Status','Produkcja, blokada albo status odpowiedni do funkcji.','Przy produkcji: produkt, ilość oraz ETA złożenia w lokalnym magazynie.'],
      ['Wszystkie widoki','Podatki','Bez wyświetlania podatków i ich opisu.','Rachunek i premia kaplicy pozostają bez zmiany.'],
      ['House','Housing capacity','Populacja / limit ratusza / łóżka / przyrost lub spadek.','Opisy znajdują się pod i.'],
      ['House','House happiness','Pula i zegar widoczne; opis pod i.','Manage all workers otwierane wyłącznie tutaj.'],
      ['Town Hall','Construction and queue','Jedyny panel z licznikiem i zarządzaniem kolejką.','Inne budynki zachowują własny Upgrade.'],
      ['Town Hall','Active effects','Aktywne błogosławieństwa, premie i kary u góry karty.','Źródło, wartość, odbiorcy, czas lub warunek zakończenia; bez podatków.'],
      ['Producent','Building upgrade','Upgrade, dwa małe czasy i sąsiednia ramka Premium z dwoma przyciskami.','Buy building package / Deliver local stock now. Opisy Coins/Instant pod i.'],
      ['Etykieta budynku','Transport · ETA / Construction · ETA','Dwa osobne wiersze przy budowanym obiekcie.','Obecne ustawienia dostaw; brak wykonalnej estymaty ma wyjaśnienie.'],
      ['Warehouse','Warehouse workers','Osoba, ładunek, ilość, źródło → cel i pozostały czas.','Dojście i bezczynność są odróżnione od dostawy.'],
      ['Warehouse','Transport priority','Pięć produktów oraz Construction; najwyżej trzy wybrane, z numerami 1/2/3.','Powtarzane bezterminowo: 8/2, 5/3/2 albo 5/3/2.'],
      ['Resources','Pięć kart produktów','Wood / Planks / Stone / Grain / Food.','Klik otwiera producenta; Buildings przywraca budynki.'],
      ['Resources','Górny numerek','Faktycznie pracujący w zakładzie producenta.','Po dojściu, bez blokady. Głodny, ale produkujący pracownik nadal się liczy.'],
      ['Resources','Dolna para','Zapas lokalny / zapas tego produktu w centrali.','Np. 50 / 800; to już nie wolni / limit pracy.'],
      ['Tasks','Active / postęp','Jedno aktywne zadanie z postępem każdego produktu w obu widokach.','Pozostałe dwie karty są szare; nie można rozpocząć ani zasilać drugiej.'],
      ['Task complete','Share your achievement','Fajerwerki, krótka wibracja telefonu, Share on Facebook / Skip.','500 monet raz na zadanie. Pominięcie zachowuje zwykłą nagrodę.'],
      ['Decorations','Pięć kart atrakcji','Dolny poziomy pasek; 5 z katalogu 20, rotacja co 24 h.','Cena, premia, liczba osób, czas i przyzwyczajenie.'],
      ['Paths','Pięć poziomów','Dolny poziomy pasek poziomów 1–5, każdy z postępem 0–100.','Następny poziom zablokowany do ukończenia całej setki poprzedniego.'],
      ['Nagłówek','Marka / Level / nawigacja','Oddzielne obszary bez przebijania starych obrazów.','Komputer i telefon.'],
      ['Settings','Animations and highlights','Jawne włączenie nadpisuje systemowe reduced-motion i jest zapisywane.','Woda, dym, postacie, wiatrak, mewy, zwierzęta, chmury i budowa.'],
      ['New player','profile=new-player','Oddzielny zapis, normalny start po 100 zasobów i 12 osób.','Nie usuwa wyspy testowej; bez grantu 10 000 monet.'],
    ]
    for r,values in enumerate(ui_copy,6):
      for c,value in enumerate(values,1):w['Frontend v8'].cell(r,c,value)
    effect_copy = [
      ['Jedzenie','+1 pp/os./h','Najedzeni mieszkańcy','Dopóki mają posiłek','Zmiana stanu karmienia','Tak; bez informacji podatkowych'],
      ['Praca','+1 pp/os./h','Osoby po dotarciu, faktycznie pracujące','Dopóki pracują bez blokady','Start, koniec lub zmiana liczby osób','Tak'],
      ['Głód','Po 12 h: −5 pp/os./h; wydajność od razu ×0,5','Głodni mieszkańcy; zamrożenie H zachowane','Do posiłku; licznik głodu jest resetowany','Początek, próg kary, koniec','Tak'],
      ['Bezrobocie','Po 12 h: −5 pp/os./h','Niepracujący; oddzielny licznik','Do podjęcia pracy','Początek, próg kary, koniec','Tak'],
      ['Przestój','−10 pp od razu, potem −10 co 24 h; minimum −50 tego składnika','Przypisani do zablokowanego producenta','Do odblokowania lub zwolnienia','Zmiana blokady albo stopnia kary','Tak'],
      ['Przeludnienie','−1 pp/os./h za każdą osobę bez łóżka','Cała populacja','Do zapewnienia miejsc; zależne od braków','Początek, zmiana liczby braków, koniec','Tak'],
      ['Pula domu','Ręczny przydział dodatnich pp; automat reszty po 23 h 58 min','Wybrani / najmniej szczęśliwi','Punkty trwałe, pula odnawiana co 24 h','Przydział, automat, odnowienie','Zdarzenia w historii; bez fikcyjnego czasu wygaśnięcia'],
      ['Kaplica: rozbudowa','+10 × osiągnięty poziom pp','Populacja','10 h','Początek, odświeżenie, koniec','Tak; podatki mechanicznie ×2, lecz opis ukryty w UI'],
      ['Kaplica: błogosławieństwo','+50 każdego zasobu oraz +20 pp','Populacja','1 h; kolejne płatne 4 h po końcu','Początek, koniec','Tak'],
      ['Latarnia: rozbudowa','100% H i produkcja materiałów ×2','Populacja; nowe osoby nadal 0% przez 24 h','24 h × osiągnięty poziom','Początek, odświeżenie, koniec','Tak'],
      ['Latarnia: uczestnik wizyty','100% H','Uczestnicy po powrocie','24 h od powrotu; wizyta trwa 4 h','Powrót, koniec','Tak'],
      ['Latarnia: premia grupy','+10 pp dla populacji','Cała populacja','6 h od powrotu każdej grupy','Każdy powrót i koniec premii grupy','Tak; niezależne premie grup sumują się'],
      ['Atrakcja','Y × współczynnik przyzwyczajenia pp','L wybranych mieszkańców','6 h, osobno dla każdego zakupu','Początek, koniec','Tak'],
      ['Przyzwyczajenie','100 / 75 / 50 / 25 / 0% początkowej skuteczności','Konkretny rodzaj atrakcji','Reset po 24 h bez jej zakupu','Zmiana stopnia, reset','Parametr oferty, nie dodatkowa osobna kara H'],
      ['Ścieżki: szybkość','+10% za ukończony poziom','Tragarze magazynu, targu i portu','Bezterminowo','Ukończenie 100/100 poziomu','Tak'],
      ['Ścieżki: ładunek','Poziom 3: +10%, 4: +20%, 5: +30%','Tragarze magazynu, targu i portu','Bezterminowo','Ukończenie całego poziomu','Tak'],
      ['Premia reakcji zadania','×3 przed 15 s; ×2 przed 5 min; później ×1','Nagroda przyjętego zadania','Do sukcesu albo wygaśnięcia','Utrwalenie i zakończenie','Tak'],
      ['Share zadania','+500 monet, jeden raz','Skarbiec gracza','Jednorazowe zdarzenie','Po skutecznym kliknięciu Share','Powiadomienie; nie jest pasywnym efektem'],
      ['Wygaśnięcie zadania','Utrata nieodwracalnych wpłat','Zadanie','W chwili terminu','Raz na wygasłe ID','Powiadomienie'],
    ]
    for r,values in enumerate(effect_copy,6):
      for c,value in enumerate(values,1):w['Efekty i notyfikacje'].cell(r,c,value)
    for r in range(6,106):
      w['Atrakcje - kalkulator'].cell(r,11,'Każdy rodzaj ma osobny licznik; każdy zakup ma własne 6 h. Przy zerowej skuteczności nominalna cena z E nie jest pobierana.')
    for r in range(6,12):
      w['Sciezki'].cell(r,11,'Premia po ukończeniu 100/100 tego poziomu. Zakup następnego poziomu wymaga pełnej setki poprzedniego.')
    for r in range(6,30):
      w['Sciezki - transport'].cell(r,13,'Jedna osoba, jeden rodzaj towaru, 100 gotowych sztuk, bez konkurujących kursów. Rejs statku i droga do pracy nie są skracane.')
    w['Odpowiedzi i zmiany']['A1']='ODPOWIEDZI I ZMIANY • historia wcześniejszych wersji'
    w['Odpowiedzi i zmiany']['A2']='Ta zakładka zachowuje wcześniejsze odpowiedzi jako historię. Wersja 7 zebrała plik 6 i decyzje D01–D17. Najnowsze późniejsze zasady i korekty frontendu są w Zmiany v8; tam opisano także pierwszeństwo nowych ustaleń.'
    w['Decyzje i konflikty']['A2']='Dawne pytania zamknięto według pliku 6 i decyzji D01–D17. Kolumna G zachowuje oryginalne źródłowe wpisy użytkownika. Późniejsze uzgodnienia V8, w tym nowe priorytety, atrakcje, drogi i widoki, są w Zmiany v8. Tematy jawnie odłożone pozostają w Na pozniej.'
    w['Ustalenia - uwagi']['A2']='Zachowano 217 wpisów V7, poprawionych według nowszych zasad, oraz dodano pełny pakiet V8. Rejestr historyczny v7 zawiera poprzednie wartości. Źródła historyczne i komentarze pozostają dosłownie zachowane; najnowsza decyzja ma pierwszeństwo.'

    # Put new decision/UX sheets near the front without changing existing formula refs.
    front=['Start','Zmiany v8','Frontend v8','Decyzje v7','Budynki','Czlowiek','Atrakcje','Atrakcje - kalkulator','Sciezki','Sciezki - transport','Efekty i notyfikacje','Sprawdzenia v8']
    w._sheets=[w[n] for n in front]+[s for s in w if s.title not in front]
    w.active=0;w.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    w.properties.title='HARBORLIGHT8 — pełne ustalenia i kalkulatory'
    w.properties.subject='Wersja7 + zatwierdzone zmiany frontendu, atrakcji, ścieżek i efektów'
    w.properties.description='Wersja8. Zachowaneźródła,komentarze i wszystkie17decyzji. Nowereguły V8mająpierwszeństwo.'
    w.properties.modified=datetime(2026,10,23,12,0)
    OUTPUT.parent.mkdir(parents=True,exist_ok=True);w.save(OUTPUT)
    return refs


def recalculate():
    program=shutil.which('libreoffice') or shutil.which('soffice')
    if not program:raise RuntimeError('LibreOffice not found: cached results are required for the release workbook.')
    with tempfile.TemporaryDirectory(prefix='harborlight-v8-lo-') as tmp:
      d=Path(tmp);target=d/'out';target.mkdir()
      result=subprocess.run([program,f'-env:UserInstallation={(d/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(target),str(OUTPUT)],capture_output=True,text=True,timeout=120)
      final=target/OUTPUT.name
      if result.returncode or not final.exists():raise RuntimeError(f'LibreOffice failed: {result.stdout} {result.stderr}')
      OUTPUT.write_bytes(final.read_bytes())
      print(result.stdout.strip())


def validate(cached=True):
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==SOURCE_SHA
    old=load_workbook(SOURCE);new=load_workbook(OUTPUT);values=load_workbook(OUTPUT,data_only=True)
    original_comments={(s.title,c.coordinate):(c.comment.text,c.comment.author) for s in old for row in s for c in row if c.comment}
    for (name,coordinate),comment in original_comments.items():
      c=new[name][coordinate];assert c.comment and (c.comment.text,c.comment.author)==comment,(name,coordinate,'comment lost')
    histories=[s.title for s in old if s.title.startswith(('Twoje','Zrodlo','Rejestr')) or s.title=='Decyzje v7']
    for name in histories:
      for row in old[name]:
        for c in row:
          assert new[name][c.coordinate].value==c.value,('historical source changed',name,c.coordinate)
    # Source cells in current sheets are also retained, not only comments.
    for row in range(6,19):assert new['Budynki'].cell(row,11).value==old['Budynki'].cell(row,11).value
    for cell in ('E10','F16','F17'):assert new['Start'][cell].value==old['Start'][cell].value
    errors=[];formula_count=0
    for s in new:
      for row in s:
        for c in row:
          if c.data_type=='f':
            formula_count+=1;v=values[s.title][c.coordinate]
            if cached and (v.value is None or v.data_type=='e'):errors.append((s.title,c.coordinate,c.value,v.value))
    if cached:
      for r in range(6,values['Sprawdzenia v8'].max_row+1):
        if values['Sprawdzenia v8'].cell(r,4).value is not True:errors.append(('check failed',r,values['Sprawdzenia v8'].cell(r,1).value,values['Sprawdzenia v8'].cell(r,2).value,values['Sprawdzenia v8'].cell(r,3).value))
    assert not errors,json.dumps(errors[:20],ensure_ascii=False)
    with ZipFile(OUTPUT) as z:assert z.testzip() is None
    report={'file':str(OUTPUT.relative_to(ROOT)),'bytes':OUTPUT.stat().st_size,'sha256':hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),'source7_sha256':SOURCE_SHA,'sheets':len(new.sheetnames),'new_decisions':len(RULES),'rules':new['Ustalenia - uwagi'].max_row-5,'comments_preserved':len(original_comments),'formulas':formula_count,'checks':new['Sprawdzenia v8'].max_row-5,'all_checks_pass':cached,'original7_unchanged':True}
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report


def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--no-recalc',action='store_true');ap.add_argument('--validate-only',action='store_true');args=ap.parse_args()
    if not args.validate_only:
      build()
      if not args.no_recalc:recalculate()
    validate(cached=not args.no_recalc)

if __name__=='__main__':main()
