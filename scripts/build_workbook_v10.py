#!/usr/bin/env python3
"""Build Excel 10 from immutable Excel 9 and current game rules.

Only four sheets are visible. All 52 sheets of Excel 9 remain unchanged as
hidden history, including their formulas and original comments. Publication
requires LibreOffice recalculation followed by the cached-result validation.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

from build_workbook_v9 import cell_style, sheet


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs/arkusze/HARBORLIGHT9-zasady.xlsx'
OUTPUT = ROOT / 'docs/arkusze/HARBORLIGHT10-zasady.xlsx'
MANIFEST = ROOT / 'docs/decisions-v10.json'
SOURCE_HASHES = {
    7: 'bf49eef3b43ba81e0492c6eb187940c99177bbb8ae925ceca2b09ab0a2b9e945',
    8: '74ea3cc586d8d889ff34de8819694430503f45ed5fa2743873c9ad0c0b83cc3e',
    9: '3fc10a5d84df001e67e217dbc8c045cc5025b7777f91239c60b6694966864cc8',
}
VISIBLE = ['Jak działa gra v10', 'Budynki v10', 'Do decyzji v10', 'Kalkulator v10']
HIDDEN = ['Decyzje v10', 'Samouczek v10', 'Sprawdzenia v10']


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def check_sources():
    for version, digest in SOURCE_HASHES.items():
        path = ROOT / f'docs/arkusze/HARBORLIGHT{version}-zasady.xlsx'
        assert sha(path) == digest, f'Immutable Excel {version} changed'


def manifest():
    data = json.loads(MANIFEST.read_text())
    assert data['version'] == 10 and data['release'] == 'frontend-v16'
    assert not data.get('pendingDecisions'), ('Unresolved workbook decisions', data.get('pendingDecisions'))
    return data


def game_facts():
    """Read actual exported values; never parse lesson source with a regex."""
    code = """
      import { TUTORIAL_LESSONS } from './harborlight-v1-progression.js';
      import * as progression from './harborlight-v1-progression.js';
      import * as games from './harborlight-v1-minigames.js';
      import { translate } from './harborlight-v1-i18n.js?v=frontend-v16';
      import './harborlight-v1-v16-copy.js?v=frontend-v16';
      import { createState } from './harborlight-v1-state.js';
      import { productionRate, haulCapacity } from './harborlight-v1-logistics.js';
      const state = createState(Date.parse('2026-10-29T12:00:00Z'));
      const rules = games.MINI_GAME_RULES;
      console.log(JSON.stringify({
        lessons: TUTORIAL_LESSONS.map(item => ({ ...item,
          titlePL: translate(item.title, 'pl'), descriptionPL: translate(item.description, 'pl') })),
        progressionRules: Object.fromEntries(Object.entries(progression).filter(([key,value]) =>
          typeof value !== 'function' && key !== 'TUTORIAL_LESSONS')),
        miniGameRules: rules,
        tiltFirst: games.miniGameRoundRules('tilt', 1),
        tiltThird: games.miniGameRoundRules('tilt', 3),
        memoryFirst: games.miniGameRoundRules('triples', 1),
        baseProduction: productionRate(state, 'farm'), baseHaul: haulCapacity(state),
      }));
    """
    result = subprocess.run(['node', '--input-type=module', '-e', code], cwd=ROOT,
                            capture_output=True, text=True, timeout=20, check=True)
    return json.loads(result.stdout)


def markdown_table(path, heading, first_column):
    text = path.read_text()
    if heading:
        text = text.split(heading, 1)[1].split('\n## ', 1)[0]
    rows = []
    for line in text.splitlines():
        if not line.startswith('|'):
            continue
        cells = [cell.strip().replace('`', '') for cell in line.strip('|').split('|')]
        if cells[0] == first_column or re.fullmatch(r'[-: ]+', cells[0]):
            continue
        rows.append(cells)
    return rows


def mechanics(data):
    rows = markdown_table(ROOT / 'docs/ECONOMY-RULES.md', '## Jak działa gra', 'Mechanika')
    overrides = data.get('mechanicsOverrides', {})
    rows = [(title, overrides.get(title, {}).get('rule', rule),
             overrides.get(title, {}).get('source', source)) for title, rule, source in rows]
    existing = {row[0] for row in rows}
    rows += [(key, item['rule'], item['source']) for key, item in overrides.items() if key not in existing]
    return rows


def building_rows():
    rows = markdown_table(ROOT / 'docs/BUILDINGS.md', None, 'Budynek / EN')
    rows = [row for row in rows if len(row) == 7]
    assert len(rows) == 12, ('Expected all 12 buildings', len(rows))
    return [(name.split(' / ', 1)[0], name.split(' / ', 1)[1], *rest) for name, *rest in rows]


def tutorial_rows(facts):
    return [(index, item['id'], item['titlePL'], item['descriptionPL'], item.get('kind', ''),
             item.get('subject', ''), item.get('prerequisite') or '—')
            for index, item in enumerate(facts['lessons'], 1)]


def calculate_sheet(workbook, data, facts):
    k = sheet(workbook, VISIBLE[3], 'KALKULATOR V10 • jeden scenariusz',
              'Niebieskie pola można zmieniać. Wyniki są przykładami, a nie odczytem zapisu gry. '
              'Nowa baza produkcji i ładunku, koszt atrakcji i skrzynie odpowiadają bieżącym zasadom; '
              'krzywe trudności minigier pozostają jawną konfiguracją projektu.',
              ['Obszar', 'Parametr / wynik', 'Wartość', 'Jednostka', 'Znaczenie'], [], [25, 55, 27, 23, 110])

    def row(number, area, title, value, unit='', note='', editable=False, fmt=None):
        for column, item in enumerate([area, title, value, unit, note], 1):
            cell_style(k.cell(number, column, item), input=editable and column == 3)
        k.row_dimensions[number].height = 36
        if fmt:
            k.cell(number, 3).number_format = fmt

    def section(number, title):
        for column in range(1, 6):
            cell_style(k.cell(number, column, title if column == 1 else ''), header=True)
        k.row_dimensions[number].height = 29

    section(6, 'Produkcja oraz transport')
    row(7, 'Wejście', 'Własny poziom producenta / tragarzy', 1, 'poziom', editable=True)
    row(8, 'Wejście', 'Pełny poziom wszystkich dróg', 0, '0…5', editable=True)
    row(9, 'Wejście', 'Pracujący tragarze', 6, 'osób', editable=True)
    row(10, 'Wejście', 'Wyżywienie', 1, '1 lub 0,5', editable=True)
    row(11, 'Wejście', 'Każdy z pięciu materiałów do przewozu', 500, 'szt.', 'Każdy rodzaj wymaga osobnych kursów.', True)
    row(12, 'Wejście', 'Mnożnik samej fazy budowania', 1, '×', 'Zwykle 1; latarnia 4; kaplica 3 × (nowy poziom−1).', True)
    row(13, 'Wynik', 'Produkcja bazowa na osobę', '=IF(C7=1,20,2*(C7+10))', 'szt./godz.')
    row(14, 'Wejście', 'Mnożnik aktywnej gry', 1, '1 lub 2', 'Aktywna gra nie przyspiesza transportu.', True)
    row(15, 'Wejście', 'Mnożnik premii latarni', 1, '1 lub 2', editable=True)
    row(16, 'Wynik', 'Produkcja po premiach i głodzie', '=C13*C10*C14*C15', 'szt./os./godz.')
    row(17, 'Wynik', 'Ładunek na kurs', '=ROUNDDOWN(ROUNDDOWN(120*1.1^(C7-1),0)*(1+0.1*MAX(0,C8-2)),0)', 'szt.')
    row(18, 'Wynik', 'Kursy dla pięciu materiałów', '=5*ROUNDUP(C11/C17,0)', 'kursów')
    row(19, 'Wynik', 'Normalny transport bez innych zadań', '=IF(C9>0,ROUNDUP(C18/C9,0)*5/(1+0.1*C8)/C10,"brak obsady")', 'min', 'Estymata; rzeczywisty T obejmuje przerwy.')
    row(20, 'Wynik', 'Sama budowa z tego T', '=IF(ISNUMBER(C19),C19*C12,"brak estymaty")', 'min')
    row(21, 'Wynik', 'Transport i budowa', '=IF(ISNUMBER(C19),C19+C20,"brak estymaty")', 'min')
    row(22, 'Wynik', 'Natychmiastowy transport: cena', '=C18*10', 'monet', 'Dotyczy jednego wybranego etapu.')
    section(24, 'Prognozy FIFO i Premium')
    row(25, 'Wejście', 'Miniony rzeczywisty T', 10, 'min', editable=True)
    row(26, 'Wejście', 'Normalna estymata reszty transportu', 5, 'min', 'Przy zerowej obsadzie: jeden hipotetyczny tragarz.', True)
    row(27, 'Wynik', 'Budowanie po dowozie Premium', '=(C25+C26)*C12', 'min')
    row(28, 'Wejście', 'Fizycznie dostępny kamień', 200, 'szt.', editable=True)
    row(29, 'Wejście', 'Bazowy kamień każdego z dwóch planów', 300, 'szt.', editable=True)
    row(30, 'Wynik', 'Każdy plan po kolejce +10%', '=ROUNDUP(C29*1.1,0)', 'szt.')
    row(31, 'Wynik', 'Rezerwacja pierwszego planu', '=MIN(C28,C30)', 'szt.')
    row(32, 'Wynik', 'Brak pierwszego planu', '=MAX(0,C30-C31)', 'szt.')
    row(33, 'Wynik', 'Dopłata pierwszego planu', '=ROUNDUP(C32*0.5*0.25,0)', 'monet')
    row(34, 'Wynik', 'Rezerwacja drugiego planu', '=MIN(MAX(0,C28-C31),C30)', 'szt.')
    row(35, 'Wynik', 'Brak drugiego planu', '=MAX(0,C30-C34)', 'szt.')
    row(36, 'Wynik', 'Dopłata drugiego planu', '=ROUNDUP(C35*0.5*0.25,0)', 'monet')
    row(37, 'Wynik', 'Obie dopłaty', '=C33+C36', 'monet', 'Ten sam kamień jest zastawiony tylko raz.')
    section(39, 'Atrakcja • jedna na całej wyspie przez sześć godzin')
    row(40, 'Wejście', 'Cała populacja', 48, 'osób', editable=True)
    row(41, 'Wejście', 'Utrwalony w ofercie zasięg', 25, '% populacji', 'Losowanie 1–50 przy utworzeniu oferty.', True)
    row(42, 'Wejście', 'Uprawnieni mieszkańcy', 20, 'osób', 'Bez zamrożonych nowych osób i już szczęśliwych w 100%.', True)
    row(43, 'Wejście', 'Y z oferty', 50, 'pp', 'Od 5 do 100.', True)
    row(44, 'Wejście', 'Numer zakupu tego rodzaju', 2, '1…5', 'Reset przyzwyczajenia po 24 h przerwy.', True)
    row(45, 'Wynik', 'Faktyczni odbiorcy', '=IF(C40>0,MIN(C42,MAX(1,ROUNDDOWN(C40*C41/100,0))),0)', 'osób')
    row(46, 'Wynik', 'Skuteczność', '=MAX(0,1-(C44-1)*0.25)', 'udział', fmt='0%')
    row(47, 'Wynik', 'Premia na odbiorcę', '=C43*C46', 'pp')
    row(48, 'Wynik', 'Cena według sumy szczęścia', '=IF(AND(C45>0,C47>0),ROUNDUP(C45*C47,0),0)', 'monet')
    row(49, 'Wejście', 'Pozostały czas aktywnej atrakcji', 0, 's', 'Nowy zakup blokowany aż do końca 6 h efektu.', True)
    row(50, 'Wynik', 'Można uruchomić następną', '=AND(C49=0,C45>0,C47>0)', 'TAK/NIE', 'Dodatkowych 30 min po końcu efektu nie ma.')
    section(52, 'Nowe ekspedycje • nagroda maleje przed twardym końcem')
    row(53, 'Wejście', 'Poprzednia niewypłacona pula', 170, 'monet', editable=True)
    row(54, 'Wejście', 'Czas bieżącej rundy', 150, 's', editable=True)
    row(55, 'Wejście', 'Twardy limit rundy', 180, 's', 'Pierwszy labirynt 180 s; kolejne rundy zależą od krzywej.', True)
    row(56, 'Wynik', 'Początek spadku nagrody', '=C55*2/3', 's', 'Pierwszy labirynt: 120 s.')
    row(57, 'Wynik', 'Nagroda nowej wygranej', '=IF(C54<C55,MAX(0,50-ROUNDDOWN(MAX(0,C54-C56)/(C55/12),0)),0)', 'monet', '−1 za pełne 1/12 limitu po granicy; przy deadline przegrana.')
    row(58, 'Wynik', 'Pula po wygranej lub przegranej', '=IF(C54<C55,C53+C57,0)', 'monet', 'Po przegranej przepada cała niewypłacona seria.')
    row(59, 'Wejście', 'Numer rundy labiryntu', 3, 'runda', editable=True)
    row(60, 'PROJEKT', 'Limit kolejnego labiryntu', '=MAX(60,180/(1+0.1*(C59-1)))', 's', 'Krzywa konfiguracji trudności; nie mniej niż 60 s.', fmt='0.00')
    row(61, 'PROJEKT', 'Liczba kulek', '=MIN(6,1+ROUNDDOWN((C59-1)/2,0))', 'kulek')
    row(62, 'PROJEKT', 'Rozmiar wzoru labiryntu', '=MIN(8,3+ROUNDDOWN((C59-1)/3,0))', 'pól / bok')
    section(64, 'Samouczek • skrzynie i podwojenie za udostępnienie')
    row(65, 'Wynik', 'Liczba akcji samouczka', len(facts['lessons']), 'akcji', 'Bez czekania na kolejne dni.')
    row(66, 'Wejście', 'Wylosowana ilość skrzyni', 200, 'szt.', 'Losowanie 100–300; wynik zapisywany raz.', True)
    row(67, 'Wejście', 'Udostępnienie skrzyni', 0, '0 lub1', 'Lokalny prototyp; nie potwierdza posta naFacebooku.', True)
    row(68, 'Wynik', 'Łącznie pięć materiałów ze skrzyni', '=C66*5*(1+C67)', 'szt.', data['tutorial']['chestDescription'])
    row(69, 'Wynik', 'Wszystkie skrzynie — minimum', '=C65*5*100', 'szt.', 'Przed opcjonalnymi premiamiShare.')
    row(70, 'Wynik', 'Wszystkie skrzynie — maksimum', '=C65*5*300', 'szt.', 'Przed opcjonalnymi premiamiShare.')
    row(71, 'Wynik', 'Średnia suma bez Share', '=C65*5*200', 'szt.', 'Wartość oczekiwana, bez jednorazowej pomocy budowy.')
    row(72, 'Wynik', 'Średnia suma po wszystkich Share', '=C71*2', 'szt.', 'Każda skrzynia podwajana najwyżej raz.')
    section(74, 'Ścieżki, ludzie, punkty i centrala')
    row(75, 'Wejście', 'Odcinki jednego poziomu ścieżek', 100, 'odcinków', editable=True)
    row(76, 'Wynik', 'Kamień na pełny poziom', '=C75*100', 'kamienia')
    row(77, 'Wynik', 'Kamień na wszystkie pięć poziomów', '=100*100*5', 'kamienia')
    row(78, 'Wejście', 'Obecny poziom Ratusza', 2, 'poziom', editable=True)
    row(79, 'ZAŁOŻENIE', 'Cena osoby — każdego zasobu', '=500*C78', 'każdego z 6', 'Obowiązują wolne łóżko i limit populacji.')
    row(80, 'Wejście', 'Monety na migawce pełnej godziny', 1000, 'monet', editable=True)
    row(81, 'Wynik', 'Punkty', '=C78*C80', 'punktów')
    row(82, 'Wejście', 'Poziom Magazynu', 1, 'poziom', editable=True)
    row(83, 'Wynik', 'Pojemność każdego materiału', '=ROUNDDOWN(1000*1.5^(C82-1),0)', 'szt.')
    row(84, 'Wynik', 'Pojemność łączna', '=ROUNDDOWN(4000*1.5^(C82-1),0)', 'szt.')
    row(85, 'Wynik', 'Miejsce wyłącznie na żywność', '=50*C82', 'szt.')
    row(86, 'Wynik', 'Maksimum innych materiałów', '=C84-C85', 'szt.')
    section(88, 'Dopasowywanie trójek')
    row(89, 'Wejście', 'Numer etapu dopasowywania', 1, 'etap', 'Większa plansza i dłuższy limit; później przemieszczanie.', True)
    row(90, 'PROJEKT', 'Wszystkie kafelki', '=MIN(120,24+6*(C89-1))', 'kafelków')
    row(91, 'Wynik', 'Rodzaje symboli', '=C90/3', 'rodzajów', 'Każdy występuje dokładnie trzy razy.')
    row(92, 'PROJEKT', 'Limit rundy', '=MIN(900,240+30*(C89-1))', 's', 'Krzywa projektu; końcowe etapy najwyżej 15 min.')
    row(93, 'Wynik', 'Początek spadku nagrody', '=C92*2/3', 's')
    row(94, 'Wynik', 'Odstęp utraty jednej monety', '=C92/12', 's')
    row(95, 'PROJEKT', 'Ruch zakrytych pól co', '=IF(C89>=4,MAX(30,60-2*(C89-4)),0)', 's', '0 oznacza brak ruchu; odkryte i dopasowane nie są przesuwane.')
    row(96, 'Wynik', 'Trójki do usunięcia', '=C90/3', 'trójek')
    row(98, 'Uwaga', 'Lojalność — wysokości nagród', 'Nieustalone', '', 'Seria i limity są uzgodnione; arkusz nie przyznaje fikcyjnych kwot.')
    row(99, 'Uwaga', 'Prawdziwe konta, reklamy i publikacje', 'Niepodłączone', '', 'Lokalne przyciski nie potwierdzają zewnętrznych zdarzeń.')
    k.auto_filter.ref = None
    bounds = [(7,1,100),(8,0,5),(9,0,1200),(10,.5,1),(11,0,10000000),(12,1,300),
              (14,1,2),(15,1,2),(25,0,1000000),(26,0,1000000),(28,0,10000000),(29,0,10000000),
              (40,0,2400),(41,1,50),(42,0,2400),(43,5,100),(44,1,5),(49,0,21600),
              (53,0,10000000),(54,0,86400),(55,1,86400),(59,1,10000),(66,100,300),(67,0,1),
              (75,0,100),(78,1,100),(80,0,1000000000),(82,1,100),(89,1,10000)]
    for number, minimum, maximum in bounds:
        validation = DataValidation(type='decimal' if minimum == .5 else 'whole', operator='between',
                                    formula1=minimum, formula2=maximum)
        validation.showErrorMessage = True
        validation.error = 'Wpisz wartość w podanym zakresie.'
        k.add_data_validation(validation)
        validation.add(k.cell(number,3))
    return k


def checks_sheet(workbook, data, facts):
    checks = []
    def check(title, formula, expected):
        row_number = 6 + len(checks)
        checks.append((title, formula, expected, f'=B{row_number}=C{row_number}'))
    expected = {13:20,16:20,17:120,18:25,19:25,20:25,21:50,22:250,27:15,30:330,31:200,
                32:130,33:17,34:0,35:330,36:42,37:59,45:12,46:.75,47:37.5,48:450,50:True,
                56:120,57:48,58:218,60:150,61:2,62:3,65:25,68:1000,69:12500,70:37500,
                71:25000,72:50000,76:10000,77:50000,79:1000,81:2000,83:1000,84:4000,85:50,86:3950}
    for number, value in expected.items():
        check(f'Kalkulator C{number}', f"='{VISIBLE[3]}'!C{number}", value)
    for number, value in {90:24,91:8,92:240,93:160,94:20,95:0,96:8}.items():
        check(f'Dopasowywanie C{number}', f"='{VISIBLE[3]}'!C{number}", value)
    for kind, limit, grace, interval, cases in [
        ('Labirynt',180,120,15,[(120,50),(135,49),(150,48),(165,47),(180,0)]),
        ('Trójki',240,160,20,[(160,50),(180,49),(220,47),(240,0)])]:
        for elapsed, reward in cases:
            check(f'{kind}: nagroda w{elapsed}s',
                  f'=IF({elapsed}<{limit},MAX(0,50-ROUNDDOWN(MAX(0,{elapsed}-{grace})/{interval},0)),0)',reward)
    check('Minimalny limit labiryntu', '=MAX(60,180/(1+0.1*(100-1)))',60)
    check('Maksymalny limit trójek', '=MIN(900,240+30*(100-1))',900)
    check('Liczba akcji samouczka', "=COUNTA('Samouczek v10'!B6:B30)", len(facts['lessons']))
    check('Liczba opisanych budynków', "=COUNTA('Budynki v10'!A6:A17)",12)
    q = sheet(workbook, HIDDEN[2], 'SPRAWDZENIA V10 • kontrola arytmetyki',
              'Każdy wiersz ma dać TRUE po przeliczeniu LibreOffice. To kontrola arkusza, '
              'nie zastępstwo testów kodu gry ani potwierdzenia publicznej publikacji.',
              ['Scenariusz', 'Wynik', 'Oczekiwane', 'Zgodność'], checks, [60,28,28,20])
    q.sheet_state = 'hidden'



def assert_runtime(data, facts):
    assert len(facts['lessons']) == data['tutorial']['lessons'] == 25
    assert facts['baseProduction'] == 20 and facts['baseHaul'] == 120
    rules = facts['miniGameRules']
    assert rules['baseReward'] == data['expeditions']['baseReward'] == 50
    assert rules['fullRewardRatio'] == data['expeditions']['fullRewardFraction']
    assert rules['penaltyIntervals'] == 12
    assert rules['tilt']['baseSeconds'] == 180 and rules['tilt']['minimumSeconds'] == 60
    assert rules['tilt']['difficultyPerStage'] == .1
    assert rules['tilt']['maximumBalls'] == 6 and rules['tilt']['maximumCells'] == 8
    memory = data['expeditions']['memory']
    actual = rules['triples']
    for game_key, workbook_key in [('firstCards','firstCards'),('cardsPerStage','cardsPerStage'),
                                  ('maximumCards','maxCards'),('baseSeconds','firstSeconds'),
                                  ('extraSecondsPerStage','secondsPerStage'),('maximumSeconds','maxSeconds'),
                                  ('shuffleFromStage','shuffleFromStage'),('firstShuffleSeconds','firstShuffleInterval'),
                                  ('minimumShuffleSeconds','minShuffleInterval')]:
        assert actual[game_key] == memory[workbook_key], (game_key, actual[game_key], memory[workbook_key])
    progression = facts['progressionRules']
    assert progression['TUTORIAL_CHEST_MIN'] == data['tutorial']['chestMin'] == 100
    assert progression['TUTORIAL_CHEST_MAX'] == data['tutorial']['chestMax'] == 300
    assert progression['TUTORIAL_BUILDING_SUPPLIES'] == data['tutorial']['starterSupplies']
    for item in facts['lessons']:
        assert item['titlePL'] != item['title'], ('Untranslated tutorial title', item['id'])
        assert item['descriptionPL'] != item['description'], ('Untranslated tutorial description', item['id'])

def build():
    check_sources()
    data, facts = manifest(), game_facts()
    assert_runtime(data, facts)
    workbook = load_workbook(BASE)
    for old in workbook:
        old.sheet_state = 'hidden'
        old.sheet_properties.tabColor = 'AAB5B9'
    rules = mechanics(data)
    explanation = ('Czytaj cztery widoczne zakładki v10. Bieżące zasady obejmują frontend-v15 i v16; '
                   'wszystkie 52 zakładki Excela 9 zachowano jako ukryte archiwum. '
                   'Dawne kalkulatory i opisy nie zastępują tych nowszych zasad.')
    main = sheet(workbook, VISIBLE[0], 'HARBORLIGHT10 • jak działa gra', explanation,
                 ['Mechanika', 'Zasada w najwyżej trzech zdaniach', 'Źródło / status'], rules, [30,124,42])
    for number in range(6,main.max_row+1):
        main.row_dimensions[number].height = 76
    first = main.max_row + 3
    for column, value in enumerate(['Samouczek', '25 konkretnych akcji — bez czekania na kolejne dni', 'Skrzynia 100–300 / Share×2'],1):
        cell_style(main.cell(first,column,value),header=True)
    for number, item in enumerate(facts['lessons'],1):
        line = first+number
        for column,value in enumerate([f"{number}. {item['titlePL']}",item['descriptionPL'],
                                       data['tutorial']['chestDescription']],1):
            cell_style(main.cell(line,column,value))
        main.row_dimensions[line].height = 62
    main.auto_filter.ref = f'A5:C{first-3}'
    sheet(workbook, VISIBLE[1], 'BUDYNKI V10 • po co są i co kosztują',
          'L to obecny własny poziom, N to nowy poziom. Koszty kolejki rosną o 10% i są zaokrąglane w górę. '
          'T to rzeczywisty transport od pierwszego odbioru do ostatniej dostawy; podany czas budowania nie zawiera samej dostawy.',
          ['Budynek','EN','Po co jest','Ludzie: start / limit','Korzyść poziomu','Koszt ulepszenia','Sama budowa','Ważna zasada'],
          building_rows(),[24,23,70,28,55,58,28,86])
    sheet(workbook, VISIBLE[2], 'DO DECYZJI V10 • tylko to, co nadal odłożone',
          'Uzgodnione zasady i jawne ustawienia projektu nie wracają jako nowe pytania. '
          'Wysokości lojalności nie wymyślamy, a lokalne prototypy nie są prawdziwą integracją dostawcy.',
          ['ID','Temat','Co już wiadomo / co pozostaje','Status'],
          [(item['id'],item['title'],item['rule'],item['status']) for item in data['futureOnly']], [15,55,130,42])
    detail = [(item['id'],item['title'],item['rule'],item['status']) for item in data['changes']]
    detail += [('Baza9','Zachowane rozstrzygnięcia',
                '58 decyzji i 19 uwag frontendu pozostaje w niezmienionym ukrytym źródle Decyzje v9; '
                'nadpisania z tej zakładki i bieżące zasady v10 mają pierwszeństwo.','ARCHIWUM + NADPISANIA')]
    d = sheet(workbook,HIDDEN[0],'DECYZJE V10 • różnice względem Excela 9',
              'Manifest docs/decisions-v10.json odróżnia decyzje użytkownika, konfigurację projektu i przyszłe integracje.',
              ['ID','Temat','Bieżąca zasada','Status'],detail,[22,55,135,44])
    d.sheet_state='hidden'
    t=sheet(workbook,HIDDEN[1],'SAMOUCZEK V10 • rzeczywiste akcje',
            'Kolejność, identyfikatory i warunki są odczytywane z eksportu modułu gry. Skrzynie są rozliczane raz; Share daje tylko drugą taką samą nagrodę.',
            ['Krok','ID','Tytuł PL','Opis PL','Rodzaj','Temat','Poprzednik'],tutorial_rows(facts),[12,33,48,120,24,35,35])
    t.sheet_state='hidden'
    calculate_sheet(workbook,data,facts)
    checks_sheet(workbook,data,facts)
    workbook._sheets=[workbook[name] for name in VISIBLE]+[item for item in workbook if item.title not in VISIBLE]
    workbook.active=0
    workbook.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    workbook.properties.title='HARBORLIGHT10 — bieżące zasady i 25 akcji samouczka'
    workbook.properties.subject='Cztery czytelne zakładki; wszystkie źródła i komentarze zachowane'
    workbook.properties.description='Frontend-v16 z pełnym zakresem v15. Bieżące zasady v10 mają pierwszeństwo przed ukrytym archiwum 9.'
    workbook.properties.modified=datetime(2026,10,29,23,0)
    workbook.save(OUTPUT)


def recalc():
    program=shutil.which('libreoffice') or shutil.which('soffice')
    if not program:
        raise RuntimeError('LibreOffice is required to store cached formula results')
    with tempfile.TemporaryDirectory(prefix='harborlight-v10-lo-') as temporary:
        directory=Path(temporary)
        out=directory/'out'
        out.mkdir()
        result=subprocess.run([program,f'-env:UserInstallation={(directory/"profile").as_uri()}',
                               '--headless','--convert-to','xlsx','--outdir',str(out),str(OUTPUT)],
                              capture_output=True,text=True,timeout=120)
        converted=out/OUTPUT.name
        if result.returncode or not converted.exists():
            raise RuntimeError(result.stdout+'\n'+result.stderr)
        OUTPUT.write_bytes(converted.read_bytes())
        print(result.stdout.strip())


def validate(cached=True):
    check_sources()
    data,facts=manifest(),game_facts()
    assert_runtime(data, facts)
    source=load_workbook(BASE)
    result=load_workbook(OUTPUT)
    values=load_workbook(OUTPUT,data_only=True)
    assert [item.title for item in result if item.sheet_state=='visible']==VISIBLE
    comments=0
    for old in source:
        assert result[old.title].sheet_state=='hidden'
        for row in old:
            for cell in row:
                actual=result[old.title][cell.coordinate]
                assert actual.value==cell.value,(old.title,cell.coordinate,'archived value changed')
                if cell.comment:
                    comments+=1
                    assert actual.comment and actual.comment.text==cell.comment.text and actual.comment.author==cell.comment.author,(old.title,cell.coordinate,'comment changed')
    assert comments==48
    assert len(facts['lessons'])==25
    for number,item in enumerate(facts['lessons'],6):
        assert result[HIDDEN[1]].cell(number,2).value==item['id']
        assert result[HIDDEN[1]].cell(number,4).value==item['descriptionPL']
    for title,rule,_ in mechanics(data):
        assert len([part for part in re.split(r'[.!?](?:\s|$)',rule) if part.strip()])<=3,('more than three sentences',title)
    for number,item in enumerate(data['changes'],6):
        assert result[HIDDEN[0]].cell(number,1).value==item['id']
        assert result[HIDDEN[0]].cell(number,3).value==item['rule']
    errors=[]
    formulas=0
    for item in result:
        for row in item:
            for cell in row:
                cached_cell=values[item.title][cell.coordinate]
                if cached_cell.data_type=='e':
                    errors.append((item.title,cell.coordinate,cached_cell.value))
                if cell.data_type=='f':
                    formulas+=1
                    if cached and cached_cell.value is None:
                        errors.append((item.title,cell.coordinate,'missing cached result'))
    checks={}
    if cached:
        for name in ['Sprawdzenia v8','Sprawdzenia v9',HIDDEN[2]]:
            count=0
            for row in range(6,values[name].max_row+1):
                if values[name].cell(row,1).value is None:
                    continue
                count+=1
                if values[name].cell(row,4).value is not True:
                    errors.append((name,row,values[name].cell(row,1).value))
            checks[name]=count
    assert not errors,errors[:30]
    with ZipFile(OUTPUT) as archive:
        assert archive.testzip() is None
    report={'file':str(OUTPUT.relative_to(ROOT)),'bytes':OUTPUT.stat().st_size,'sha256':sha(OUTPUT),
            'visible_sheets':VISIBLE,'total_sheets':len(result.sheetnames),
            'archive9_sheets_unchanged':len(source.sheetnames),'comments_preserved':comments,
            'mechanics':len(mechanics(data)),'mechanics_at_most_three_sentences':True,
            'current_changes':len(data['changes']),'tutorial_actions':len(facts['lessons']),
            'formulas':formulas,'checks':checks,'cached_and_checks_pass':cached,
            'source7_8_9_unchanged':True}
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--validate-only',action='store_true')
    parser.add_argument('--no-recalc',action='store_true')
    options=parser.parse_args()
    if not options.validate_only:
        build()
        if not options.no_recalc:
            recalc()
    validate(not options.no_recalc)


if __name__=='__main__':
    main()
