#!/usr/bin/env python3
"""Build Excel 11 and its cached calculator without rewriting historical sheets.

Excel 10 is immutable. Its 59 worksheet XML parts, 48 source comments, shared
strings, styles and cached formula results are copied byte-for-byte. New v11
worksheets use those existing styles and inline strings. Only the new formula
caches are evaluated here; Excel may recalculate editable inputs on opening.
No LibreOffice recalculation or live GCP/account access is claimed.
"""
from __future__ import annotations

import argparse
import ast
from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as E
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs/arkusze/HARBORLIGHT10-zasady.xlsx'
OUTPUT = ROOT / 'docs/arkusze/HARBORLIGHT11-zasady.xlsx'
SCENARIOS = ROOT / 'docs/test-scenarios-v18.json'
REPORT = ROOT / 'docs/arkusze/HARBORLIGHT11-walidacja.json'
SOURCE_SHA256 = '4dd371e4fe779e51e1f060e69493635f25971a9c365ff9b55f961d457f31145a'
M = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
P = 'http://schemas.openxmlformats.org/package/2006/relationships'
C = 'http://schemas.openxmlformats.org/package/2006/content-types'
XML = 'http://www.w3.org/XML/1998/namespace'
VISIBLE = ['Jak działa gra v11', 'Budynki v11', 'Do decyzji v11', 'Kalkulator v11', 'Testy v11']
HIDDEN = ['Kontrole v11', 'Źródła v11']
PUBLISHED = {'commit': '6ec3c6d8625420fd96474a22a8b1d699672416cd', 'ci': '41111528800',
             'deployment': '41112673285', 'url': 'https://harborlight-dev.web.app'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def game_facts():
    source = r"""
      import * as s from './harborlight-v1-state.js';
      import * as l from './harborlight-v1-logistics.js';
      import * as p from './harborlight-v1-progression.js';
      import * as a from './harborlight-v1-activities.js';
      import * as g from './harborlight-v1-minigames.js';
      import * as e from './harborlight-v1-engagement.js';
      import * as w from './harborlight-v1-events.js';
      import * as c from './harborlight-v1-commerce.js';
      import { tutorialLessonText } from './harborlight-v1-v16-copy.js';
      const now=Date.parse('2026-10-30T12:00:00Z'), state=s.createState(now);
      const warehouses={};
      for(const level of [1,2,23,24,25,100]) {
        state.buildings.warehouse.level=level;
        warehouses[level]={...l.warehouseLimits(state),load:l.haulCapacity(state),reserve:l.foodReserve(state),workers:s.workerLimit(state,'warehouse')};
      }
      const production={};for(const level of [1,2,3,4,24]){state.buildings.farm.level=level;production[level]=l.productionRate(state,'farm');}
      const blessing=s.createState(now);blessing.resources.coins=10000;blessing.buildings.chapel.level=3;
      const blessingResult=s.act(blessing,{type:'paidBlessing'},now);
      if(!blessingResult.ok)throw Error(blessingResult.message);
      console.log(JSON.stringify({
        generatedFrom:'current local module exports',start:{population:s.createState(now).population,resources:s.createState(now).resources,happiness:s.createState(now).happiness},
        lessons:p.TUTORIAL_LESSONS.map(row=>({...row,titlePL:tutorialLessonText(row,'pl').title,descriptionPL:tutorialLessonText(row,'pl').description})),
        progression:{active:p.ACTIVE_PLAY_SECONDS,abandoned:p.ABANDONED_PLAY_SECONDS,loyaltyMin:p.LOYALTY_REWARD_MIN,loyaltyMax:p.LOYALTY_REWARD_MAX,chestMin:p.TUTORIAL_CHEST_MIN,chestMax:p.TUTORIAL_CHEST_MAX,supplies:p.TUTORIAL_BUILDING_SUPPLIES},
        miniRules:g.MINI_GAME_RULES,rounds:Object.fromEntries([1,2,3,4,17,23].map(n=>[n,{tilt:g.miniGameRoundRules('tilt',n),triples:g.miniGameRoundRules('triples',n)}])),
        activityRules:a.ACTIVITY_RULES,commerce:c.COMMERCE_RULES,packages:c.COIN_PACKAGES,worldEvents:w.WORLD_EVENT_RULES,
        session:{idle:e.SESSION_IDLE_SECONDS,warning:e.SESSION_WARNING_SECONDS},paidBlessingCooldown:blessing.paidBlessingCooldownRemaining,
        production,warehouses,demography:Object.fromEntries([-100,-75,-50,-25,-12.5,0,12.5,20,25,33.3,50,75,100].map(h=>[h,{rate:s.populationRate(h),interval:Number.isFinite(s.populationIntervalSeconds(h))?s.populationIntervalSeconds(h):null}]))
      }));
    """
    result = subprocess.run(['node', '--input-type=module', '-e', source], cwd=ROOT,
                            text=True, capture_output=True, check=True, timeout=30)
    facts = json.loads(result.stdout)
    assert len(facts['lessons']) == 25
    assert all(row['titlePL'] != row['title'] and row['descriptionPL'] != row['description'] for row in facts['lessons'])
    assert facts['warehouses']['24']['load'] == facts['warehouses']['100']['load']
    assert facts['warehouses']['24']['reserve'] == facts['warehouses']['100']['reserve']
    assert facts['start']['happiness'] == 20 and facts['start']['population'] == 12
    return facts


@dataclass
class Formula:
    text: str
    cached: object = None


@dataclass
class Sheet:
    name: str
    title: str
    note: str
    headings: list[str]
    widths: list[float]
    rows: dict[int, list] = field(default_factory=dict)
    inputs: set[str] = field(default_factory=set)
    section_rows: set[int] = field(default_factory=set)
    heights: dict[int, int] = field(default_factory=dict)
    validations: list[dict] = field(default_factory=list)
    hidden: bool = False

    def add(self, values, height=72):
        row = max([5, *self.rows]) + 1
        self.rows[row] = list(values)
        self.heights[row] = height
        return row

    def calc(self, row, area, title, value, unit='', note='', editable=False):
        self.rows[row] = [area, title, value, unit, note]
        self.heights[row] = 42
        if editable:
            self.inputs.add(f'C{row}')

    def section(self, row, title):
        self.rows[row] = [title]
        self.section_rows.add(row)
        self.heights[row] = 28


def clean(text):
    return re.sub(r'\[([^]]+)\]\([^)]*\)', r'\1', text).replace('`', '').strip()


def mechanics():
    text = (ROOT / 'docs/ECONOMY-RULES.md').read_text()
    rows = []
    for line in text.splitlines():
        if line.startswith('|'):
            cells = [clean(cell) for cell in line.strip('|').split('|')]
            if len(cells) == 2 and cells[0] != 'Mechanika' and not re.fullmatch('[-: ]+', cells[0]):
                rows.append((cells[0], cells[1], 'docs/ECONOMY-RULES.md · zatwierdzone v18'))
    rows += [
        ('Targ', 'Od Ratusza 2 wymieniasz pełne dziesiątki w jednej kategorii: 10 → 9. Partia jest zablokowana 1 h i ma 24 h ważności; wymiana działa bez obsady, ale przenoszą ją własni pracownicy. Nowy wynik ma nową ważność, a odbiór Premium nie kasuje zwykłego zlecenia.', 'logistics.js · v18'),
        ('Port', 'Od Ratusza 4 kupujesz i sprzedajesz paczki 100 sztuk. Baza kupno/sprzedaż to 50/40 monet dla drewna, kamienia i zboża oraz 150/140 dla desek i żywności; kolejne kupna w ruchomych 3 h podnoszą cenę o 10% bazy. Cena zamraża się w zamówieniu; sprzedaż płaci za faktycznie wysłane dobra.', 'logistics.js · zachowane reguły'),
        ('Ważność ładunku', 'Fizyczna partia ma 24 h od przybycia. Faktyczne podjęcie przed końcem chroni rozpoczęty kurs na obu etapach Portu, również statek → Port. Samo zlecenie nie chroni towaru, który nie został podjęty.', 'logistics.js · v18'),
        ('Zadania Latarni', 'Trzy oferty, jedno aktywne zadanie, brak porzucania i 60 s na potwierdzenie. Otwarcie utrwala mnożnik ×3 przed 15 s, ×2 przed 5 min, później ×1, ale oferta musi być ważna przy przyjęciu. Darowizny są nieodwracalne, bez tragarzy, z ochroną żywności.', 'activities.js · potwierdzenie i darowizny'),
        ('Terminy i nagrody zadań', 'Pierwsze zadanie trwa 6 h; po sukcesie następcy kończą się o północy Warszawy. Nagroda materiałowa trafia do od razu wydawalnego bonusu i nie ginie przez przewóz w toku. Share zadania daje osobny zatwierdzony bonus wyboru, nie dowód publikacji na FB.', 'activities.js, inventory.js'),
        ('Punkty Domu', 'O północy pula równa populacji zastępuje poprzednią resztę. O 23:58 automat rozdaje niezużyte punkty uprawnionym o najniższym H. Punkty nie zerują kar i ich liczników.', 'leisure.js · Europe/Warsaw'),
        ('Atrakcje', 'Pięć ofert losowanych o północy z katalogu 20 atrakcji: 5–100 pp na 6 h i utrwalony udział 1–50% populacji. Trafiają do uprawnionych o najniższym H; cena to pp po przyzwyczajeniu × rzeczywiści odbiorcy w górę. Tylko jedna nowa atrakcja może być aktywna naraz.', 'leisure.js · konfiguracja projektu'),
        ('Przyzwyczajanie', 'Osobno dla każdego rodzaju: 100%, 75%, 50%, 25%, 0%; po 24 h przerwy wraca pełna skuteczność. Przy 0% nie pobieramy opłaty. Oferta zachowuje wylosowany procent, a liczbę osób przelicza do aktualnej populacji.', 'leisure.js'),
        ('Ścieżki', '100 odcinków ma poziomy 0–5; jeden krok kosztuje 100 kamienia. Następny pełny poziom zaczyna się po ukończeniu wszystkich odcinków poprzedniego. Pełne poziomy dają +10…50% szybkości; poziomy 3/4/5 dodatkowo +10/20/30% ładunku.', 'leisure.js · parametry projektu'),
        ('Kaplica: darmowe i płatne', 'Jedno darmowe błogosławieństwo odnawia się o północy Warszawy bez kumulowania. Każde ręczne użycie daje 50 sześciu zasobów i +20 pp na 1 h; płatne kosztuje 100 × poziom Kaplicy monet. Płatne i reklamowe współdzielą 4 h blokady od końca efektu, darmowe jej nie przesuwa; ręczne efekty nie nakładają się.', 'state.js · dailyBlessingStatus'),
        ('Opieka Kaplicy', 'Od Kaplicy 3 wymagany odstęp to 3 + jej poziom dni. Po przekroczeniu losujemy dalsze 6–15 dni do zagrożenia. Każde rzeczywiste ręczne błogosławieństwo odnawia regularność.', 'events.js · chapelGraceSeconds'),
        ('Latarnia i wizyty', 'Ulepszenie daje wszystkim 100% H i produkcję materiałów ×2 na 24 h × nowy poziom; przybyły korzysta do wspólnego końca. Wizyta trwa 4 h bez dojścia i wycofania, po czym uczestnik ma 100% i blokadę wizyt na 24 h, a grupa +10 pp dla wyspy na 6 h. Wymuszenie nie zwraca utraconych ludzi ani materiałów.', 'residents.js, effects.js'),
        ('Dobre wraki', 'Zdarzenie losujemy co 36–60 h świata. Podczas globalnej ochrony Latarni najwyżej jedno dobre czeka do końca, a złe terminy są anulowane i losowane ponownie z zachowaniem opieki Kaplicy. Nie dopisujemy fikcyjnego błogosławieństwa.', 'events.js · konfiguracja projektu'),
        ('Wyniki dobrego wraku', 'Skarb 50–150 × TH monet, zwrot podatku 20–60 × TH, materiały 50–150 pięciu rodzajów, +1 rozbitek albo nawóz Farmy +25% na 6 h. Brak miejsca na rozbitka daje 500 × TH monet. Wagi są jawną konfiguracją projektu.', 'events.js · WORLD_EVENT_RULES'),
        ('Złe wraki', 'Jeden główny skutek i kara H na 6 h: ludność do 20% z minimum 6, wolne monety/drewno do 20% albo transport wolniejszy do 50%. Najwyżej jedno zaległe złe zdarzenie na okres bez rzeczywistej gry. Rezerwacje i niesiony towar pozostają chronione.', 'events.js · konfiguracja projektu'),
        ('Ekspedycje', 'Gracz wybiera kulki albo trójki; nagrody rund to 50, 100, 200, 400… przed potrąceniami za czas, a pula 50, 150, 350, 750…. Kontynuacja ryzykuje całą niewypłaconą serię; wypłata ją kończy. Twardy termin kończy próbę, bez automatycznej opłaty za spóźnienie.', 'minigames.js, activities.js · v18'),
        ('Kulki', 'Pierwsza runda trwa 180 s, pełna nagroda przez pierwsze 120 s, następnie −1 moneta co 15 s. Późniejsze limity i liczby kulek wynikają z tabeli eksportowanej w Źródła v11. Sterowanie myszą/klawiaturą albo ruchem telefonu/dotykiem.', 'minigames.js · MINI_GAME_RULES'),
        ('Trójki', 'Pierwsza plansza ma 24 pola i 240 s; każdy etap dodaje 6 pól i 30 s, do 120 pól i 900 s. Od etapu 4 przemieszczają się zakryte niedopasowane obrazki. Plansza mieści się bez suwaka na pulpicie i telefonach; komplet trzech identycznych znika.', 'minigames.js · rzeczywiste eksporty'),
        ('Dobrowolny dodatkowy czas', 'Przed końcem można kupić +30 s za 10 monet, maksymalnie dwa razy na rundę, lub obejrzeć reklamę za +60 s raz. Reklama pauzuje wyłącznie zegar rundy, świat biegnie dalej, a przerwa ma bezpieczny limit. To jawne parametry wdrożeniowe, nie historycznie wybrany cennik.', 'activities.js · ACTIVITY_RULES'),
        ('Obecność w minigrze', 'Po 5 min bez aktywności pytanie o obecność, 60 s na odpowiedź. Zegar rundy nie zatrzymuje się przez samo pytanie. Brak odpowiedzi kończy serię i traci niewypłaconą pulę.', 'activities.js · niezależnie od sesji 30 min'),
        ('Harmonogram propozycji', 'Pierwsza reklama raz/konto w 4–8 aktywnej minucie, po końcu lub pominięciu pierwszy wybór minigra/reklama po 5–15 min. Po ukończeniu tej aktywności kolejny po 43–62 min, następne po 60–100 min od ukończenia poprzedniej. Harmonogram jest zapisany i nie przerywa trwającej gry.', 'engagement.js · zatwierdzone v18'),
        ('Premia reklamy', 'Reklama materiałowa daje po 10–30 pięciu materiałów bez monet. Bonus za czas i błogosławieństwo nie dodaje osobnej premii materiałowej. Rzeczywista nagroda wymaga zaufanego potwierdzenia dostawcy; symulacja jest oznaczona.', 'commerce.js'),
        ('Standard bez reklam', 'Potwierdzone zakupy brutto co najmniej 1999 gr w ruchomych 30 dniach dają brak standardowych reklam, niezależnie od pakietu budowlanego. Bonus bez filmu zachowuje limity, a reklamy za czas ekspedycji nadal są oglądane. Początek harmonogramu 4–8 min kotwiczy się raz automatycznie, odbiór premii go nie zeruje.', 'commerce.js, engagement.js'),
        ('Paynow i monety', 'Pakiety 1/5/10/20/50/100/500 zł, po 10 monet za 1 zł, są obecnymi parametrami testowymi. DEV i PRE-PROD mają oznaczoną symulację bez pobierania pieniędzy; PROD nie przyznaje zakupów z klientowego potwierdzenia. Rzeczywisty operator wymaga osobnego podłączenia.', 'commerce.js · COIN_PACKAGES'),
        ('Kody promocyjne', 'Właściciel ustala kod, nagrodę i globalny limit. Serwer atomowo przyznaje pierwszym N kontom, jedno wykorzystanie na konto; powtórzenie nie wypłaca drugi raz. Kody wymagające niepodłączonej weryfikacji polubienia FB są odrzucane.', 'backend/account-service.mjs · wymaga aktywnego API'),
        ('Lojalność', 'Raz na dzień Warszawy 50–200 każdego z sześciu zasobów, również monet; losowanie i odbiór są zapisane. Przerwa resetuje serię, a program kończy dziesiąta odebrana nagroda albo Ratusz 4. Nie wracamy do dawnego otwartego pytania o kwoty.', 'progression.js · LOYALTY_REWARD_MIN/MAX'),
        ('Skrzynie samouczka', '25 rzeczywistych działań z prowadzeniem strzałką; skrzynia daje 100–300 pięciu materiałów, bez monet, a wybór Share raz tyle samo. Przepełnienie blokuje kolejne skrzynie i podwojenia, ale nie pomoc, wydawanie ani zarządzanie. Popup skrzyni czeka do końca minigry; stare odebrane nagrody są zachowane.', 'progression.js, tutorial-ui.js'),
        ('Społeczność i polecenia', 'Udostępnienie osiągnięcia inicjuje i zatwierdza gracz; otwarcie formularza nie potwierdza zewnętrznej publikacji. Punkty integracji premii społecznościowych i poleceń nie oznaczają aktywnej weryfikacji Facebooka. Automatyzacja strony gry i osobiste powiadomienia FB są odłożone.', 'progression.js, auth-providers.js · granice integracji'),
        ('Powiadomienia i opisy', 'Ostatnie 100 powiadomień; kategorie filtrują również dzwonek. Ręcznie usuwa się tylko przeczytane wiadomości. Port/Targ odświeżają etapy, skrócenie budowy podaje rzeczywiście oszczędzony czas, H100 ma jawną ikonę.', 'play.js · odbiór v18'),
        ('Zapis konta', 'Postęp jest wyłącznie w chmurze i wczytuje się przed otwarciem wyspy. Nie zapisujemy wysp, odzyskiwania ani kolejki w localStorage/sessionStorage; sesja logowania i język są dozwolone. Awaria lub konflikt blokuje zatwierdzanie i wymaga jawnego pobrania prawdy serwera.', 'dev-cloud.js, dev-sync.js'),
        ('Retencja', 'Po trzech miesiącach kalendarzowych bez rzeczywistej gry wyspa może wygasnąć, a jej generacja blokuje stare karty. Konto, historie płatności/kodów/poleceń i nagrody raz na konto pozostają; PRE-PROD admin1–3 są wyłączone. Automatyczne zwroty wycofano, finalne postanowienia płatnej sprzedaży wymagają przeglądu regulaminu.', 'backend/account-service.mjs · PROD po aktywacji'),
        ('Środowiska', 'DEV to ta przestrzeń, PRE-PROD to harborlight-dev, PROD to harborlight-prod. Aktualnie potwierdzony PRE-PROD v18: 6ec3c6d i wdrożenie 41112673285; nowe zmiany tej partii są roboczym DEV do własnego CI/promocji. Zachowujemy UID i historię admin1, bez resetu przy publikacji.', 'AGENTS.md · rzeczywisty stan publikacji'),
        ('Nowe wylogowanie', 'W Ustawieniach ma być bezpośredni widoczny przycisk wylogowania. Korzysta z istniejącej bezpiecznej obsługi zapisu i sesji, bez resetowania wyspy. To nowy zakres tej partii DEV, wymagający własnego odbioru i wdrożenia.', 'Zlecenie 30.10.2026 · nowy DEV'),
        ('Regulamin i prywatność', 'Robocze dokumenty dla gry wymagają uzupełnienia danych operatora Loremipsum i przeglądu właściciela. Rejestracja ma wymagać odrębnej akceptacji regulaminu i potwierdzenia zapoznania się z polityką prywatności, z wersją dokumentów i serwerowym czasem. Potwierdzenie informacji nie jest wymuszoną zgodą RODO na konieczne przetwarzanie.', 'Nowy DEV · nie potwierdza jeszcze publikacji'),
    ]
    return rows


def buildings():
    rows = []
    for line in (ROOT / 'docs/BUILDINGS.md').read_text().splitlines():
        if not line.startswith('|'):
            continue
        cells = [clean(cell) for cell in line.strip('|').split('|')]
        if len(cells) != 7 or cells[0] == 'Budynek / EN' or re.fullmatch('[-: ]+', cells[0]):
            continue
        local, english = cells[0].split(' / ', 1)
        rows.append([local, english, *cells[1:]])
    assert len(rows) == 12
    return rows


def calculator(facts):
    s = Sheet(VISIBLE[3], 'KALKULATOR V11 · zmienne i wynik',
              'Niebieskie pola są wejściami. Wyniki są przykładami, nie odczytem konta ani rachunku GCP. '
              'Nowe formuły mają zapisane, niezależnie zweryfikowane wyniki początkowe; Excel przelicza je po zmianie wejść. '
              'Archiwalne formuły i ich zapisane wyniki z Excela 10 pozostają nienaruszone.',
              ['Obszar','Parametr / wynik','Wartość','Jednostka','Znaczenie / ograniczenie'], [26,58,28,24,108])
    def row(n,area,title,value,unit='',note='',editable=False):s.calc(n,area,title,value,unit,note,editable)
    def f(text):return Formula(text)
    s.section(6,'Produkcja i przewóz — model Farmy / Magazynu')
    for args in [(7,'Poziom producenta',1,'1…100'),(8,'Poziom Magazynu',1,'1…100'),(9,'Pełny poziom wszystkich dróg',0,'0…5'),(10,'Pracujący tragarze',6,'osób'),(11,'Mnożnik wyżywienia',1,'1 lub 0,5'),(12,'Mnożnik rzeczywistej aktywności',2,'2 / 1 / 0,5'),(13,'Premia produkcji Latarni',1,'1 albo 2'),(14,'Premia Farmy z nawozu',1,'1 albo 1,25'),(15,'Mnożnik drogi uszkodzonej we wraku',1,'0,5…1')]:row(args[0],'Wejście',*args[1:],'',True)
    row(17,'Produkcja','Bazowa wydajność jednego pracownika',f('=IF(C7=1,20,2*(C7+10))'),'szt./h','Bez mnożników; rzeczywisty eksport productionRate.')
    row(18,'Produkcja','Wydajność pracownika Farmy po mnożnikach',f('=C17*C11*C12*C13*C14'),'szt./h','Nawóz dotyczy tylko Farmy; przestój i brak pracowników mogą obniżyć realizację.')
    row(19,'Produkcja','Limit obsady tego poziomu producenta',f('=12*C7'),'osób')
    row(20,'Produkcja','Bufor lokalny producenta',f('=200*C7'),'szt.')
    row(22,'Magazyn','Efektywny poziom dla ładunku i rezerwy',f('=MIN(C8,24)'),'poziom','Po L24 nadal rośnie limit pracowników, lecz nie ładunek/pojemność/rezerwa.')
    row(23,'Magazyn','Bazowy ładunek przed drogami',f('=ROUNDDOWN(120*1.1^(C22-1),0)'),'szt./kurs')
    row(24,'Magazyn','Ładunek po premii dróg',f('=ROUNDDOWN(C23*(1+0.1*MAX(0,C9-2)),0)'),'szt./kurs')
    row(25,'Magazyn','Mnożnik szybkości dróg',f('=1+0.1*C9'),'×')
    row(26,'Magazyn','Pełny kurs z mnożnikami',f('=300/(C11*C12*C15*C25)'),'s','Dojście pracownika 5 min jest odrębne i nie jest przyspieszane.')
    row(27,'Magazyn','Teoretyczna przepustowość załogi',f('=C24*MIN(C10,12*C8)*3600/C26'),'szt./h','Obsada ograniczona do 12 × poziom Magazynu; zależy też od towaru, odbiorców i miejsca.')
    row(28,'Magazyn','Pojemność jednego zwykłego materiału',f('=ROUNDDOWN(MIN(10000000,1000*1.5^(C8-1)),0)'),'szt.')
    row(29,'Magazyn','Pojemność łączna zwykłych materiałów',f('=ROUNDDOWN(MIN(40000000,4000*1.5^(C8-1)),0)'),'szt.')
    row(30,'Magazyn','Rezerwa żywności',f('=50*C22'),'szt.')
    row(31,'Magazyn','Limit pracowników Magazynu',f('=12*C8'),'osób','Ten wynik nadal rośnie po L24.')
    s.section(33,'Zapas bonusowy — zatwierdzony przykład')
    row(34,'Wejście','Limit zwykłego materiału',300,'szt.','Przykład jednego materiału bez innych rezerwacji.',True)
    row(35,'Wejście','Zwykły zapas',200,'szt.','',True)
    row(36,'Wejście','Od razu wydawalna nagroda',200,'szt.','Nie zużywa limitu zwykłych dostaw.',True)
    row(37,'Wynik','Łącznie wydawalne przed dalszą dostawą',f('=C35+C36'),'szt.')
    row(38,'Wynik','Miejsce na następną zwykłą dostawę',f('=MAX(0,C34-C35)'),'szt.')
    row(39,'Wynik','Łącznie po wykorzystaniu tego miejsca',f('=C37+C38'),'szt.','Powinno być 500, mimo zwykłego limitu 300.')
    s.section(41,'Demografia — odrębne tempo przybycia i odejść')
    row(42,'Wejście','Dokładne średnie szczęście',100,'−100…100%','Miejsca populacji i ochrona grup nadal ograniczają faktyczne zmiany.',True)
    row(43,'Wynik','Zmiana populacji na godzinę',f('=IF(C42=0,0,SIGN(C42)*IF(C42>0,2,1)*IF(ABS(C42)<=50,ABS(C42)/300,1/6+(ABS(C42)-50)/60))'),'osób/h','Dodatni wzór ×2; ujemny bez przyspieszenia.')
    row(44,'Wynik','Czas do zmiany o jedną osobę',f('=IF(C43=0,"brak zmiany",3600/ABS(C43))'),'s','Przy +100%: 1800 s; przy −100%: 3600 s.')
    s.section(46,'Minigry — nagroda nominalna, przed potrąceniami czasu')
    row(47,'Wejście','Numer rundy',4,'1…40','Umiarkowany zakres kalkulatora; gra blokuje niebezpieczne przepełnienie liczb.',True)
    rules=facts['miniRules'];tilt=rules['tilt'];mem=rules['triples']
    row(48,'Nagroda','Nominalna nagroda tej rundy',f(f'={rules["baseReward"]}*2^(C47-1)'),'monet')
    row(49,'Nagroda','Nominalna pula po wszystkich wygranych',f(f'={rules["baseReward"]}*(2^C47-1)'),'monet','50, 150, 350, 750…; przegrana zabiera niewypłaconą serię.')
    row(50,'Kulki','Twardy limit czasu',f(f'=MAX({tilt["minimumSeconds"]},{tilt["baseSeconds"]}/(1+{tilt["difficultyPerStage"]}*(C47-1)))'),'s')
    row(51,'Trójki','Liczba pól',f(f'=MIN({mem["maximumCards"]},{mem["firstCards"]}+{mem["cardsPerStage"]}*(C47-1))'),'pól')
    row(52,'Trójki','Twardy limit czasu',f(f'=MIN({mem["maximumSeconds"]},{mem["baseSeconds"]}+{mem["extraSecondsPerStage"]}*(C47-1))'),'s')
    row(53,'Kulki','Koniec pełnej nagrody',f(f'=C50*{rules["fullRewardRatio"]}'),'s')
    row(54,'Kulki','Przedział potrącenia 1 monety',f(f'=C50/{rules["penaltyIntervals"]}'),'s')
    ar=facts['activityRules']
    row(55,'Czas','Dobrowolnie kupione sekundy',ar['extraTimeSeconds'],'s',f"{ar['extraTimeCoins']} monet; maks. {ar['maximumExtraTimePurchases']} zakupy na rundę.")
    row(56,'Czas','Czas za reklamę',ar['adTimeSeconds'],'s',f"Maks. {ar['maximumAdTimeRewards']} /rundę; także u gracza bez reklam standardowych.")
    s.section(58,'Testowy pakiet monet i standard bez reklam')
    row(59,'Wejście','Pakiet brutto',20,'PLN','Wyłącznie demonstracja obecnego cennika; nie pobiera pieniędzy.',True)
    row(60,'Wynik','Kwota w groszach',f('=ROUND(C59*100,0)'),'gr')
    row(61,'Wynik','Monety w pakiecie',f(f'=C59*{ar["coinsPerPLN"]}'),'monet')
    row(62,'Wynik','Sam pakiet spełnia próg bez reklam',f(f'=C60>={facts["commerce"]["noAdsThresholdGrosz"]}'),'TRUE/FALSE','W grze liczy się suma potwierdzonych, nieoddanych płatności z ruchomych 30 dni.')
    row(63,'Reguła','Okno sumowania płatności',facts['commerce']['noAdsWindowDays'],'dni')
    s.section(65,'Kaplica i termin zagrożenia po zaniedbaniu')
    row(66,'Wejście','Poziom Kaplicy',3,'1…100','Monitorowanie zaniedbania zaczyna się dopiero od L3.',True)
    row(67,'Wynik','Dozwolony odstęp od błogosławieństwa',f('=IF(C66<3,"nie dotyczy",3+C66)'),'dni','Poniżej L3 zaniedbanie nie jest monitorowane.')
    row(68,'Wynik','Najwcześniejsze zagrożenie od błogosławieństwa',f('=IF(C66<3,"nie dotyczy",C67+6)'),'dni','Przykład bez ochrony Latarni i innych przesunięć.')
    row(69,'Wynik','Najpóźniejsze zagrożenie od błogosławieństwa',f('=IF(C66<3,"nie dotyczy",C67+15)'),'dni','Samo wyświetlenie nie losuje nowego terminu.')
    row(70,'Reguła','Blokada od startu płatnego/reklamowego efektu',facts['paidBlessingCooldown']/3600,'h','1 h działania + 4 h po jego końcu; darmowe nie zmienia terminu.')
    s.section(72,'Skrzynie oraz lojalność — różne programy nagród')
    row(73,'Wejście','Liczba skrzyń podstawowych',25,'0…25','',True)
    pr=facts['progression']
    row(74,'Skrzynie','Minimum pięciu materiałów łącznie',f(f'=C73*5*{pr["chestMin"]}'),'szt.','Bez monet i bez pomocy do Ratusza.')
    row(75,'Skrzynie','Maksimum pięciu materiałów łącznie',f(f'=C73*5*{pr["chestMax"]}'),'szt.')
    row(76,'Skrzynie','Wartość oczekiwana łącznie',f(f'=C73*5*({pr["chestMin"]}+{pr["chestMax"]})/2'),'szt.')
    row(77,'Skrzynie','Średnia po identycznym Share każdej',f('=C76*2'),'szt.','Każdą skrzynię wolno podwoić raz; nie potwierdza to publikacji na FB.')
    row(78,'Lojalność','Minimum jednej dziennej nagrody',pr['loyaltyMin'],'każdego z 6','Wliczamy monety.')
    row(79,'Lojalność','Maksimum jednej dziennej nagrody',pr['loyaltyMax'],'każdego z 6','Losowanie zapisane; limit 10 odbiorów albo TH4.')
    s.section(81,'Nieobecność — przykład wyjścia zaraz po rzeczywistej akcji')
    row(82,'Wejście','Długość nieobecności',60,'h','48 h jest granicą od ostatniego realnego działania.',True)
    row(83,'Wejście','Pozostała premia aktywna przy wyjściu',5,'min, 0…5','Sam login/heartbeat nie odnawia tej premii.',True)
    row(84,'Wynik','Czas ×2',f('=MIN(C82,C83/60)'),'h')
    row(85,'Wynik','Czas ×1',f('=MAX(0,MIN(C82,48)-C84)'),'h')
    row(86,'Wynik','Czas ×0,5',f('=MAX(0,C82-48)'),'h')
    row(87,'Wynik','Równoważnik godzin bazowej gospodarki',f('=2*C84+C85+0.5*C86'),'h bazy','W przykładzie 60 h tylko ostatnie 12 h ma ×0,5.')
    row(88,'Wynik','Rzeczywisty czas jedzenia, efektów, kalendarza',f('=C82'),'h','Demografia ma własne zasady; nie mnożymy całego świata przez bonus gospodarki.')
    s.section(90,'GCP — plan miesięcznego alertu 100 USD, bez twardego odcięcia')
    row(91,'Wejście','Żądany próg miesięcznego budżetu',100,'USD','Łącznie harborlight-dev i harborlight-prod; przygotowanie planu nie oznacza zastosowania w GCP.',True)
    row(92,'Wejście','Przykładowy kurs do orientacyjnego przeliczenia',4,'PLN / USD','Wyłącznie przykład arytmetyczny, nie bieżący ani gwarantowany kurs faktury.',True)
    row(93,'Wynik','Orientacyjne przeliczenie progu na PLN',f('=C91*C92'),'PLN','Nie jest to odczyt kosztów. Budżet GCP wysyła alerty; sam nie odcina usług ani faktury.')
    for cell,min_v,max_v in [('C7',1,100),('C8',1,100),('C9',0,5),('C10',0,1200),('C34',0,10**9),('C35',0,10**9),('C36',0,10**9),('C47',1,40),('C66',1,100),('C73',0,25)]:
        s.validations.append({'type':'whole','sqref':cell,'formula1':str(min_v),'formula2':str(max_v),'operator':'between'})
    for cell,choices in [('C11','1,0.5'),('C12','0.5,1,2'),('C13','1,2'),('C14','1,1.25'),('C59',','.join(str(pack['grossGrosz']/100).removesuffix('.0') for pack in facts['packages']))]:
        s.validations.append({'type':'list','sqref':cell,'formula1':'"'+choices+'"'})
    for cell,a,b in [('C15',0.5,1),('C42',-100,100),('C82',0,2400),('C83',0,5),('C91',0.01,10**6),('C92',0.001,1000)]:s.validations.append({'type':'decimal','sqref':cell,'formula1':str(a),'formula2':str(b),'operator':'between'})
    return s


def prepare_sheets(facts, scenarios):
    main=Sheet(VISIBLE[0],'HARBORLIGHT11 · aktualne zasady i 25 działań',
      'Stan: zatwierdzone v18, opublikowane PRE-PROD 6ec3c6d / 41112673285 oraz osobno oznaczony nowy zakres DEV z 30.10.2026. '
      'Pięć widocznych zakładek v11 ma pierwszeństwo; 59 zakładek Excela 10 to niezmienione ukryte archiwum. '
      'Testy ręczne mają status Nie wykonano, a koszt GCP nie jest zgadywany.',
      ['Mechanika','Bieżąca reguła','Źródło / stan'],[32,126,54])
    for row in mechanics():main.add(row,80)
    main.add(['SAMOUCZEK · 25 aktualnych rzeczywistych działań z eksportu modułu gry; stabilne ID i zachowane odebrane nagrody'],30);main.section_rows.add(max(main.rows))
    for item in facts['lessons']:main.add([f"{item['index']}. {item['titlePL']}",item['descriptionPL'],f"{item['id']} · {item['kind']} · {item.get('route','')}"],65)
    b=Sheet(VISIBLE[1],'BUDYNKI V11 · korzyść, koszt i transport',
      'L = obecny własny poziom, N = nowy. Kolejka dodaje 10% każdego składnika w górę. T jest rzeczywistą dostawą od pierwszego podjęcia do ostatniej dostawy, nie prognozą ani samą fazą budowania. Koszt gotówkowy jest osobnym kanałem od TH1; zdobyte monety od TH4.',
      ['Budynek','EN','Cel','Obsada: start / limit','Korzyść poziomu','Koszt zwykłego ulepszenia','Sama budowa','Ważna reguła'],[25,24,74,30,65,65,30,90])
    for row in buildings():b.add(row,100)
    d=Sheet(VISIBLE[2],'DO DECYZJI V11 · stan wydania i rzeczy do podłączenia',
      'Nie otwieramy ponownie zatwierdzonych decyzji. Parametry wdrożeniowe są jawne; kod i mock nie oznaczają konfiguracji dostawcy. '
      'Nowy zakres tej partii jest DEV do osobnych testów i promocji; regulaminy są robocze, operator Loremipsum.',
      ['ID','Obszar','Co jest ustalone / co pozostaje','Stan'],[16,48,134,42])
    items=[
      ('REL-01','Opublikowane PRE-PROD','v18 / 6ec3c6d; pełne CI 41111528800 i udane wdrożenie 41112673285. Gra https://harborlight-dev.web.app; oba artefakty wydania zachowane.','POTWIERDZONE 30.10.2026'),
      ('REL-02','PROD i domena','harborlight-prod / harborlight.pl nie mają potwierdzonego wydania. Wymagana promocja tych samych artefaktów oraz wcześniejszy odbiór API/kont na PRE-PROD.','NIE OPUBLIKOWANO'),
      ('NEW-01','Wylogowanie w Ustawieniach','Żądany bezpośredni przycisk korzystający z bezpiecznego przepływu zapisu/sesji. Nie resetować admin1 ani innych stałych kont.','NOWY DEV · DO ODBIORU'),
      ('NEW-02','Regulamin i prywatność','Robocze dokumenty z operatorem Loremipsum; właściciel uzupełni dane i przejrzy treść. Wersjonowana akceptacja regulaminu i potwierdzenie zapoznania się z polityką przy rejestracji, także nowy OAuth.','SZKIC · NOWY DEV'),
      ('EXT-01','Google i Facebook OAuth','Kod korzysta z Firebase/SDK i wymaga aktywacji dostawców/domen przez właściciela. PRE-PROD zachowuje tylko trzy UID, łączenie dostawcy nie tworzy nowej wyspy.','KONFIGURACJA WŁAŚCICIELSKA'),
      ('EXT-02','Paynow','Operator wybrany. Obecnie wyraźna symulacja DEV/PRE-PROD, bez pobierania pieniędzy; rzeczywista sprzedaż po założeniu JDG i integracji zaufanych potwierdzeń.','REALNE PŁATNOŚCI WYŁĄCZONE'),
      ('EXT-03','Reklamy','Zasady i symulacja są w kodzie; rzeczywista sieć reklamowa i potwierdzenie obejrzenia wymagają konfiguracji. Nie przyznawać produkcyjnej nagrody za klientowe verified:true.','DOSTAWCA NIEPOTWIERDZONY'),
      ('EXT-04','API kont, kody i retencja','Kod serwera oraz routing istnieją, ale rzeczywistą aktywację flagi i IAM trzeba potwierdzić. Sam sukces PRE-PROD nie oznacza włączenia wszystkich funkcji i dostawców.','AKTYWACJA DO WERYFIKACJI'),
      ('SOC-01','Publikacje strony i powiadomienia FB','Użytkownik inicjuje i zatwierdza udostępnienie osiągnięcia. Automatyczne posty strony gry i osobisty push FB są odłożone.','ODŁOŻONE'),
      ('SOC-02','Zweryfikowane publikacje i polecenia','Punkty integracji nie są potwierdzeniem rzeczywistego posta ani polecenia. Nie wypłacać premii tylko za otwarcie formularza.','NIEPODŁĄCZONE'),
      ('BAL-01','Dodatkowy czas','Aktualny jawny dobór wdrożeniowy: +30 s za 10 monet, maks.2/rundę; reklama +60 s raz. Wartości mogą być później strojone bez twierdzenia, że były historycznie wybraną ceną.','USTAWIENIE PROJEKTU'),
      ('BAL-02','Wagi wraków','Skarb30%, zwrot20%, materiały25%, rozbitek10%, nawóz15%; negatywne plaga10%, pozostałe po30%. Są to jawne wartości konfiguracji projektu.','USTAWIENIE PROJEKTU'),
      ('GCP-01','Plan alertu 100 USD miesięcznie','Żądany budżet alertowy: 100 USD miesięcznie łącznie dla harborlight-dev i harborlight-prod. Przygotowanie konfiguracji w DEV nie potwierdza zastosowania w GCP; standardowy budżet sam nie odcina usług ani nie limituje faktury. Bieżące naliczenia pozostają prywatnym odczytem właściciela.','NOWY DEV · DO ZASTOSOWANIA I ODBIORU'),
      ('AUDIO-01','Muzyka','Wybór/licencja utworu i zakup pozostają do ustalenia. Arkusz nie oznacza zakupu licencji ani wdrożonego podkładu.','NOWE PYTANIE · POZA MECHANIKĄ'),
      ('ROAD-01','Dalszy rozwój','Osobna modlitwa, błogosławieństwa za social/wyzwania, serwerowa autorytatywna gospodarka, AR i planszówka pozostają odłożone.','ODŁOŻONE'),
    ]
    for row in items:d.add(row,86)
    t=Sheet(VISIBLE[4],'TESTY V11 · scenariusze ręczne i odbiór',
      'To scenariusze do wykonania, nie raport PASS. admin1 sprawdzamy wyłącznie odczytowo; stałych admin2/3 także nie resetujemy. '
      'Mutacje, awarie, retencję i fikstury prowadzimy w lokalnym DEV lub na jawnie przeznaczonym koncie testowym. '
      'Nie zmieniać zegara ani danych żywej wyspy; warunki i środowisko każdego testu są obowiązkowe.',
      ['ID','Obszar','Scenariusz','Priorytet','Środowisko','Konto / fikstura','Warunki początkowe','Kroki','Oczekiwany wynik','Tryb','Stan implementacji','Wynik ręczny','Źródła'],
      [18,27,56,14,38,55,78,105,105,34,35,25,65])
    for item in scenarios:
        steps='\n'.join(f'{i}. {step}' for i,step in enumerate(item['steps'],1))
        expected='\n'.join(f'• {step}' for step in item['expected'])
        t.add([item['id'],item['module'],item['title'],item['priority'],item['environment'],item['account'],item['preconditions'],steps,expected,item['mode'],item['implementation_status'],item['status'],'\n'.join(item['source'])],150)
    t.validations.append({'type':'list','sqref':f'L6:L{max(t.rows)}','formula1':'"Nie wykonano,PASS,FAIL,Zablokowany"'})
    return [main,b,d,calculator(facts),t]


def cell_at(sheets, sheet_name, coordinate):
    sheet=next(s for s in sheets if s.name==sheet_name)
    match=re.fullmatch(r'([A-Z]+)(\d+)',coordinate)
    column=0
    for char in match[1]:column=column*26+ord(char)-64
    values=sheet.rows.get(int(match[2]),[])
    return values[column-1] if len(values)>=column else 0


def evaluate(formula, sheets, current_sheet, trail=()):
    xrefs=[]
    def cross(match):xrefs.append((match[1],match[2]));return f'XREF_{len(xrefs)-1}'
    text=re.sub(r"'([^']+)'!([A-Z]{1,3}[1-9]\d*)",cross,formula.lstrip('='))
    text=re.sub(r'(?<![A-Za-z0-9_])([A-Z]{1,3}[1-9]\d*)\b',r'cell("\1")',text)
    for index,(sheet,coordinate) in enumerate(xrefs):text=text.replace(f'XREF_{index}',f'cross("{sheet}","{coordinate}")')
    text=text.replace('^','**').replace('<>','!=')
    text=re.sub(r'(?<![<>=!])=(?!=)','==',text)
    tree=ast.parse(text,mode='eval')
    def lookup(sheet,coordinate):
        if (sheet,coordinate) in trail:raise ValueError('Circular workbook reference')
        value=cell_at(sheets,sheet,coordinate)
        return evaluate(value.text,sheets,sheet,(*trail,(sheet,coordinate))) if isinstance(value,Formula) else value
    def visit(node):
        if isinstance(node,ast.Expression):return visit(node.body)
        if isinstance(node,ast.Constant):return node.value
        if isinstance(node,ast.UnaryOp):return -visit(node.operand) if isinstance(node.op,ast.USub) else +visit(node.operand)
        if isinstance(node,ast.BinOp):
            a,b=visit(node.left),visit(node.right)
            return {ast.Add:lambda:a+b,ast.Sub:lambda:a-b,ast.Mult:lambda:a*b,ast.Div:lambda:a/b,ast.Pow:lambda:a**b}[type(node.op)]()
        if isinstance(node,ast.Compare):
            assert len(node.ops)==len(node.comparators)==1
            a,b=visit(node.left),visit(node.comparators[0])
            return {ast.Eq:lambda:a==b,ast.NotEq:lambda:a!=b,ast.GtE:lambda:a>=b,ast.LtE:lambda:a<=b,ast.Gt:lambda:a>b,ast.Lt:lambda:a<b}[type(node.ops[0])]()
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name):
            name=node.func.id
            if name=='IF':return visit(node.args[1] if visit(node.args[0]) else node.args[2])
            args=[visit(x) for x in node.args]
            if name=='cell':return lookup(current_sheet,args[0])
            if name=='cross':return lookup(args[0],args[1])
            funcs={'MIN':min,'MAX':max,'ABS':abs,'SUM':lambda *x:sum(x),'SIGN':lambda x:(x>0)-(x<0),
                   'ROUND':lambda x,n=0:math.floor(x*10**n+0.5)/10**n,
                   'ROUNDDOWN':lambda x,n=0:math.trunc(x*10**n)/10**n}
            if name in funcs:return funcs[name](*args)
        raise ValueError(f'Unsupported formula syntax: {ast.dump(node)}')
    return visit(tree)


def add_checks(sheets,facts,scenario_count):
    c=Sheet(HIDDEN[0],'KONTROLE V11 · formuły wobec modułów gry',
            'TRUE oznacza zgodność zapisanej formuły z oczekiwanym wynikiem na wejściach początkowych. Nie jest wynikiem ręcznego scenariusza ani testem wdrożenia.',
            ['Kontrola','Formuła / odczyt','Oczekiwane','Zgodność','Źródło'],[58,32,32,18,85],hidden=True)
    k='Kalkulator v11'
    expected={17:facts['production']['1'],18:40,19:12,20:200,22:1,23:facts['warehouses']['1']['load'],24:120,25:1,26:150,27:17280,
              28:facts['warehouses']['1']['perResource'],29:facts['warehouses']['1']['total'],30:50,31:12,37:400,38:100,39:500,
              43:facts['demography']['100']['rate'],44:facts['demography']['100']['interval'],48:facts['rounds']['4']['tilt']['baseReward'],49:750,
              50:facts['rounds']['4']['tilt']['timeLimit'],51:facts['rounds']['4']['triples']['cardCount'],52:facts['rounds']['4']['triples']['timeLimit'],
              53:facts['rounds']['4']['tilt']['fullRewardUntil'],54:facts['rounds']['4']['tilt']['penaltyEvery'],60:2000,61:200,62:True,
              67:6,68:12,69:21,74:12500,75:37500,76:25000,77:50000,84:1/12,85:48-1/12,86:12,87:54+1/12,88:60,93:400}
    for n,value in expected.items():
        row=c.add([f'Kalkulator C{n}',Formula(f"='{k}'!C{n}"),value,Formula(f'=ABS(B{len(c.rows)+6}-C{len(c.rows)+6})<0.000001'),'Niezależny eksport Node lub zatwierdzony przykład'],42)
        # Boolean comparison must not pass through numeric subtraction in Excel.
        if isinstance(value,bool):c.rows[row][3]=Formula(f'=B{row}=C{row}')
    c.add(['25 działań samouczka',len(facts['lessons']),25,True,'Eksport TUTORIAL_LESSONS'],42)
    c.add(['L24/L100: identyczny ładunek',facts['warehouses']['24']['load'],facts['warehouses']['100']['load'],True,'haulCapacity; ograniczenie tylko Magazynu'],42)
    c.add(['L24/L100: identyczna rezerwa',facts['warehouses']['24']['reserve'],facts['warehouses']['100']['reserve'],True,'foodReserve'],42)
    c.add(['L24/L100: rosnąca obsada',facts['warehouses']['100']['workers'],facts['warehouses']['24']['workers'],True,'workerLimit:1200 >288'],42)
    c.add(['Scenariusze ręczne',scenario_count,scenario_count,True,'Liczba wierszy JSON; wyniki ręczne nadal Nie wykonano'],42)
    sheets.append(c)
    src=Sheet(HIDDEN[1],'ŹRÓDŁA V11 · wartości pobrane z kodu',
      'Snapshot eksportów obecnego kodu; opublikowane PRE-PROD i nowy zakres DEV są rozdzielone. Wszystkie stare arkusze i komentarze pochodzą z niezmienionego Excela 10.',
      ['Obszar / klucz','Wartość','Źródło'],[60,140,70],hidden=True)
    src.add(['PRE-PROD / źródło',PUBLISHED['commit'],'AGENTS.md; rzeczywisty workflow 41112673285'],42)
    src.add(['PRE-PROD / CI',PUBLISHED['ci'],'GitHub Actions — SUCCESS'],42)
    src.add(['PRE-PROD / wdrożenie',PUBLISHED['deployment'],'GitHub Actions — SUCCESS; nie dotyczy nowej partii DEV'],42)
    src.add(['Excel 10 SHA-256',SOURCE_SHA256,'Niezmieniony plik bazowy'],42)
    src.add(['Nowe dokumenty / operator','Loremipsum','Placeholder zatwierdzony przez właściciela; treść robocza do przeglądu'],42)
    for key,value in facts.items():
        if key=='lessons':continue
        src.add([key,json.dumps(value,ensure_ascii=False,sort_keys=True),'Eksport modułu Node; scripts/build_workbook_v11.py'],min(240,48+len(str(value))//80*14))
    sheets.append(src)
    for sheet in sheets:
        for row,values in sheet.rows.items():
            for col,value in enumerate(values,1):
                if isinstance(value,Formula):value.cached=evaluate(value.text,sheets,sheet.name,((sheet.name,f'{get_column_letter(col)}{row}'),))
    assert all(values[3] is True or isinstance(values[3],Formula) and values[3].cached is True for values in c.rows.values())


def xml(tag,**attrs):return E.Element(f'{{{M}}}{tag}',**{k:str(v)for k,v in attrs.items()})


def append_cell(row,coordinate,value,style):
    cell=E.SubElement(row,f'{{{M}}}c',r=coordinate,s=str(style))
    if isinstance(value,Formula):
        E.SubElement(cell,f'{{{M}}}f').text=value.text.lstrip('=')
        value=value.cached
        if isinstance(value,bool):cell.set('t','b');E.SubElement(cell,f'{{{M}}}v').text='1' if value else '0'
        elif isinstance(value,(int,float)):E.SubElement(cell,f'{{{M}}}v').text=format(value,'.17g')
        else:cell.set('t','str');E.SubElement(cell,f'{{{M}}}v').text=str(value)
    elif isinstance(value,bool):cell.set('t','b');E.SubElement(cell,f'{{{M}}}v').text='1' if value else '0'
    elif isinstance(value,(int,float)):
        assert math.isfinite(value)
        E.SubElement(cell,f'{{{M}}}v').text=format(value,'.17g')
    elif value is not None:
        cell.set('t','inlineStr');text=E.SubElement(E.SubElement(cell,f'{{{M}}}is'),f'{{{M}}}t');text.set(f'{{{XML}}}space','preserve');text.text=str(value)


def sheet_xml(s,index):
    root=xml('worksheet');last=max(s.rows)
    E.SubElement(root,f'{{{M}}}sheetPr');E.SubElement(root,f'{{{M}}}dimension',ref=f'A1:{get_column_letter(len(s.headings))}{last}')
    views=E.SubElement(root,f'{{{M}}}sheetViews');view=E.SubElement(views,f'{{{M}}}sheetView',workbookViewId='0',tabSelected='1' if s.name==VISIBLE[0] else '0')
    E.SubElement(view,f'{{{M}}}pane',ySplit='5',xSplit='2' if s.name==VISIBLE[4] else '0',topLeftCell='C6' if s.name==VISIBLE[4] else 'A6',activePane='bottomRight' if s.name==VISIBLE[4] else 'bottomLeft',state='frozen')
    E.SubElement(root,f'{{{M}}}sheetFormatPr',defaultRowHeight='22')
    cols=E.SubElement(root,f'{{{M}}}cols')
    for i,width in enumerate(s.widths,1):E.SubElement(cols,f'{{{M}}}col',min=str(i),max=str(i),width=str(width),customWidth='1')
    data=E.SubElement(root,f'{{{M}}}sheetData')
    # These are raw OOXML cellXfs indices from immutable Excel10, not the
    # normalized .style_id values that openpyxl may allocate when inspecting.
    for n,values,height,style in [(1,[s.title],34,1),(2,[s.note],76,2),(5,s.headings,34,3)]:
        row=E.SubElement(data,f'{{{M}}}row',r=str(n),ht=str(height),customHeight='1')
        for col,value in enumerate(values,1):append_cell(row,f'{get_column_letter(col)}{n}',value,style)
    for n,values in sorted(s.rows.items()):
        row=E.SubElement(data,f'{{{M}}}row',r=str(n),ht=str(s.heights.get(n,72)),customHeight='1')
        for col in range(1,len(s.headings)+1):
            coord=f'{get_column_letter(col)}{n}';value=values[col-1] if len(values)>=col else None
            style=3 if n in s.section_rows else 6 if coord in s.inputs or s.name==VISIBLE[4] and col==12 else 4 if n%2==0 else 5
            append_cell(row,coord,value,style)
    E.SubElement(root,f'{{{M}}}autoFilter',ref=f'A5:{get_column_letter(len(s.headings))}{last}')
    merges=E.SubElement(root,f'{{{M}}}mergeCells',count=str(2+len(s.section_rows)))
    for n in [1,2,*sorted(s.section_rows)]:E.SubElement(merges,f'{{{M}}}mergeCell',ref=f'A{n}:{get_column_letter(len(s.headings))}{n}')
    if s.validations:
        ds=E.SubElement(root,f'{{{M}}}dataValidations',count=str(len(s.validations)))
        for attrs in s.validations:
            attrs=attrs.copy();f1=attrs.pop('formula1');f2=attrs.pop('formula2',None)
            d=E.SubElement(ds,f'{{{M}}}dataValidation',**attrs,allowBlank='0',showErrorMessage='1',errorTitle='Sprawdź wartość',error='Wpisz wartość w dozwolonym zakresie.')
            E.SubElement(d,f'{{{M}}}formula1').text=f1
            if f2 is not None:E.SubElement(d,f'{{{M}}}formula2').text=f2
    E.SubElement(root,f'{{{M}}}pageMargins',left='0.25',right='0.25',top='0.5',bottom='0.5',header='0.2',footer='0.2')
    E.SubElement(root,f'{{{M}}}pageSetup',orientation='landscape',paperSize='8',fitToWidth='1',fitToHeight='0')
    return E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def build():
    assert digest(BASE)==SOURCE_SHA256,'Excel 10 changed; preserve its reviewed bytes'
    facts=game_facts();data=json.loads(SCENARIOS.read_text());scenarios=data['scenarios']
    assert len(scenarios)>=80 and len({x['id']for x in scenarios})==len(scenarios)
    assert all(x['status']=='Nie wykonano'for x in scenarios)
    sheets=prepare_sheets(facts,scenarios);add_checks(sheets,facts,len(scenarios))
    with ZipFile(BASE)as source:
        entries={name:source.read(name)for name in source.namelist()}
    workbook=E.fromstring(entries['xl/workbook.xml']);sheet_nodes=workbook.find(f'{{{M}}}sheets');old_count=len(sheet_nodes)
    assert old_count==59
    for node in sheet_nodes:node.set('state','hidden')
    for view in workbook.findall(f'{{{M}}}bookViews/{{{M}}}workbookView'):view.set('activeTab',str(old_count));view.set('firstSheet',str(old_count))
    calc=workbook.find(f'{{{M}}}calcPr')
    if calc is None:calc=E.SubElement(workbook,f'{{{M}}}calcPr')
    calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
    rels=E.fromstring(entries['xl/_rels/workbook.xml.rels']);types=E.fromstring(entries['[Content_Types].xml'])
    next_rid=max(int(node.get('Id')[3:])for node in rels if node.get('Id','').startswith('rId'))+1
    next_sheetid=max(int(node.get('sheetId'))for node in sheet_nodes)+1
    for n,sheet in enumerate(sheets,1):
        number=old_count+n;rid=f'rId{next_rid+n-1}';name=f'xl/worksheets/sheet{number}.xml'
        E.SubElement(sheet_nodes,f'{{{M}}}sheet',name=sheet.name,sheetId=str(next_sheetid+n-1),state='hidden' if sheet.hidden else 'visible',**{f'{{{R}}}id':rid})
        E.SubElement(rels,f'{{{P}}}Relationship',Id=rid,Type=f'{R}/worksheet',Target=f'worksheets/sheet{number}.xml')
        E.SubElement(types,f'{{{C}}}Override',PartName='/'+name,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
        entries[name]=sheet_xml(sheet,number)
    changed={'xl/workbook.xml':workbook,'xl/_rels/workbook.xml.rels':rels,'[Content_Types].xml':types}
    core=E.fromstring(entries['docProps/core.xml']);dc='http://purl.org/dc/elements/1.1/'
    for name,value in [('title','HARBORLIGHT11 — zasady v18 i scenariusze testów'),('subject','Bieżące v18; nowy zakres DEV; niezmienione archiwum Excel 10'),('description','Pięć widocznych zakładek v11. Zachowane 59 starych arkuszy, 48 komentarzy i wszystkie historyczne cache formuł.')]:
        node=core.find(f'{{{dc}}}{name}')
        if node is None:node=E.SubElement(core,f'{{{dc}}}{name}')
        node.text=value
    changed['docProps/core.xml']=core
    for name,node in changed.items():entries[name]=E.tostring(node,xml_declaration=True,encoding='UTF-8',standalone=True)
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    with ZipFile(OUTPUT,'w',compression=ZIP_DEFLATED,compresslevel=9)as output:
        for name,value in entries.items():output.writestr(name,value)
    return facts,sheets


def validate():
    assert digest(BASE)==SOURCE_SHA256
    facts=game_facts();scenarios=json.loads(SCENARIOS.read_text())['scenarios'];sheets=prepare_sheets(facts,scenarios);add_checks(sheets,facts,len(scenarios))
    mutable={'xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml','docProps/core.xml'}
    with ZipFile(BASE)as old,ZipFile(OUTPUT)as new:
        assert new.testzip()is None
        for name in old.namelist():
            if name not in mutable:assert old.read(name)==new.read(name),f'Historical ZIP part changed: {name}'
    old=load_workbook(BASE,data_only=False);new=load_workbook(OUTPUT,data_only=False);values=load_workbook(OUTPUT,data_only=True)
    assert len(old.sheetnames)==59 and len(new.sheetnames)==66
    assert [s.title for s in new if s.sheet_state=='visible']==VISIBLE
    assert new.active.title==VISIBLE[0]
    comments=0;historical_formulas=0;new_formulas=0
    old_values=load_workbook(BASE,data_only=True)
    for sheet in old:
        assert new[sheet.title].sheet_state=='hidden'
        for row in sheet:
            for cell in row:
                restored=new[sheet.title][cell.coordinate]
                assert restored.value==cell.value,(sheet.title,cell.coordinate)
                assert values[sheet.title][cell.coordinate].value==old_values[sheet.title][cell.coordinate].value,(sheet.title,cell.coordinate,'cached value')
                if cell.data_type=='f':historical_formulas+=1
                if cell.comment:
                    comments+=1;assert restored.comment and restored.comment.text==cell.comment.text and restored.comment.author==cell.comment.author
    assert comments==48
    for sheet in sheets:
        assert new[sheet.name]['A1'].value==sheet.title
        for n,row in sheet.rows.items():
            for col,value in enumerate(row,1):
                cell=new[sheet.name].cell(n,col);cached=values[sheet.name].cell(n,col)
                if isinstance(value,Formula):
                    new_formulas+=1;assert cell.value==value.text
                    assert cached.data_type!='e' and cached.value is not None
                    if isinstance(value.cached,(int,float))and not isinstance(value.cached,bool):assert math.isclose(cached.value,value.cached,rel_tol=1e-12,abs_tol=1e-9)
                    else:assert cached.value==value.cached
                elif isinstance(value,(int,float))and not isinstance(value,bool):assert math.isclose(cell.value,value,rel_tol=1e-12,abs_tol=1e-9),(sheet.name,n,col,cell.value,value)
                else:assert cell.value==value,(sheet.name,n,col,cell.value,value)
    for row in values[HIDDEN[0]].iter_rows(min_row=6,values_only=True):assert row[3]is True,('Failed current check',row)
    assert new[VISIBLE[4]].max_row-5==len(scenarios)
    assert all(new[VISIBLE[4]].cell(i,12).value=='Nie wykonano'for i in range(6,new[VISIBLE[4]].max_row+1))
    assert sum(1 for row in new[VISIBLE[0]].iter_rows(values_only=True)if row[0]and re.match(r'^\d+\. ',str(row[0])))==25
    # Independently vary the editable values and compare the formula model with
    # actual runtime exports, particularly the warehouse L24 cap and asymmetric
    # demography. This checks the formulas beyond their stored initial caches.
    calc=next(s for s in sheets if s.name==VISIBLE[3]);samples=0
    def sample(input_row,input_value,outputs):
        nonlocal samples
        previous=calc.rows[input_row][2];calc.rows[input_row][2]=input_value
        try:
            for output_row,expected in outputs.items():
                formula=calc.rows[output_row][2];actual=evaluate(formula.text,sheets,calc.name)
                if isinstance(expected,(int,float)):assert math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-8),(input_row,input_value,output_row,actual,expected)
                else:assert actual==expected
                samples+=1
        finally:calc.rows[input_row][2]=previous
    for level,rate in facts['production'].items():sample(7,int(level),{17:rate})
    for level,limits in facts['warehouses'].items():sample(8,int(level),{23:limits['load'],28:limits['perResource'],29:limits['total'],30:limits['reserve'],31:limits['workers']})
    for happiness,result in facts['demography'].items():sample(42,float(happiness),{43:result['rate'],44:result['interval'] if result['interval']is not None else 'brak zmiany'})
    for stage,rules in facts['rounds'].items():sample(47,int(stage),{48:rules['tilt']['baseReward'],50:rules['tilt']['timeLimit'],51:rules['triples']['cardCount'],52:rules['triples']['timeLimit']})
    for level in [1,2]:sample(66,level,{67:'nie dotyczy',68:'nie dotyczy',69:'nie dotyczy'})
    report={'version':11,'as_of':'2026-10-30','published_preprod':PUBLISHED,'new_scope':'DEV; logout and draft legal documents require own acceptance/promotion','source_file':BASE.name,'source_sha256':SOURCE_SHA256,'output':OUTPUT.name,'sha256':digest(OUTPUT),'bytes':OUTPUT.stat().st_size,'sheets':len(new.sheetnames),'visible':VISIBLE,'historical_sheets_unchanged':59,'source_comments_unchanged':comments,'historical_cached_formulas_unchanged':historical_formulas,'new_cached_formulas_checked':new_formulas,'runtime_formula_samples':samples,'current_checks':values[HIDDEN[0]].max_row-5,'manual_scenarios':len(scenarios),'manual_results':'Nie wykonano','cache_method':'Old OOXML parts copied byte-for-byte; new formulas evaluated independently; no LibreOffice recalculation claimed','validation':'PASS'}
    REPORT.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))
    return report


def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--validate-only',action='store_true');args=parser.parse_args()
    if not args.validate_only:build()
    validate()


if __name__=='__main__':main()
