#!/usr/bin/env python3
"""Create simplified Excel 9 from immutable Excel 8 and decisions-v9.json.

Four visible sheets; all 45 source sheets, their formulas and 48 comments are
retained as hidden history. openpyxl and LibreOffice are required to publish.
"""
from __future__ import annotations
import argparse
from copy import copy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[1]
BASE=ROOT/'docs/arkusze/HARBORLIGHT8-zasady.xlsx'
OLD7=ROOT/'docs/arkusze/HARBORLIGHT7-zasady.xlsx'
OUTPUT=ROOT/'docs/arkusze/HARBORLIGHT9-zasady.xlsx'
MANIFEST=ROOT/'docs/decisions-v9.json'
HASH7='bf49eef3b43ba81e0492c6eb187940c99177bbb8ae925ceca2b09ab0a2b9e945'
HASH8='74ea3cc586d8d889ff34de8819694430503f45ed5fa2743873c9ad0c0b83cc3e'
NAVY,TEAL,GOLD,LIGHT,BLUE='123140','146D75','F0CE80','F0F6F8','135CAA'
VISIBLE=['Jak działa gra','Budynki v9','Do decyzji','Kalkulator']

# Each explanation has at most three concise sentences. Detailed exceptions are
# in the hidden decision manifest; historical sheets never override version 9.
MECHANICS=[
('Start','Zaczynasz z 12 osobami w domu, Ratuszem 1 i po 100 każdego z sześciu zasobów. Nie ma przydzielonych pracowników, a szczęście wynosi 0%. Oddzielny profil nowego gracza nie kasuje poprzedniej wyspy.','Start / oddzielny zapis'),
('Poziom i mieszkanie','Poziom Ratusza jest poziomem gry i ogranicza ulepszanie innych budynków. Ratusz daje limit 24 mieszkańców na poziom, a Dom osobno zapewnia 24 łóżka na swój poziom. Brak łóżek zmniejsza szczęście wszystkich o 1 pp/h za każdą osobę bez miejsca.','Reguły V7/V8 zachowane'),
('Przydział i podróż','Każde z ośmiu miejsc pracy ma limit 12 × własny poziom, ze wspólnej puli wolnych osób w domu. Do pracy i z powrotem idzie się po 5 minut. Po zwolnieniu osoba bez ładunku od razu rozpoczyna powrót.','V9-19 / F04'),
('Tragarz z ładunkiem','Odwołany tragarz najpierw kończy rozpoczętą dostawę i do tego czasu nadal pracuje. Nie bierze kolejnego kursu, potem wraca do domu przez 5 minut. Dopiero powrót czyni go dostępnym.','V9-19'),
('Produkcja','Poziomy 1, 2, 3, 4 dają na osobę 10, 12, 13, 14 sztuk/h, a dalej +1 sztukę/h na poziom. Każda osoba ma własny postęp gotowej sztuki. Zwolnienie kasuje jej niedokończoną sztukę.','V9-50'),
('Młyn i tartak','Młyn zużywa 2 zboża z centrali na 1 żywność, a tartak 2 drewna na 1 deskę. Surowiec schodzi dopiero przy ukończeniu produktu. Brak wejścia lub miejsca wstrzymuje postęp bez jego utraty.','V9-50'),
('Dwa magazyny','Gotowe produkty czekają lokalnie: pojemność wynosi 200 × własny poziom producenta. Dopiero tragarze zanoszą je do centrali, skąd można je wydawać i karmić ludzi. Monety są osobno w skarbcu.','V8 / V9-45'),
('Pojemność centrali','Magazyn 1 mieści 1000 każdego materiału i 4000 łącznie; oba limity rosną ×1,5 na poziom. 50 × poziom Magazynu pojemności łącznej jest zarezerwowane dla żywności. Materiały przypisane budowie nadal zajmują miejsce aż do odbioru przez tragarza.','V9-14,45'),
('Jedzenie','Każdy mieszkaniec zużywa 1 żywność/h, automatycznie i równomiernie z centrali. Jedna sztuka wystarcza 12 osobom na 5 minut. Po wyczerpaniu posiłku wszystkich dotyczy brak jedzenia.','V9-46'),
('Głód i ponowne karmienie','Brak jedzenia od razu zmniejsza wydajność do połowy; po 12 h nalicza −5 pp/h. Posiłek od razu zatrzymuje karę i przywraca wydajność. Dopiero pełna godzina nieprzerwanego karmienia zeruje wcześniejszy licznik głodu.','V9-1'),
('Szczęście','Każda osoba ma własne szczęście ograniczone do −100…100%, a wyspa pokazuje średnią. Jedzenie i rzeczywista praca dodają osobno +1 pp/h; brak pracy po 12 h daje −5 pp/h. Prezentacja zaokrągla w górę do pełnego procentu, lecz rachunek zachowuje ułamki.','V7/V8 + V9-1,2'),
('Nowy mieszkaniec','Nowa osoba przez pierwsze 24 h ma zamrożone 0% szczęścia. Nadal je i może pracować, ale liczniki głodu i bezrobocia zaczynają działać dopiero po tym czasie. Startowa dwunastka nie ma zamrożenia.','V9-2'),
('Przestój','Brak wejścia lub miejsca zatrzymuje produkcję i jej premię szczęścia za pracę. Kara wynosi −10 pp od początku, kolejne −10 co 24 h, maksymalnie −50 z tego składnika. Nie dodajemy za ten sam przestój drugiej kary bezrobocia.','V7 D14 / V9-50'),
('Przyrost i odejścia','Dokładna średnia szczęścia wyznacza tempo przyrostu albo odejść; minimum populacji to 6. Zachowujemy kolejność grup odejścia, a w wybranej grupie odchodzi najmniej szczęśliwy, z losowaniem remisów. Gdy wszyscy są chronieni, ubytek przepada i nie tworzy długu.','V9-48,49'),
('Punkty Domu','O północy Dom dostaje pulę pp równą aktualnej populacji, bez przenoszenia starej reszty. O 23:58 automat rozdaje niewykorzystane punkty najmniej szczęśliwym uprawnionym osobom. Punkty nie usuwają kar ani ich liczników.','V9-57'),
('Kaplica','Ulepszenie daje +10 × nowy poziom pp na 10 h; kolejna rozbudowa zastępuje poprzedni efekt pełnym nowym czasem. Ręczne błogosławieństwo daje +50 każdego zasobu i +20 pp na godzinę, bez nakładania dwóch ręcznych efektów. Darmowe użycia zależą od poziomu; płatne kosztuje 100 × poziom i ma 4 h blokady od końca poprzedniego płatnego efektu.','V9-4,56; mechaniczne podatki w szczegółach'),
('Darmowe błogosławieństwo','Darmowego użycia można dokonać podczas blokady płatnego, jeżeli inne błogosławieństwo już nie trwa. Darmowe użycie nie przesuwa terminu następnego płatnego. Błogosławieństwa pozostają odrębne od nieustalonej aktywności modlitewnej.','V9-56'),
('Latarnia','Ulepszenie daje wymuszone 100% szczęścia i produkcję ×2 przez 24 h × osiągnięty poziom. Każdy efekt ma własny zegar, a nakładające się wymuszenia 100% korzystają z jednego zapisanego stanu sprzed pierwszego. Stan wraca po końcu ostatniego wymuszenia, z nadal aktywnymi dodatkami i karami.','V9-3,4'),
('Zwiedzanie latarni','Wizyta trwa 4 h bez dojścia i wycofania; jednocześnie są 2 × poziom latarni miejsca. Po powrocie uczestnik ma 100% H i blokadę wizyty na 24 h, a jego grupa daje +10 pp populacji na 6 h. Pula darmowych wejść 2L działa w ruchomych 24 h, a trzy dodatkowe pakiety według zatwierdzonego cennika.','V7 D01–D02 + V9-3'),
('Atrakcje','O północy pojawia się 5 ofert z katalogu 20 dekoracji i wydarzeń. Atrakcja daje Y pp na 6 h osobom o najniższym uprawnionym H, najwyżej tylu, jaki jest poziom Ratusza; cena wynosi Y × faktyczna liczba odbiorców. Y jest losową liczbą całkowitą 5–100, a brak odbiorców blokuje zakup.','V9-5,6,8'),
('Przyzwyczajanie','Każdy rodzaj atrakcji ma osobny ciąg skuteczności: 100%, 75%, 50%, 25%, 0%. Po 24 h bez zakupu tego rodzaju wraca pełna skuteczność, a przy 0% nie pobieramy opłaty. Każdy zakup ma własne 6 h, a efekty sumują się przed ograniczeniem H.','V9-7,8 + V8'),
('Ścieżki','Jest 100 odcinków na poziomach 0–5; jedno ulepszenie kosztuje 100 kamienia i działa po kliknięciu. Kolejny poziom wolno zacząć po ukończeniu wszystkich 100 odcinków poprzedniego. Pełne poziomy dają +10…50% szybkości tragarzy, a poziomy 3/4/5 również +10/20/30% ładunku.','V8 + F10'),
('Kursy i priorytety','Bazowy pełny kurs trwa 5 min i zabiera jeden rodzaj, do 60 sztuk na poziomie miejsca pracy 1. Trzy trwałe priorytety towarów/Construction dzielą kursy 8+2 losowe, 5+3+2 losowe albo 5/3/2, z ochroną ratunkowej żywności. Construction wysyła najpierw wszystkie ładunki najstarszej aktywnej budowy, a potem następnej.','V9-20 + V8'),
('Rozbudowa','Ratusz L daje L aktywnych budów i L miejsc kolejki; kolejka kosztuje +10% każdego składnika. T biegnie od pierwszego odbioru materiałów do ostatniej dostawy, także przez przerwy; obsada zmienia ETA i daje powiadomienie. Sama budowa trwa T, latarnia 4T, kaplica 3T × (nowy poziom−1).','V9-9,45'),
('Plany i anulowanie','Pakiet 1000 × poziom Ratusza/24 h daje +4 budowy, dowolną kolejkę i planowanie braków za 25% ich bazowej wartości portowej po +10% kolejki. Dostępne materiały przypisujemy planom FIFO, bez podwójnego zastawiania. Anulować wolno przed pierwszym przewozem/budową, ze zwrotem rzeczywistych rezerwacji, lecz bez opłaty prognozy.','V9-11,12'),
('Nagrody i zwroty','Materiały wypełniają centralę zgodnie z limitami i przestrzenią żywności. Nadwyżka czeka bezterminowo i wchodzi automatycznie bez tragarzy. Przed przyjęciem nadwyżki nie można wydać.','V9-13,14,44'),
('Przyspieszanie za monety','Skrócenie 25/50/75/100% budowania kosztuje 1 monetę za pełną minutę oszczędności, minimum 5 monet. Natychmiastowy przewóz kosztuje 10 monet za każdy potrzebny kurs, także niepełny, i dotyczy tylko wybranego etapu. Budowanie i dostawa są osobnymi usługami.','V9-15,54'),
('Premium dowóz budowy','Dostawa Premium przywozi brakujące materiały od razu i zamraża budowanie według minionego T oraz normalnej estymaty reszty. Bez magazynierów przyjmujemy w kalkulacji jednego hipotetycznego tragarza z obecnymi parametrami. Sama faza budowania nie znika i może być skrócona oddzielnie.','V9-54,55'),
('Pełne kupno budynku','Kompletny budynek kosztuje portową wartość wszystkich materiałów ×1,25, nominalne monety i potrzebne kursy Premium. Jest gotowy po 10 s w osobnej usłudze, poza zwykłą kolejką i bez pobrania własnych materiałów. Obowiązują limit Ratusza i kolejność poziomów tego samego budynku.','V9-10'),
('Targ','Od Ratusza 2 wymieniasz pełne dziesiątki tylko w grupie drewno/kamień/zboże albo deski/żywność: 10 daje 9. Dostawa jest zablokowana 1 h; fizyczna partia ma ważność 24 h, a wynik wymiany dostaje nowy termin. Wymiana działa bez obsady, lecz to własna załoga nosi towary w obie strony.','V9-51,52,53'),
('Port','Od Ratusza 4 handlujesz paczkami 100: podstawowe sprzedaż/kupno 40/50, przetworzone 140/150 monet. Każda kolejna paczka kupna podnosi cenę liniowo o 10% bazowej ceny, osobno dla towaru w ruchomych 3 h. Cena zamraża się przy transakcji, a monety za sprzedaż przychodzą przy wypłynięciu.','V7 D16'),
('Ważność towaru','Każda fizyczna partia targowa i portowa ma 24 h od przybycia. Faktyczne podjęcie ładunku przed terminem chroni rozpoczęty kurs. Samo zlecenie odbioru nie chroni towaru pozostawionego na miejscu.','V9-52'),
('Zadania latarni','Są 3 oferty, ale tylko 1 aktywne zadanie i nie można go porzucić. Potwierdzenie ma 60 s; mnożnik utrwala otwarcie okna: ×3 przed 15 s, ×2 przed 5 min, później ×1. Oferta musi nadal być ważna przy potwierdzeniu, a zamknięcie usuwa szkic.','V9-16,17'),
('Darowizny i terminy','Towary przekazujesz w nieodwracalnych ratach bez tragarzy z centrali, zachowując rezerwę jedzenia. Pierwsze zadania trwają 6 h; po pierwszym sukcesie w ciągu następcy mają termin o północy, nawet gdy pojawiają się wcześniej po kolejnym sukcesie. Norma kosztów i skalowanie nagród z poziomu 3 pozostają zachowane.','V9-47 + V8'),
('Ukończenie i Share','Sukces zadania przyznaje nagrodę, pokazuje fajerwerki i krótką wibrację obsługiwanego telefonu. Share daje osobne 500 monet raz, a Skip zachowuje normalną nagrodę. Nierozstrzygnięta propozycja pozostaje bezterminowo w powiadomieniach, bez rzeczywistej integracji Facebook.','V9-18'),
('Pula ekspedycji','Nowa bezpłatna seria zaczyna od pustej puli; każda terminowa wygrana dodaje 50 monet. Dalej ryzykujesz całą niewypłaconą pulę, a przegrana ją zeruje. Wypłata kończy serię; wcześniej wypłaconych monet nie tracisz.','V9-21,22,23,31,58'),
('Zegar ekspedycji','Zegar działa także w tle i podczas pytania o obecność; po limicie potencjalna pula maleje o 10 monet za pełną sekundę spóźnienia. Po 5 min bez aktywności jest pytanie i 60 s na odpowiedź, po czym seria przegrywa. Jedna seria blokuje wyspę do wypłaty albo przegranej; dokładne plansze i czasy to parametry projektu.','V9-24…30,58'),
('Premie społecznościowe','Docelowo pierwsza publikacja daje 1000 monet raz, awans od Ratusza 2 daje 500 raz na poziom, a każda trzecia kwalifikowana publikacja 1000. Liczymy najwyżej jeden wpis dziennie ze wszystkich platform, a spełnione premie mogą się sumować. W grze są reguły i sprawdzone punkty integracji, bez prawdziwego potwierdzania zewnętrznych publikacji.','V9-37…41'),
('Polecenia','Nowe konto z indywidualnego linku ma jednego stałego polecającego, który dostaje 1000 monet po pierwszym logowaniu. Trzy aktywne dni w ciągu 30 dni od rejestracji dają polecającemu po 500 pięciu materiałów, bez monet w tej drugiej premii. Prawdziwe konta i potwierdzenia wymagają backendu.','V9-33,42,43'),
('Lojalność','Liczymy kolejne dni logowania, a przerwa rozpoczyna serię od początku. Program kończy się przy Ratuszu 4 albo po odebraniu dziesiątej dziennej nagrody. Wartości nagród nie zostały ustalone, więc nie przyznajemy fikcyjnych kwot.','V9-35,36'),
('Samouczek','25 mikrozadań pojawia się po 5 w pierwszych 5 aktywnych dniach gry; postęp i odbiór nagród są trwałe. Uczą obsady, jedzenia, przewozu, szczęścia, budów i pozostałych podstaw. Kwoty dobrano projektowo: łącznie 1000 monet i 400 materiałów, niezależnie od nieustalonej lojalności.','F07 / jawny dobór projektowy'),
('Aktywna gra','Rzeczywista akcja w widocznej aktywnej karcie odnawia bonus produkcji i podatków ×2 na 5 minut. Czas się nie sumuje; z latarnią produkcja może wynosić ×4, z kaplicą podatki ×4. Bonus nie przyspiesza transportu, a informacje podatkowe pozostają ukryte w interfejsie.','F08 / jawne założenie przed GO'),
('Kupno osoby','Osoba kosztuje 500 × obecny poziom Ratusza każdego z sześciu zasobów, jak najbliższa rozbudowa Ratusza. Potrzebne są wolne łóżko i miejsce w limicie ratusza, a osoba zaczyna z zasadami pierwszych 24 h. Tę cenę oznaczamy jako ogłoszone założenie wdrożeniowe, nie historyczną cenę ustaloną w V7.','F16 / jawne założenie przed GO'),
('Punkty i zasoby','Punkty to migawka poziom Ratusza × monety przeliczana o pełnej godzinie. Strzałka ↑/↓/→ pokazuje bilans godzinowy produkcji i zużycia całej wyspy. Samo przeniesienie towaru między magazynami nie zmienia tego bilansu.','F18,F19'),
('Dzień i zapis','Odnowienia dzienne i dni aktywności używają 00:00 Europe/Warsaw, także offline. Czasy trwania efektów 6 h i 24 h, ruchome okna oraz 30 dni nadal liczymy od zdarzenia. Przeglądarka zachowuje oba oddzielne profile, bez synchronizacji z prawdziwym kontem.','V9-34'),
('Nowy interfejs','Nie ma lewego paska ani pięciu materiałów w nagłówku; są pracujący, ETA nowej osoby i Punkty. Kafelki budynków mają ustaloną kolejność i stany osób, a zasoby pokazują lokalny / centralny (wszystkie przewozy w toku). Ustawienia obejmują EN/PL, filtry powiadomień oraz animacje, a rozwijane szczegóły efektu są pod i.','F01…F06,F09…F15,F17'),
]

BUILDINGS=[
('Ratusz','Town Hall','Poziom gry i limit innych budynków','0','24 × L mieszkańców; L budów + L kolejki','500 × L każdego z 6 zasobów','T','Efekty u góry; zarządzanie budową tylko tutaj'),
('Dom','House','Łóżka, mieszkańcy i pula szczęścia','0','24 × L łóżek; dzienna pula według populacji','500 × L każdego z 6 zasobów','T','Zbiorcze zarządzanie pracą tylko tutaj; automat 23:58'),
('Farma','Farm','Zboże','0 / 12 × własny L','10/12/13/14… szt./os./h; bufor 200L','L × (72 drewna +72 kamienia +120 zboża)','T','Osobne postępy; magazyn lokalny i centrala'),
('Młyn','Windmill','2 zboża → 1 żywność','0 / 12 × własny L','Wydajność i bufor własnego poziomu','2L × (72 drewna +72 kamienia +120 żywności)','T','Wejście pobierane przy gotowej sztuce'),
('Kamieniołom','Quarry','Kamień','0 / 12 × własny L','Wydajność i bufor własnego poziomu','2L × (72 drewna +192 kamienia)','T','Suma kamienia konstrukcyjnego i własnego produktu'),
('Drwale','Lumber Camp','Drewno','0 / 12 × własny L','Wydajność i bufor własnego poziomu','2L × (192 drewna +72 kamienia)','T','Suma drewna konstrukcyjnego i własnego produktu'),
('Tartak','Sawmill','2 drewna → 1 deska','0 / 12 × własny L','Wydajność i bufor własnego poziomu','2L × (72 drewna +72 kamienia +120 desek)','T','Wejście pobierane przy gotowej sztuce'),
('Magazyn','Warehouse','Centrala i kursy produkcji/budów','0 / 12 × własny L','Pojemność ×1,5; ładunek ×1,1/poziom','300 × L każdego z 6 zasobów','T','Miejsce 50L dla Food; fizyczne rezerwacje zajmują skład'),
('Targ','Market','Barter w 2 grupach, od Ratusza 2','0 / 12 × własny L','1000 łącznie na L1, dalej ×1,5','300 × L każdego z 6 zasobów','T','10→9; 1 h blokady; własny transport; wymiana przy 0 obsady'),
('Port','Harbor','Handel monetowy od Ratusza 4','0 / 12 × własny L','1000 na rodzaj /4000 łącznie na L1, dalej ×1,5','400 × L każdego z 6 zasobów','T','Paczki 100; 2 etapy transportu; osobne partie 24 h'),
('Kaplica','Chapel','Premia rozbudowy i błogosławieństwa','0','+10 × N pp oraz mechaniczne podatki ×2 na 10 h','N × (180 monet +20 drewna +60 kamienia +10 desek)','3T × (N−1)','Darmowe użycie nie przesuwa blokady płatnego'),
('Latarnia','Lighthouse','Premia rozbudowy, wizyty i zadania','0','100% H i produkcja ×2 przez 24 h × N','500 × max(1,L) każdego z 6 zasobów','4T','Wizyty 4 h, osobiste 24 h, grupa +10 pp/6 h'),
]


def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def cell_style(c,header=False,input=False):
 c.font=Font(name='Calibri',size=11 if not header else 10,bold=header,color='FFFFFF' if header else BLUE if input else NAVY)
 c.fill=PatternFill('solid',fgColor=TEAL if header else 'EAF2FF' if input else LIGHT if c.row%2==0 else 'FFFFFF')
 c.alignment=Alignment(vertical='top',wrap_text=True);c.border=Border(bottom=Side(style='hair',color='D5E2E6'))
def sheet(w,name,title,desc,headers,rows,widths):
 s=w.create_sheet(name);n=len(headers)
 for r in (1,2):s.merge_cells(start_row=r,start_column=1,end_row=r,end_column=n)
 s['A1']=title;s['A2']=desc
 for r in (1,2):
  c=s.cell(r,1);c.font=Font(name='Calibri',size=19 if r==1 else 11,bold=r==1,color=GOLD if r==1 else NAVY);c.fill=PatternFill('solid',fgColor=NAVY if r==1 else LIGHT);c.alignment=Alignment(wrap_text=True,vertical='center')
 s.row_dimensions[1].height=35;s.row_dimensions[2].height=68
 for ci,h in enumerate(headers,1):cell_style(s.cell(5,ci,h),True)
 s.row_dimensions[5].height=30
 for ri,values in enumerate(rows,6):
  for ci,v in enumerate(values,1):cell_style(s.cell(ri,ci,v))
  s.row_dimensions[ri].height=max(38,min(130,18*(max((len(str(x)) for x in values if x is not None),default=0)//95+1)))
 for ci,width in enumerate(widths,1):s.column_dimensions[get_column_letter(ci)].width=width
 s.freeze_panes='B6';s.auto_filter.ref=f'A5:{get_column_letter(n)}{s.max_row}';s.sheet_view.showGridLines=False;s.sheet_properties.tabColor=TEAL
 s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.sheet_properties.pageSetUpPr.fitToPage=True;s.print_title_rows='1:5'
 return s


def tutorial_rows():
 code=(ROOT/'harborlight-v1-progression.js').read_text()
 pattern=r"lesson\('([^']+)', (\d+), '([^']+)', '([^']+)', '([^']+)', '([^']+)', '([^']+)'(?:, '([^']+)')?\)"
 entries=re.findall(pattern,code)
 if len(entries)!=25:raise AssertionError(f'Expected25tutoriallessons,found{len(entries)}')
 return [(id,int(day),title,description,kind,subject,pre or '—',25 if int(day)<=2 else 50,resource,10 if int(day)<=2 else 20,'DOBÓR PROJEKTOWY: nie są to nagrody lojalnościowe') for id,day,title,description,kind,subject,resource,pre in entries]


def build():
 assert sha(OLD7)==HASH7 and sha(BASE)==HASH8,'Source7/8changed'
 manifest=json.loads(MANIFEST.read_text());assert len(manifest['decisions'])==58 and len(manifest['frontend'])==19
 w=load_workbook(BASE)
 for s in w:s.sheet_state='hidden';s.sheet_properties.tabColor='AAB5B9'
 m=sheet(w,'Jak działa gra','HARBORLIGHT9 • jak działa gra','Czytaj tylko cztery widoczne zakładki. Każda mechanika ma najwyżej trzy zdania; 45 zakładek wersji 8 zachowano jako ukryte archiwum, którego nie trzeba czytać. Zasady V9 i nowsza rozmowa mają pierwszeństwo; stan testów i publikacji gry znajduje się w dokumentacji repozytorium.',
  ['Mechanika','Zasada w najwyżej trzech zdaniach','Źródło / status'],MECHANICS,[28,122,38])
 for r in range(6,m.max_row+1):m.row_dimensions[r].height=74
 b=sheet(w,'Budynki v9','BUDYNKI • jedna czytelna tabela','L oznacza obecny własny poziom budynku, N jego nowy poziom. Koszt kolejki: każdy składnik +10%, zaokrąglony w górę. T to zakończony rzeczywisty transport od pierwszego odbioru do ostatniej dostawy; kolumna budowania nie zawiera samego transportu. Wszystkie ceny bazowe pozostają zgodne z V8, z osobnymi nowymi zasadami usług i rezerwacji.',
  ['Budynek','EN','Po co jest','Ludzie: start / limit','Co daje poziom','Koszt ulepszenia','Sama budowa','Najważniejsza uwaga'],BUILDINGS,[23,22,40,27,46,54,26,70])
 future=[]
 for i,item in enumerate(manifest['futureOnly'],1):
  reasons=[
   'Reguły puli i potrąceń są zatwierdzone; krzywa czasu 180/(1+0,1×(runda−1)) jest wyłącznie konfiguracją prototypu.',
   'Nie podłączono odtwarzania ani weryfikacji reklam; nie pokazujemy fikcyjnego potwierdzenia obejrzenia.',
   'Zasady serii i końca programu są zatwierdzone; brak kwot blokuje wypłatę, nie liczenie dni.',
   'Nie ustalono czasu, miejsc ani efektu osobnej modlitwy; ręczne błogosławieństwo kaplicy jest gotową odrębną regułą.',
   'Reguły i czyste punkty integracji nie są prawdziwymi kontami, publikacjami, poleceniami ani transakcjami.',
  ]
  future.append((f'P{i:02}',item,reasons[i-1],'JAWNIE NA PÓŹNIEJ'))
 sheet(w,'Do decyzji','DO DECYZJI • tylko rzeczy pozostawione na później','Wszystkie 58 decyzji logicznych z rozmowy są rozstrzygnięte i zachowane w ukrytej zakładce Decyzje v9. Poniżej nie wracają stare pytania ani rozstrzygnięte konflikty. Ogłoszone przed GO założenia wdrożeniowe są jawnie opisane, a nie przedstawiane jako nieustalone reguły.',
  ['ID','Co pozostało','Co już wiadomo / granica realizacji','Status'],future,[12,86,110,27])
 detail=[]
 for d in manifest['decisions']:detail.append((f'V9-{d["id"]:02}',d['title'],d['rule'],d['status']))
 for d in manifest['frontend']:detail.append((d['id'],d['title'],d['rule'],d['status']))
 for k,v in manifest['implementationDefaults'].items():detail.append((k,'Jawne ustawienie projektu',str(v),'KONFIGURACJA / DOBÓR WDROŻENIOWY'))
 detail.extend([
  ('I01','Zależne poziomy przy anulowaniu','Najpierw anuluj późniejsze zależne poziomy tego samego budynku; nie anulujemy niczego niejawnie kaskadowo.','DOPRECYZOWANIE WDROŻENIOWE'),
  ('I02','Port: początek ważności','Zakupiona partia ma 24 h od przybycia statku i rozładunek nie odnawia terminu; partie sprzedażowe mają czas od każdego przybycia tragarza.','DOPRECYZOWANIE WDROŻENIOWE'),
  ('I03','Port: częściowe wygaśnięcie','Płatność tylko za faktycznie wysłaną ilość, proporcjonalnie do zamrożonej ceny, w dół do całych monet; zamówienie nadal wchodzi pełnymi paczkami 100.','DOPRECYZOWANIE WDROŻENIOWE'),
 ])
 d=sheet(w,'Decyzje v9','DECYZJE V9 • pełny manifest','Pełna treść docs/decisions-v9.json: 58 decyzji, 19 uwag frontendowych i jawne ustawienia projektu. To szczegółowy rejestr, nie kolejna lista pytań. Źródło V8 jest ukrytym archiwum i nie zastępuje tych nowszych rozstrzygnięć.',
  ['ID','Temat','Ostateczna zasada','Status'],detail,[28,42,128,55]);d.sheet_state='hidden'
 lessons=tutorial_rows()
 t=sheet(w,'Samouczek v9','SAMOUCZEK • 25 mikrozadań','Pięć lekcji przez pierwsze pięć aktywnych dni Europe/Warsaw. Kwoty dobrano projektowo przy wdrożeniu: dni 1–2 po 25 monet i 10 materiału, dni 3–5 po 50 monet i 20 materiału. Razem 1000 monet i 400 materiałów; nie zastępują nieustalonych nagród lojalnościowych.',
  ['ID','Dzień','Tytuł EN','Co uczy','Warunek','Temat','Poprzednik','Monety','Materiał','Sztuk','Status kwot'],lessons,[27,12,38,95,18,28,28,15,20,15,62]);t.sheet_state='hidden'
 # A single calculator with short, clearly separated blocks and blue inputs.
 k=sheet(w,'Kalkulator','KALKULATOR • jeden scenariusz zamiast wielu tabel','Niebieskie wartości możesz zmieniać. To przykłady, nie odczyt zapisu gry; założenia są podane przy wynikach. Krzywa czasu ekspedycji oraz kwoty samouczka są jawnie projektowe. Koszty transportu i budowy nie obejmują wcześniejszego zbierania zasobów ani czekania w kolejce.',
  ['Obszar','Parametr / wynik','Wartość','Jednostka','Znaczenie'],[],[26,54,25,22,100])
 def row(r,a,b,c,unit='',note='',input=False,fmt=None):
  for ci,v in enumerate([a,b,c,unit,note],1):cell_style(k.cell(r,ci,v),input=input and ci==3)
  k.row_dimensions[r].height=34
  if fmt:k.cell(r,3).number_format=fmt
 def section(r,title):
  for ci in range(1,6):cell_style(k.cell(r,ci,title if ci==1 else ''),header=True)
  k.row_dimensions[r].height=28
 section(6,'Transport i budowa')
 inputs=[(7,'Własny poziom magazynu',1,'poziom'),(8,'Ukończony poziom dróg',0,'0…5'),(9,'Dostępni magazynierzy',6,'osób'),(10,'Wyżywienie',1,'1 albo 0,5'),(11,'Drewno do przewozu',500,'szt.'),(12,'Kamień do przewozu',500,'szt.'),(13,'Zboże do przewozu',500,'szt.'),(14,'Żywność do przewozu',500,'szt.'),(15,'Deski do przewozu',500,'szt.'),(16,'Mnożnik fazy budowania',1,'1 /4 /3×(N−1)')]
 for r,title,value,unit in inputs:row(r,'Wejście',title,value,unit,'Scenariusz edytowalny',True)
 row(17,'Wynik','Ładunek na kurs','=ROUNDDOWN(ROUNDDOWN(60*1.1^(C7-1),0)*(1+0.1*MAX(0,C8-2)),0)','szt.','Własny poziom i dodatkowa pojemność dróg')
 row(18,'Wynik','Potrzebne kursy','=SUM(ROUNDUP(C11/C17,0),ROUNDUP(C12/C17,0),ROUNDUP(C13/C17,0),ROUNDUP(C14/C17,0),ROUNDUP(C15/C17,0))','kursów','Każdy rodzaj osobno; także ostatni niepełny')
 row(19,'Wynik','Normalny transport bez innych zadań','=IF(C9>0,ROUNDUP(C18/C9,0)*5/(1+0.1*C8)/C10,"brak obsady")','min','To estymata; rzeczywisty T obejmuje przerwy i zmiany obsady')
 row(20,'Wynik','Sama budowa z tego T','=IF(ISNUMBER(C19),C19*C16,"brak estymaty")','min')
 row(21,'Wynik','Transport + budowa','=IF(ISNUMBER(C19),C19+C20,"brak estymaty")','min')
 row(22,'Wynik','Instant transport: cena','=C18*10','monet','Wyłącznie wskazany etap dostawy')
 section(24,'Premium dowóz w trakcie transportu')
 row(25,'Wejście','Rzeczywisty miniony T',10,'min','Od pierwszego odbioru, także przerwy',True)
 row(26,'Wejście','Normalna estymata reszty',5,'min','Przy 0 osób użyj jednego hipotetycznego tragarza',True)
 row(27,'Wynik','Zamrożona sama budowa','=(C25+C26)*C16','min','Dostawa jest natychmiastowa; budowanie nie zeruje się')
 section(29,'Prognozy FIFO — ten sam kamień tylko raz')
 row(30,'Wejście','Kamień fizycznie dostępny',200,'szt.','Przed pierwszym planem',True)
 row(31,'Wejście','Bazowy koszt każdego planu',300,'szt.','Dwa kolejne plany tego samego materiału',True)
 row(32,'Wynik','Każdy plan po kolejce +10%','=ROUNDUP(C31*1.1,0)','szt.')
 row(33,'Wynik','Rezerwacja pierwszego planu','=MIN(C30,C32)','szt.','Nadal zajmuje miejsce do odbioru')
 row(34,'Wynik','Brak pierwszego planu','=MAX(0,C32-C33)','szt.')
 row(35,'Wynik','Dopłata pierwszego','=ROUNDUP(C34*0.5*0.25,0)','monet')
 row(36,'Wynik','Rezerwacja drugiego','=MIN(MAX(0,C30-C33),C32)','szt.')
 row(37,'Wynik','Brak drugiego','=MAX(0,C32-C36)','szt.')
 row(38,'Wynik','Dopłata drugiego','=ROUNDUP(C37*0.5*0.25,0)','monet')
 row(39,'Wynik','Obie dopłaty','=C35+C38','monet','Osobne zakupy zaokrąglane osobno;17+42=59')
 section(41,'Atrakcja: rzeczywiści odbiorcy')
 row(42,'Wejście','Poziom Ratusza',4,'poziom','Maksymalna liczba odbiorców',True)
 row(43,'Wejście','Uprawnieni odbiorcy',2,'osób','Bez100%H i bez zamrożonych nowych osób',True)
 row(44,'Wejście','Wylosowane Y',50,'5…100','Zapisane w ofercie',True)
 row(45,'Wejście','Kolejny zakup tego rodzaju',2,'1…5','Reset po 24 h bez zakupu tego rodzaju',True)
 row(46,'Wynik','Faktyczni odbiorcy','=MIN(C42,C43)','osób')
 row(47,'Wynik','Skuteczność','=MAX(0,1-(C45-1)*0.25)','udział',fmt='0%')
 row(48,'Wynik','Pobierana cena','=IF(AND(C46>0,C47>0),C44*C46,0)','monet','Przy 0 odbiorców lub 0 skuteczności zakup zablokowany')
 row(49,'Wynik','Dokładna premia na osobę','=C44*C47','pp','Na 6 h; końcowe osobiste H nadal ograniczone do ±100')
 section(51,'Ekspedycja: +50 monet, nie podwojenie')
 row(52,'Wejście','Faktyczna poprzednia pula',170,'monet','Po wcześniejszych potrąceniach',True)
 row(53,'Wejście','Spóźnienie pełnych sekund',3,'s','Po końcu limitu rundy',True)
 row(54,'Wynik','Potencjalna pula kolejnej wygranej','=C52+50','monet')
 row(55,'Wynik','Pula po potrąceniu','=MAX(0,C54-10*ROUNDDOWN(C53,0))','monet','Przy 0 seria przegrywa')
 row(56,'Wejście','Numer rundy do krzywej',3,'runda','Krzywa docelowa pozostaje na później',True)
 row(57,'PROJEKT','Przykładowy limit prototypu','=180/(1+0.1*(C56-1))','s','KONFIGURACJA, nie zatwierdzona docelowa trudność',fmt='0.00')
 section(59,'Ścieżki i kupno osoby')
 row(60,'Wejście','Liczba odcinków poziomu',100,'odcinków','Maksimum 100; kolejny poziom po pełnej setce',True)
 row(61,'Wynik','Kamień na ukończenie poziomu','=C60*100','kamienia')
 row(62,'Wynik','Wszystkie pięć poziomów','=100*100*5','kamienia')
 row(63,'Wejście','Obecny poziom Ratusza',2,'poziom','Ogłoszone założenie kosztu jednej osoby',True)
 row(64,'ZAŁOŻENIE','Cena jednej osoby — każdego zasobu','=500*C63','każdego z 6','Pełny koszt najbliższego Ratusza; limity łóżek i populacji')
 section(66,'Samouczek i Punkty')
 row(67,'PROJEKT','Łączne monety 25 lekcji',"=SUM('Samouczek v9'!H6:H30)",'monet','Kwoty dobrane projektowo; nie są nagrodami lojalności')
 row(68,'PROJEKT','Łączne materiały 25 lekcji',"=SUM('Samouczek v9'!J6:J30)",'szt.','Suma różnych materiałów; bez monet')
 row(69,'Wejście','Monety w chwili pełnej godziny',1000,'monet','Migawka nie zmienia się po każdej transakcji',True)
 row(70,'Wynik','Punkty z migawki','=C63*C69','punktów')
 section(72,'Miejsce w centrali')
 row(73,'Wejście','Poziom Magazynu',1,'poziom','Limit na rodzaj i łącznie rosną ×1,5',True)
 row(74,'Wynik','Limit na materiał','=ROUNDDOWN(1000*1.5^(C73-1),0)','szt.')
 row(75,'Wynik','Limit łączny','=ROUNDDOWN(4000*1.5^(C73-1),0)','szt.')
 row(76,'Wynik','Miejsce wyłącznie na Food','=50*C73','szt.','Nie mogą go zajmować inne towary, zwroty ani nagrody')
 row(77,'Wynik','Maksimum pozostałych materiałów','=C75-C76','szt.','Fizyczne rezerwacje do odbioru nadal zajmują tę przestrzeń')
 row(79,'Uwaga','Prawdziwe płatności i integracje','Niepodłączone','','Kalkulator nie tworzy kont, transakcji ani zewnętrznych publikacji')
 k.auto_filter.ref=None
 # Input ranges: explicit bounds prevent accidental impossible calculator scenarios.
 for cell,minv,maxv in [('C7',1,100),('C8',0,5),('C9',0,1200),('C42',1,100),('C43',0,2400),('C44',5,100),('C45',1,5),('C56',1,10000),('C60',0,100),('C63',1,100),('C73',1,100)]:
  dv=DataValidation(type='whole',operator='between',formula1=minv,formula2=maxv);dv.showErrorMessage=True;dv.error='Wpisz wartość w podanym zakresie.';k.add_data_validation(dv);dv.add(k[cell])
 checks=[]
 def check(title,formula,expected):
  r=6+len(checks);checks.append((title,formula,expected,f'=B{r}=C{r}'))
 for cell,expected,title in [('C17',60,'Ładunek bazowy'),('C18',45,'Pięć rodzajów po500'),('C19',40,'Transport45kursów/6osób'),('C20',40,'Budowanie standardowe'),('C21',80,'Cały scenariusz'),('C22',450,'Instant45kursów'),('C27',15,'PremiumT10+5'),('C35',17,'Prognoza1'),('C38',42,'Prognoza2'),('C39',59,'Brak podwójnego zastawu'),('C46',2,'Rzeczywiści odbiorcy'),('C47',.75,'Drugie użycie'),('C48',100,'CenaY×2osoby'),('C49',37.5,'Dokładna ułamkowa premia'),('C54',220,'170+50'),('C55',190,'3sekundyspóźnienia'),('C57',150,'Krzywa projektowa runda3'),('C61',10000,'Kamieńpoziom'),('C62',50000,'Kamień0→5'),('C64',1000,'CenaosobyTH2każdego'),('C67',1000,'Samouczekmonety'),('C68',400,'Samouczekmateriały'),('C70',2000,'PunktyTH2×1000'),('C74',1000,'CentralaL1per'),('C75',4000,'CentralaL1total'),('C76',50,'MiejscenaFood'),('C77',3950,'Innetowarymaks')]:check(title,f'=Kalkulator!{cell}',expected)
 check('58decyzji','=COUNTA(\'Decyzje v9\'!A6:A63)',58)
 check('19uwagfrontu','=COUNTA(\'Decyzje v9\'!A64:A82)',19)
 check('25lekcji','=COUNTA(\'Samouczek v9\'!A6:A30)',25)
 q=sheet(w,'Sprawdzenia v9','SPRAWDZENIA V9 • kontrola liczb','Formuły sprawdzają arytmetykę przykładowego kalkulatora i kompletność manifestu. Wynik TRUE nie zastępuje testów silnika ani publikacji gry. Kolumny i historyczne źródła są zachowane osobno.',
  ['Przykład','Wynik','Oczekiwane','Zgodność'],checks,[55,28,28,20]);q.sheet_state='hidden'
 w._sheets=[w[n] for n in VISIBLE]+[s for s in w if s.title not in VISIBLE]
 w.active=0;w.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
 w.properties.title='HARBORLIGHT9 — uproszczone zasady po58decyzjach';w.properties.subject='Cztery widoczne zakładki; zachowane źródła, komentarze i kalkulatory';w.properties.description='Wersja9. docs/decisions-v9.json ma pierwszeństwo przed ukrytym archiwumV8.';w.properties.modified=datetime(2026,10,23,16,0)
 w.save(OUTPUT)


def recalc():
 program=shutil.which('libreoffice') or shutil.which('soffice')
 if not program:raise RuntimeError('LibreOffice required for cachedformularesults')
 with tempfile.TemporaryDirectory(prefix='harborlight-v9-lo-') as tmp:
  d=Path(tmp);out=d/'out';out.mkdir()
  result=subprocess.run([program,f'-env:UserInstallation={(d/"profile").as_uri()}','--headless','--convert-to','xlsx','--outdir',str(out),str(OUTPUT)],capture_output=True,text=True,timeout=120)
  file=out/OUTPUT.name
  if result.returncode or not file.exists():raise RuntimeError(result.stdout+'\n'+result.stderr)
  OUTPUT.write_bytes(file.read_bytes());print(result.stdout.strip())


def validate(cached=True):
 assert sha(OLD7)==HASH7 and sha(BASE)==HASH8
 source=load_workbook(BASE);result=load_workbook(OUTPUT);values=load_workbook(OUTPUT,data_only=True)
 assert [s.title for s in result if s.sheet_state=='visible']==VISIBLE
 count_comments=0
 for s in source:
  assert result[s.title].sheet_state=='hidden'
  for row in s:
   for cell in row:
    actual=result[s.title][cell.coordinate]
    assert actual.value==cell.value,(s.title,cell.coordinate,'source value changed')
    if cell.comment:
     count_comments+=1;assert actual.comment and actual.comment.text==cell.comment.text and actual.comment.author==cell.comment.author,(s.title,cell.coordinate,'comment changed')
 assert count_comments==48
 manifest=json.loads(MANIFEST.read_text())
 for i,decision in enumerate(manifest['decisions'],6):
  assert result['Decyzje v9'].cell(i,1).value==f'V9-{decision["id"]:02}'
  assert result['Decyzje v9'].cell(i,3).value==decision['rule'],('missingdecision',decision['id'])
 for i,item in enumerate(manifest['frontend'],64):
  assert result['Decyzje v9'].cell(i,1).value==item['id']
  assert result['Decyzje v9'].cell(i,3).value==item['rule'],('missingfrontend',item['id'])
 for title,explanation,_ in MECHANICS:
  assert len([p for p in re.split(r'[.!?](?:\s|$)',explanation) if p.strip()])<=3,('morethan3sentences',title)
 errors=[];formulas=0
 for s in result:
  for row in s:
   for c in row:
    if c.data_type=='f':
     formulas+=1;v=values[s.title][c.coordinate]
     if cached and (v.value is None or v.data_type=='e'):errors.append((s.title,c.coordinate,v.value))
 if cached:
  for row in range(6,values['Sprawdzenia v8'].max_row+1):
   if values['Sprawdzenia v8'].cell(row,4).value is not True:errors.append(('historicalcheck',row))
  for row in range(6,values['Sprawdzenia v9'].max_row+1):
   if values['Sprawdzenia v9'].cell(row,4).value is not True:errors.append(('check',row,values['Sprawdzenia v9'].cell(row,1).value))
 assert not errors,errors[:20]
 with ZipFile(OUTPUT) as z:assert z.testzip() is None
 manifest=json.loads(MANIFEST.read_text())
 report={'file':str(OUTPUT.relative_to(ROOT)),'bytes':OUTPUT.stat().st_size,'sha256':sha(OUTPUT),'visible_sheets':VISIBLE,'total_sheets':len(result.sheetnames),'historical_sheets_unchanged':len(source.sheetnames),'comments_preserved':count_comments,'complete_manifest_58_plus_19':True,'mechanics_at_most_3_sentences':True,'manifest_decisions':len(manifest['decisions']),'frontend_requirements':len(manifest['frontend']),'mechanics':len(MECHANICS),'tutorial_lessons':25,'formulas':formulas,'new_checks':values['Sprawdzenia v9'].max_row-5,'cached_and_checks_pass':cached,'source7_unchanged':sha(OLD7)==HASH7,'source8_unchanged':sha(BASE)==HASH8}
 print(json.dumps(report,ensure_ascii=False,indent=2));return report


def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--no-recalc',action='store_true');p.add_argument('--validate-only',action='store_true');a=p.parse_args()
 if not a.validate_only:
  build()
  if not a.no_recalc:recalc()
 validate(not a.no_recalc)
if __name__=='__main__':main()
